"""Quantity-based location allocations; fixed costs are apportioned, never duplicated."""
import json, math
FIELDS={'material':'total_material','hours':'total_hours','labor':'total_labor','quoted':'quoted_total'}
DEFAULT='Unallocated'
COMMON='Common to All Areas'
RESERVED={DEFAULT.casefold(),COMMON.casefold(),'shared / unallocated'}
def allocations(value):
    if isinstance(value,str):
        try:value=json.loads(value)
        except ValueError:raise ValueError('Invalid area allocation JSON') from None
    if not isinstance(value,list) or len(value)>200:raise ValueError('Use up to 200 areas per item')
    result=[];seen=set()
    for r in value:
        if not isinstance(r,dict):raise ValueError('Invalid area allocation')
        name=' '.join(str(r.get('area','')).split());q=r.get('quantity')
        if name.casefold()==COMMON.casefold():name=COMMON
        if not name or len(name)>120:raise ValueError('Area names must be 1–120 characters')
        if name.casefold() in seen:raise ValueError('An area can appear only once per item')
        if type(q) not in (float,int) or not math.isfinite(q) or not 0<=q<=1e9:raise ValueError('Area quantities must be finite and non-negative')
        if r.get('common') is True and name.casefold()!=COMMON.casefold():raise ValueError('Common item must use the Common to All Areas allocation')
        seen.add(name.casefold());result.append({'area':name,'quantity':float(q),**({'shared':True} if r.get('shared') is True else {}),**({'common':True} if name.casefold()==COMMON.casefold() or r.get('common') is True and name==COMMON else {})})
    if result and sum(r['quantity'] for r in result)<=0:raise ValueError('Allocated items need a positive total quantity; clear allocations for a zero-quantity item')
    return result

def parts(line):
    rows=allocations(line.get('area_allocations','[]'))
    total=sum(r['quantity'] for r in rows)
    return [(r['area'],r['quantity'],r['quantity']/total) for r in rows] if rows else [(DEFAULT,line['quantity'],1)]

def summary(lines, area_names=None):
    groups={}
    for name in area_names or []:
        groups[name.casefold()]={'area':name,'included':dict.fromkeys(FIELDS,0.),'optional':dict.fromkeys(FIELDS,0.),'items':[]}
    for line in lines:
        for name,qty,weight in parts(line):
            group=groups.setdefault(name.casefold(),{'area':name,'included':dict.fromkeys(FIELDS,0.),'optional':dict.fromkeys(FIELDS,0.),'items':[]})
            values={k:line[f]*weight for k,f in FIELDS.items()}
            target=group['optional' if line.get('included_in_total') is False else 'included']
            for k in FIELDS:target[k]+=values[k]
            group['items'].append({'id':line['id'],'name':line['name_snapshot'],'category':line['category_key'],'quantity':qty,'uom':line['uom'],'optional':line.get('included_in_total') is False,**values})
    for group in groups.values():
        categories={}
        for item in group['items']:
            key=item['category']
            category=categories.setdefault(key,{'category':key,'items':[], 'included':dict.fromkeys(FIELDS,0.),'optional':dict.fromkeys(FIELDS,0.)})
            category['items'].append(item)
            target=category['optional' if item['optional'] else 'included']
            for field in FIELDS:target[field]+=item[field]
        group['categories']=list(categories.values())
        group['item_count']=len(group['items'])
        group['quantity_total']=sum(item['quantity'] for item in group['items'])
        units={}
        for item in group['items']:units[item['uom']]=units.get(item['uom'],0)+item['quantity']
        group['quantity_by_unit']=units
    return list(groups.values())


def initialize_breakdown_areas(conn):
    conn.execute('''CREATE TABLE IF NOT EXISTS project_breakdown_areas(
        project_id INTEGER NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
        area TEXT NOT NULL COLLATE NOCASE,
        sort_order INTEGER NOT NULL,
        PRIMARY KEY(project_id,area))''')


def names_for_project(conn,pid,lines=None):
    saved=[row[0] for row in conn.execute('SELECT area FROM project_breakdown_areas WHERE project_id=? ORDER BY sort_order,area',(pid,))]
    seen={s.casefold() for s in saved}
    if lines is None:
        lines=conn.execute('SELECT area_allocations FROM project_lines WHERE project_id=? ORDER BY sort_order,id',(pid,))
    for line in lines:
        for row in allocations(line['area_allocations'] if isinstance(line,dict) or 'area_allocations' in line.keys() else line[0]):
            name=row['area']
            if name.casefold() not in seen and name.casefold() not in RESERVED:
                saved.append(name);seen.add(name.casefold())
    return saved


def checked_names(value):
    if not isinstance(value,list) or len(value)>200:raise ValueError('Choose up to 200 areas')
    seen=set();names=[]
    for raw in value:
        if not isinstance(raw,str):raise ValueError('Area name must be text')
        name=' '.join(raw.split())
        if not name or len(name)>120 or name.casefold() in RESERVED or name.casefold() in seen:
            raise ValueError('Area names must be unique, 1-120 characters and not reserved')
        seen.add(name.casefold());names.append(name)
    return names


def save_breakdown(conn,pid,user,payload):
    """Atomic project Breakdown update. Never rewrite unchanged line costs or catalogue."""
    import sqlite3
    from .workspace import require_project, check_revision, record_activity
    from .projects import calculate_project
    from .utils import utc_now
    project=require_project(conn,pid,user,'edit')
    if 'expected_revision' not in payload:raise ValueError('Refresh the estimate before saving its breakdown')
    check_revision(project,payload)
    names=checked_names(payload.get('areas'))
    raw=payload.get('rows')
    if not isinstance(raw,list) or len(raw)>2000:raise ValueError('Invalid breakdown items')
    existing={r['id']:r for r in conn.execute('SELECT id,quantity,area_allocations,revision,category_snapshot FROM project_lines WHERE project_id=?',(pid,))}
    if len(raw)!=len(existing):raise ValueError('Estimate items have changed. Refresh Breakdown and retry.')
    changed=[];seen=set();tolerance=1e-7
    for item in raw:
        if not isinstance(item,dict) or type(item.get('id')) is not int or item['id'] not in existing or item['id'] in seen:
            raise ValueError('Invalid or repeated breakdown item; refresh and retry')
        lid=item['id'];seen.add(lid);old=existing[lid]
        if type(item.get('expected_revision')) is not int or old['revision']!=item['expected_revision']:
            raise ValueError('An item has changed; refresh the breakdown before saving')
        q=item.get('quantity')
        if type(q) not in (int,float) or not math.isfinite(q) or not 0<=q<=1e9:
            raise ValueError('Main quantities must be finite and non-negative')
        assigned=allocations(item.get('allocations',[]))
        if assigned and not math.isclose(sum(row['quantity'] for row in assigned),q,rel_tol=0,abs_tol=tolerance):
            raise ValueError('Allocation does not match main quantity for item '+str(lid))
        if q==0 and assigned:raise ValueError('Zero-quantity items cannot have allocations')
        for allocation in assigned:
            area=allocation['area'];key=area.casefold()
            if key=='shared / unallocated' and assigned==allocations(old['area_allocations']):
                continue  # Preserve an unchanged legacy allocation, without reinterpreting pricing.
            if key!=COMMON.casefold() and key not in {n.casefold() for n in names}:
                raise ValueError('An allocation references a deleted area: '+area)
            if key==COMMON.casefold() and len(assigned)!=1:
                raise ValueError('Common to all areas cannot be mixed with individual area quantities')
        before=allocations(old['area_allocations'])
        if not math.isclose(q,old['quantity'],abs_tol=tolerance,rel_tol=0) or assigned!=before:
            changed.append((lid,q,assigned,old))
    if not names and any(r for r in raw if r.get('allocations') and any(a.get('area')!=COMMON for a in r['allocations'])):
        raise ValueError('Add an area for allocated estimate items')
    # Verify everything before mutating. The HTTP connection is a write transaction.
    conn.execute('DELETE FROM project_breakdown_areas WHERE project_id=?',(pid,))
    conn.executemany('INSERT INTO project_breakdown_areas(project_id,area,sort_order) VALUES(?,?,?)',((pid,name,i) for i,name in enumerate(names)))
    stamp=utc_now()
    for lid,q,assigned,old in changed:
        # Lighting's special supply quantity override is normalized on main-quantity edits,
        # identically to the regular estimate item editor.
        reset_supply=(old['category_snapshot'] in ('Lighting','Emergency & Exit Lighting')
                      and not math.isclose(q,old['quantity'],abs_tol=tolerance,rel_tol=0))
        if reset_supply:
            conn.execute('''UPDATE project_lines SET quantity=?,area_allocations=?,supply_quantity_override=NULL,
                revision=revision+1,updated_at=? WHERE id=? AND project_id=?''',
                (q,json.dumps(assigned),stamp,lid,pid))
        else:
            conn.execute('''UPDATE project_lines SET quantity=?,area_allocations=?,revision=revision+1,
                updated_at=? WHERE id=? AND project_id=?''',
                (q,json.dumps(assigned),stamp,lid,pid))
    conn.execute('UPDATE projects SET updated_at=? WHERE id=?',(stamp,pid))
    record_activity(conn,pid,user,'Updated area breakdown',f'{len(names)} areas, {len(changed)} changed items',touch=False)
    return calculate_project(conn,pid)


def excel(wb,line_parent_rows,area_names=None):
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.comments import Comment
    ws=wb.create_sheet('AREA ALLOCATIONS');overview=wb.create_sheet('AREA SUMMARY')
    ws.append(['Area / floor / room','Item','Category','Quantity (edit here)','Material','Hours','Labour','Quoted','Included / Optional'])
    area_rows={};records=[];detail=wb['CONNECTED COST DETAIL']
    for name in area_names or []:area_rows.setdefault(name.casefold(),(name,[]))
    for line,parent in line_parent_rows:
        entries=parts(line);start=ws.max_row+1;end=start+len(entries)-1
        if allocations(line.get('area_allocations','[]')):detail.cell(parent,4,f"=SUM('AREA ALLOCATIONS'!D{start}:D{end})")
        if allocations(line.get('area_allocations','[]')):detail.cell(parent,4).comment=Comment('Quantity is linked to AREA ALLOCATIONS. Edit the area quantities there. Fixed-total components are distributed proportionally, once per item.','TEC')
        for name,qty,_ in entries:
            allocated=bool(allocations(line.get('area_allocations','[]')))
            row=ws.max_row+1;area_rows.setdefault(name.casefold(),(name,[]))[1].append(row)
            ws.append([name,line['name_snapshot'],line['category_key'],qty,None,None,None,None,'Optional' if line.get('included_in_total') is False else 'Included'])
            records.append((name,line['category_key'],line['name_snapshot'],line['uom'],row))
            # When all allocations are zero, retain fixed costs in the first area.
            weight=f'IF(SUM($D${start}:$D${end})=0,{1 if row==start else 0},D{row}/SUM($D${start}:$D${end}))'
            for col,source in [(5,'G'),(6,'I'),(7,'J'),(8,'L')]:ws.cell(row,col,f"='CONNECTED COST DETAIL'!{source}{parent}*{weight}")
            ws.cell(row,1).data_type='s'
            ws.cell(row,2,f"='CONNECTED COST DETAIL'!B{parent}")
            ws.cell(row,3,f"='CONNECTED COST DETAIL'!A{parent}")
            ws.cell(row,9,f"=IF('CONNECTED COST DETAIL'!P{parent}=1,\"Included\",\"Optional\")")
            if not allocated:ws.cell(row,4,f"='CONNECTED COST DETAIL'!D{parent}")
            ws.cell(row,4).fill=PatternFill('solid',fgColor='FFF7E4')
    rule=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1=0);rule.errorTitle='Invalid quantity';rule.error='Enter zero or a positive quantity.';rule.showErrorMessage=True;rule.errorStyle='stop';ws.add_data_validation(rule)
    if ws.max_row>1:rule.add(f'D2:D{ws.max_row}')
    overview.append(['Area / floor / room','Material','Hours','Labour','Quoted total','Optional material','Optional hours','Optional labour','Optional quoted'])
    for name,rows in area_rows.values():
        row=overview.max_row+1;overview.cell(row,1,name).data_type='s'
        for offset,status in [(0,'Included'),(4,'Optional')]:
            for col,source in [(2,'E'),(3,'F'),(4,'G'),(5,'H')]:
                n=ws.max_row
                overview.cell(row,col+offset,f"=SUMPRODUCT(--('AREA ALLOCATIONS'!$A$2:$A${n}=$A{row}),--('AREA ALLOCATIONS'!$I$2:$I${n}=\"{status}\"),'AREA ALLOCATIONS'!${source}$2:${source}${n})")
    last=overview.max_row;overview.append(['TOTAL'])
    for col in range(2,10):overview.cell(last+1,col,f'=SUM({get_column_letter(col)}2:{get_column_letter(col)}{last})' if last>1 else '=0')
    for sheet in (ws,overview):
        sheet.freeze_panes='D2' if sheet==ws else 'B2';sheet.auto_filter.ref=sheet.dimensions
        sheet.sheet_view.showGridLines=False
        for c in sheet[1]:c.fill=PatternFill('solid',fgColor='173B55');c.font=Font(color='FFFFFF',bold=True)
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if c.column>1:c.number_format='#,##0.00'
        for col in range(1,10):sheet.column_dimensions[get_column_letter(col)].width=32 if col==1 or (sheet==ws and col in (2,3)) else 19
        sheet.print_title_rows='1:1';sheet.page_setup.orientation='landscape';sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0

    structured_excel(wb,records,list(area_names or []))
    main=wb['COST ESTIMATE'];start=main.max_row+3
    main.cell(start,1,'BREAKDOWN COST SUMMARY').font=Font(bold=True,color='087F75',size=14)
    for col,title in enumerate(['Breakdown','Material','Hours','Labour','Quoted total','Optional material','Optional hours','Optional labour','Optional quoted'],1):main.cell(start+1,col,title).font=Font(bold=True)
    for r in range(2,overview.max_row+1):
        for col in range(1,10):
            c=main.cell(start+r,col,f"='AREA SUMMARY'!{get_column_letter(col)}{r}");c.number_format='#,##0.00' if col>1 else 'General'
    main.print_area=f'A1:N{start+overview.max_row}'



def structured_excel(wb,records,area_names):
    """Formula-linked hierarchical workbook: AREA > CATEGORY > ITEM, common priced once."""
    from collections import OrderedDict
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    ws=wb.create_sheet('STRUCTURED BREAKDOWN')
    ws.append(['Area / category / item','Item qty','Unit','Included material','Included hours','Included labour','Included quoted','Optional quoted','Cost treatment','Optional material','Optional hours','Optional labour'])
    ws.freeze_panes='D2';ws.sheet_view.showGridLines=False
    ws.row_dimensions[1].height=36
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='173B55')
        cell.font=Font(name='Aptos',size=10,bold=True,color='FFFFFF')
        cell.alignment=Alignment(vertical='center',wrap_text=True)
    ws.column_dimensions['A'].width=53;ws.column_dimensions['B'].width=15;ws.column_dimensions['C'].width=12
    for col in 'DEFGHIJKL':ws.column_dimensions[col].width=20
    colors={'area':'173B55','category':'E5F4F0','item':'FFFFFF','total':'087F75'}
    grouped=OrderedDict((name,OrderedDict()) for name in area_names)
    for name,cat,title,uom,source in records:
        grouped.setdefault(name,OrderedDict()).setdefault(cat,[]).append((title,uom,source))
    first=ws.max_row+1;area_rows=[]
    for area,categories in grouped.items():
        area_row=ws.max_row+1;area_rows.append(area_row)
        ws.cell(area_row,1,area)
        cat_rows=[]
        for cat,items in categories.items():
            cat_row=ws.max_row+1;cat_rows.append(cat_row)
            ws.cell(cat_row,1,'    '+cat)
            item_rows=[]
            for name,uom,source in items:
                row=ws.max_row+1;item_rows.append(row)
                ws.cell(row,1,'        '+name)
                ws.cell(row,2,f"='AREA ALLOCATIONS'!D{source}")
                ws.cell(row,3,uom)
                ws.cell(row,9,f"='AREA ALLOCATIONS'!I{source}")
                for target,source_col,status in [(4,'E','Included'),(5,'F','Included'),(6,'G','Included'),(7,'H','Included'),(8,'H','Optional'),(10,'E','Optional'),(11,'F','Optional'),(12,'G','Optional')]:
                    ws.cell(row,target,f'=IF(I{row}="{status}",' + f"'AREA ALLOCATIONS'!{source_col}{source},0)")
                ws.cell(row,1).alignment=Alignment(indent=2,vertical='center',wrap_text=True)
                ws.row_dimensions[row].height=24
                if ws.cell(row,9).value:
                    ws.cell(row,9).font=Font(name='Aptos',size=9,color='6D7284')
            for col in (4,5,6,7,8,10,11,12):
                letter=get_column_letter(col)
                ws.cell(cat_row,col,'=SUM('+','.join(letter+str(r) for r in item_rows)+')' if item_rows else '=0')
            ws.cell(cat_row,1).font=Font(name='Aptos',bold=True,color='065F58')
            for cell in ws[cat_row]:cell.fill=PatternFill('solid',fgColor=colors['category'])
            ws.row_dimensions[cat_row].height=27
        for col in (4,5,6,7,8,10,11,12):
            letter=get_column_letter(col)
            ws.cell(area_row,col,'=SUM('+','.join(letter+str(r) for r in cat_rows)+')' if cat_rows else '=0')
        for cell in ws[area_row]:cell.fill=PatternFill('solid',fgColor=colors['area']);cell.font=Font(name='Aptos',bold=True,color='FFFFFF')
        ws.row_dimensions[area_row].height=30
    row=ws.max_row+2
    ws.cell(row,1,'PROJECT TOTAL (COMMON COUNTED ONCE)')
    for col in (4,5,6,7,8,10,11,12):
        letter=get_column_letter(col)
        ws.cell(row,col,'=SUM('+','.join(letter+str(r) for r in area_rows)+')' if area_rows else '=0')
    for cell in ws[row]:cell.fill=PatternFill('solid',fgColor=colors['total']);cell.font=Font(name='Aptos',bold=True,color='FFFFFF')
    ws.row_dimensions[row].height=34
    ws.cell(row+2,1,'Note: Common to All Areas is priced once. Optional totals are shown separately from the included quoted total.')
    ws.merge_cells(start_row=row+2,start_column=1,end_row=row+2,end_column=12)
    ws.cell(row+2,1).font=Font(name='Aptos',size=10,italic=True,color='64748B')
    for r in ws.iter_rows(min_row=2,max_row=row):
        for cell in r:
            if cell.column in (4,6,7,8,10,12):cell.number_format='#,##0.00;[Red](#,##0.00);-'
            if cell.column in (2,5,11):cell.number_format='#,##0.00;[Red](#,##0.00);-'
            cell.alignment=Alignment(vertical='center',wrap_text=cell.column==1)
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_options.horizontalCentered=True
    ws.print_area=f'A1:L{row+2}'
    return ws
