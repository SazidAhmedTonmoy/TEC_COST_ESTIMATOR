from __future__ import annotations

import argparse
import json
import re
import sqlite3
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from .database import connect, initialize, refresh_catalog_index, write_lock
from .utils import as_float, clean_text, json_dumps, normalize_text, sha256_file, stable_hash, utc_now


HEADER_ALIASES = {
    "item": "item_no",
    "qty": "qty",
    "u m": "uom",
    "q m": "q_m",
    "size": "size",
    "description": "description",
    "material unit": "unit_material",
    "mat factor": "material_factor",
    "material result": "material_result",
    "labor unit": "unit_hours",
    "labour unit": "unit_hours",
    "labor factor": "labor_factor",
    "labour factor": "labor_factor",
    "labor result": "labor_result",
    "labour result": "labor_result",
}

CODE_PATTERNS = [
    re.compile(r"\bTYPE\s+([A-Z]{1,8}[-/]?[A-Z0-9]*\d[A-Z0-9./-]*)\b", re.I),
    re.compile(r"\b(CID\s*[-#]?\s*\d+[A-Z]?)\b", re.I),
    re.compile(r"\b(DB[- ]?[A-Z0-9]+)\b", re.I),
    re.compile(r"\b([A-Z]{1,6}[-/]?\d+[A-Z0-9./-]*)\b", re.I),
]


@dataclass
class ParsedComponent:
    line_order: int
    item_no: str
    description: str
    quantity_per_unit: float
    source_quantity: float
    uom: str = ""
    q_m: str = ""
    size: str = ""
    unit_material: float = 0.0
    material_factor: float = 1.0
    unit_hours: float = 0.0
    labor_factor: float = 1.0
    is_parent: int = 0
    source_row: int | None = None
    quantity_formula: str = ""
    notes: str = ""
    labor_pricing: str = 'hours'
    unit_labor_price: float = 0.0


@dataclass
class ParsedAssembly:
    code: str
    name: str
    description: str
    phase: str
    subphase: str
    scope: str
    base_quantity: float
    base_uom: str
    source_row: int
    source_type: str
    components: list[ParsedComponent] = field(default_factory=list)
    source_material_total: float = 0.0
    source_labor_hours: float = 0.0
    material_markup_hint: float | None = None
    labor_rate_hint: float | None = None
    overhead_markup_hint: float | None = None
    work_type: str = "assembly"
    category_override: str = ""


@dataclass
class ParsedIssue:
    severity: str
    issue_code: str
    message: str
    row_no: int | None = None
    payload: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedSheet:
    sheet_name: str
    project_name: str
    import_format: str
    row_count: int
    header_count: int
    assemblies: list[ParsedAssembly] = field(default_factory=list)
    issues: list[ParsedIssue] = field(default_factory=list)


def _cell_value(ws: Any, row: int, col: int | None) -> Any:
    return ws.cell(row, col).value if col else None


def _formula_text(value: Any) -> str:
    if isinstance(value, str) and value.startswith("="):
        return value
    text = getattr(value, "text", "")
    return text if isinstance(text, str) else ""


def _row_texts(ws: Any, row: int) -> list[str]:
    return [clean_text(cell.value) for cell in ws[row] if clean_text(cell.value)]


def _header_map(ws: Any, row: int) -> dict[str, int] | None:
    result: dict[str, int] = {}
    for cell in ws[row]:
        label = normalize_text(cell.value)
        if label in HEADER_ALIASES:
            result[HEADER_ALIASES[label]] = cell.column
    if "description" in result and "qty" in result and "item_no" in result:
        labor_factor = result.get("labor_factor")
        if labor_factor and "labor_result" not in result:
            result["labor_result"] = labor_factor + 2
        return result
    return None


def _project_name(ws: Any) -> str:
    for row in range(1, min(ws.max_row, 40) + 1):
        for col in range(1, ws.max_column + 1):
            text = normalize_text(ws.cell(row, col).value)
            if text in {"project", "project name"}:
                for candidate_col in range(col + 1, min(ws.max_column, col + 6) + 1):
                    candidate = clean_text(ws.cell(row, candidate_col).value)
                    if candidate:
                        return candidate
    return ""


def _extract_phase(texts: Iterable[str], prefix: str) -> str | None:
    pattern = re.compile(rf"^{re.escape(prefix)}\s*:\s*(.+)$", re.I)
    for text in texts:
        match = pattern.match(text)
        if match:
            return clean_text(match.group(1))
    return None


def _extract_code(text: str) -> str:
    upper = clean_text(text).upper()
    for pattern in CODE_PATTERNS:
        match = pattern.search(upper)
        if match:
            code = re.sub(r"\s+", "", match.group(1))
            return code.replace("CID-", "CID") if code.startswith("CID-") else code
    return ""


def _clean_assembly_name(context: str, title_description: str) -> str:
    name = clean_text(context)
    name = re.sub(r"\(\s*QTY\s*[^)]*\)", "", name, flags=re.I)
    name = re.sub(r"^TYPE\s+", "", name, flags=re.I)
    code = _extract_code(name)
    if code:
        name = re.sub(re.escape(code), "", name, count=1, flags=re.I)
    name = re.sub(r"^[\s:_-]+", "", name).strip()
    if not name or normalize_text(name) in {"allowance", "installation", "supply installation"}:
        name = clean_text(title_description)
    return name or "Unnamed ConEst assembly"


def _component_from_row(
    value_ws: Any,
    formula_ws: Any,
    row: int,
    columns: dict[str, int],
    base_quantity: float,
    line_order: int,
    is_parent: bool = False,
) -> ParsedComponent | None:
    description = clean_text(_cell_value(value_ws, row, columns.get("description")))
    if not description:
        return None
    source_qty = as_float(_cell_value(value_ws, row, columns.get("qty")), 0.0)
    qty_formula = _formula_text(_cell_value(formula_ws, row, columns.get("qty")))
    divisor = base_quantity if abs(base_quantity) > 1e-12 else 1.0
    material_factor_value = _cell_value(value_ws, row, columns.get("material_factor"))
    labor_factor_value = _cell_value(value_ws, row, columns.get("labor_factor"))
    return ParsedComponent(
        line_order=line_order,
        item_no=clean_text(_cell_value(value_ws, row, columns.get("item_no"))),
        description=description,
        quantity_per_unit=source_qty / divisor,
        source_quantity=source_qty,
        uom=clean_text(_cell_value(value_ws, row, columns.get("uom"))),
        q_m=clean_text(_cell_value(value_ws, row, columns.get("q_m"))),
        size=clean_text(_cell_value(value_ws, row, columns.get("size"))),
        unit_material=as_float(_cell_value(value_ws, row, columns.get("unit_material")), 0.0),
        material_factor=as_float(material_factor_value, 1.0) if material_factor_value not in (None, "") else 1.0,
        unit_hours=as_float(_cell_value(value_ws, row, columns.get("unit_hours")), 0.0),
        labor_factor=as_float(labor_factor_value, 1.0) if labor_factor_value not in (None, "") else 1.0,
        is_parent=1 if is_parent else 0,
        source_row=row,
        quantity_formula=qty_formula,
    )


def parse_conest_sheet(value_ws: Any, formula_ws: Any) -> ParsedSheet:
    parsed = ParsedSheet(
        sheet_name=value_ws.title,
        project_name=_project_name(value_ws),
        import_format="conest_report",
        row_count=value_ws.max_row,
        header_count=0,
    )
    columns: dict[str, int] | None = None
    phase = ""
    subphase = ""
    pending_context: list[str] = []
    active: ParsedAssembly | None = None

    def finish_active(row_no: int | None, missing_total: bool = False) -> None:
        nonlocal active
        if active is None:
            return
        active.source_material_total = sum(
            component.source_quantity * component.unit_material * component.material_factor
            for component in active.components
        )
        active.source_labor_hours = sum(
            component.source_quantity * component.unit_hours * component.labor_factor
            for component in active.components
        )
        if not active.components:
            parsed.issues.append(
                ParsedIssue("warning", "EMPTY_ASSEMBLY", "Assembly contained no usable component rows.", active.source_row)
            )
        else:
            parsed.assemblies.append(active)
        if missing_total:
            parsed.issues.append(
                ParsedIssue(
                    "warning",
                    "MISSING_TOTAL_ROW",
                    "Assembly was closed by the next assembly or end of sheet instead of a Totals row.",
                    row_no,
                    {"assembly_row": active.source_row, "assembly_name": active.name},
                )
            )
        active = None

    for row in range(1, value_ws.max_row + 1):
        possible_header = _header_map(value_ws, row)
        if possible_header:
            columns = possible_header
            parsed.header_count += 1
            continue

        texts = _row_texts(value_ws, row)
        new_phase = _extract_phase(texts, "Phase")
        new_subphase = _extract_phase(texts, "SubPhase")
        if new_phase is not None:
            phase = new_phase
        if new_subphase is not None:
            subphase = new_subphase
        if not columns:
            continue

        description = clean_text(_cell_value(value_ws, row, columns.get("description")))
        item_no = clean_text(_cell_value(value_ws, row, columns.get("item_no")))
        qty_raw = _cell_value(value_ws, row, columns.get("qty"))
        qty = as_float(qty_raw, 0.0)
        description_norm = normalize_text(description)

        if description_norm.endswith("totals") and not description_norm.startswith("phase totals"):
            finish_active(row)
            continue

        if normalize_text(item_no) == "title":
            if active is not None:
                finish_active(row, missing_total=True)
            base_qty = qty
            if base_qty <= 0:
                parsed.issues.append(
                    ParsedIssue(
                        "warning",
                        "NON_POSITIVE_BASE_QUANTITY",
                        "TITLE row had a zero or invalid quantity; a base quantity of 1 was used.",
                        row,
                        {"raw_quantity": str(qty_raw)},
                    )
                )
                base_qty = 1.0
            context = " > ".join(pending_context[-3:])
            immediate_context = pending_context[-1] if pending_context else ""
            active = ParsedAssembly(
                code=_extract_code(immediate_context or context),
                name=_clean_assembly_name(immediate_context, description),
                description=description,
                phase=phase,
                subphase=subphase,
                scope=context,
                base_quantity=base_qty,
                base_uom=clean_text(_cell_value(value_ws, row, columns.get("uom"))),
                source_row=row,
                source_type="conest",
            )
            parent = _component_from_row(value_ws, formula_ws, row, columns, base_qty, 0, True)
            if parent:
                active.components.append(parent)
            pending_context.clear()
            continue

        if active is not None:
            if description and not description_norm.startswith(("phase ", "subphase ", "takeoff report")):
                component = _component_from_row(
                    value_ws, formula_ws, row, columns, active.base_quantity, len(active.components), False
                )
                if component and (
                    component.item_no
                    or component.source_quantity != 0
                    or component.unit_material != 0
                    or component.unit_hours != 0
                ):
                    active.components.append(component)
            continue

        if description and abs(qty) < 1e-12 and not description_norm.startswith(
            ("phase totals", "takeoff report", "item", "labor result")
        ):
            if not pending_context or normalize_text(pending_context[-1]) != description_norm:
                pending_context.append(description)
                pending_context = pending_context[-3:]

    if active is not None:
        finish_active(value_ws.max_row, missing_total=True)
    if not parsed.header_count:
        parsed.issues.append(
            ParsedIssue("error", "NO_HEADER", "No ConEst Item/Qty/Description header was found.")
        )
    return parsed


def _sample_header_row(ws: Any) -> tuple[int, dict[str, int]] | None:
    expected = {
        "scopes": "scope",
        "size": "code",
        "special materials suppliers": "description",
        "take off qty": "takeoff_qty",
        "qty": "qty",
        "unit mat": "unit_material",
        "total mat total": "material_total",
        "unit hrs": "unit_hours",
        "total hrs": "hours_total",
        "total labour": "labor_total",
        "grand total": "grand_total",
        "subtotal": "subtotal",
        "comment": "comments",
    }
    for row in range(1, min(ws.max_row, 30) + 1):
        mapping: dict[str, int] = {}
        for cell in ws[row]:
            label = normalize_text(cell.value)
            if label in expected:
                # The Sample workbook contains small summary tables to the right
                # with repeated labels such as "Scopes". The first occurrence
                # belongs to the main connected estimate table.
                mapping.setdefault(expected[label], cell.column)
        if "description" in mapping and "qty" in mapping and "unit_material" in mapping:
            return row, mapping
    return None


def _sample_scope(ws: Any, row: int, scope_col: int) -> str:
    for current in range(row, 0, -1):
        value = clean_text(ws.cell(current, scope_col).value)
        if value:
            return value
    return ""


def _find_label_column(ws: Any, labels: set[str], max_row: int = 6) -> int | None:
    normalized_labels = {normalize_text(label) for label in labels}
    for row in range(1, min(ws.max_row, max_row) + 1):
        for cell in ws[row]:
            if normalize_text(cell.value) in normalized_labels:
                return int(cell.column)
    return None


def _sample_project_name(ws: Any) -> str:
    direct = _project_name(ws)
    if direct:
        return direct
    ignored = {
        "mat mark up", "material mark up", "toal labour", "total labour", "labour rate", "labor rate",
        "grand total", "unit install cost", "unit installation cost",
    }
    candidates: list[str] = []
    for row in range(1, min(ws.max_row, 5) + 1):
        for cell in ws[row]:
            text = clean_text(cell.value)
            if not text or normalize_text(text) in ignored:
                continue
            if isinstance(cell.value, str) and not cell.value.startswith("=") and len(text) >= 5:
                candidates.append(text)
    return max(candidates, key=len) if candidates else ""


def _sample_unit_install_col(ws: Any, header_row: int, columns: dict[str, int]) -> int | None:
    found = _find_label_column(
        ws,
        {"unit install cost", "unit installation cost", "unit install price", "unit installed cost"},
        max_row=max(6, header_row),
    )
    if found:
        return found
    subtotal = columns.get("subtotal")
    if subtotal and subtotal + 1 <= ws.max_column:
        return subtotal + 1
    return None


def _sample_setting_value(ws: Any, labels: set[str], default: float) -> float:
    col = _find_label_column(ws, labels, max_row=6)
    if not col:
        return default
    # Most source workbooks place the editable value directly below the label.
    for row in range(2, min(ws.max_row, 8) + 1):
        if normalize_text(ws.cell(row, col).value) in {normalize_text(label) for label in labels}:
            for candidate_row in range(row + 1, min(ws.max_row, row + 3) + 1):
                value = ws.cell(candidate_row, col).value
                if isinstance(value, (int, float)):
                    return float(value)
    return default


def _sum_range_from_formula(formula: str, total_col: int) -> tuple[int, int] | None:
    if not formula:
        return None
    letter = re.escape(get_column_letter(total_col))
    match = re.search(rf"SUM\s*\(\s*\$?{letter}\$?(\d+)\s*:\s*\$?{letter}\$?(\d+)\s*\)", formula, re.I)
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def _base_row_from_formula(formula: str, qty_col: int, start: int, end: int) -> int | None:
    if not formula:
        return None
    qty_letter = re.escape(get_column_letter(qty_col))
    # Unit install formulas commonly divide the block total by the assembly quantity cell.
    matches = re.findall(rf"/\s*\$?{qty_letter}\$?(\d+)", formula, re.I)
    for value in matches:
        row = int(value)
        if start <= row <= end:
            return row
    return None


def _formula_ratio_from_base(value_ws: Any, formula: str, qty_col: int, base_row: int, row: int) -> float | None:
    """Infer a child quantity ratio when the source estimate quantity is zero.

    Handles the common connected-estimate formulas used in the supplied workbooks, e.g.
    =$E$31, =$E$31*O35, =$C$63*M68, or =$C$40*2.
    """
    if not formula:
        return None
    text = formula.replace(" ", "").lstrip("=")
    qty_letter = get_column_letter(qty_col)
    base_pattern = re.compile(rf"\$?{re.escape(qty_letter)}\$?{base_row}", re.I)
    if not base_pattern.search(text):
        return None
    remaining = base_pattern.sub("1", text, count=1)
    # Reject formulas that contain another operator structure we cannot safely evaluate.
    if any(op in remaining for op in ("+", "-", "/", "(", ")", ",", ":")):
        if remaining in {"1", "+1"}:
            return 1.0
        return None
    factors = remaining.split("*")
    ratio = 1.0
    for factor in factors:
        factor = factor.strip()
        if not factor or factor == "1":
            continue
        try:
            ratio *= float(factor)
            continue
        except ValueError:
            pass
        ref = re.fullmatch(r"\$?([A-Z]{1,3})\$?(\d+)", factor, re.I)
        if not ref:
            return None
        col = column_index_from_string(ref.group(1))
        ref_row = int(ref.group(2))
        value = as_float(value_ws.cell(ref_row, col).value, float("nan"))
        if value != value:  # NaN
            return None
        ratio *= value
    return max(0.0, ratio)


def _sample_component(
    value_ws: Any,
    formula_ws: Any,
    row: int,
    columns: dict[str, int],
    base_row: int,
    source_base_quantity: float,
    line_order: int,
    is_parent: bool,
    code_col: int | None,
) -> ParsedComponent | None:
    description_col = columns["description"]
    description = clean_text(value_ws.cell(row, description_col).value)
    if not description:
        return None
    qty_col = columns["qty"]
    source_qty = as_float(value_ws.cell(row, qty_col).value, 0.0)
    quantity_formula = _formula_text(formula_ws.cell(row, qty_col).value)
    if source_base_quantity > 0:
        ratio = source_qty / source_base_quantity
    elif row == base_row:
        ratio = 1.0
    else:
        inferred = _formula_ratio_from_base(value_ws, quantity_formula, qty_col, base_row, row)
        ratio = inferred if inferred is not None else max(0.0, source_qty)
    return ParsedComponent(
        line_order=line_order,
        item_no=clean_text(value_ws.cell(row, code_col).value) if code_col else "",
        description=description,
        quantity_per_unit=max(0.0, ratio),
        source_quantity=max(0.0, source_qty),
        uom="EA",
        unit_material=max(0.0, as_float(value_ws.cell(row, columns["unit_material"]).value, 0.0)),
        material_factor=1.0,
        unit_hours=max(0.0, as_float(value_ws.cell(row, columns.get("unit_hours", 0)).value, 0.0)) if columns.get("unit_hours") else 0.0,
        labor_factor=1.0,
        is_parent=1 if is_parent else 0,
        source_row=row,
        quantity_formula=quantity_formula,
        notes=clean_text(value_ws.cell(row, columns.get("comments", 0)).value) if columns.get("comments") else "",
    )


def parse_sample_sheet(value_ws: Any, formula_ws: Any) -> ParsedSheet:
    header = _sample_header_row(value_ws)
    parsed = ParsedSheet(
        sheet_name=value_ws.title,
        project_name=_sample_project_name(value_ws),
        import_format="sample_cost_sheet",
        row_count=value_ws.max_row,
        header_count=1 if header else 0,
    )
    if not header:
        parsed.issues.append(ParsedIssue("error", "NO_SAMPLE_HEADER", "Sample cost-sheet header was not found."))
        return parsed

    header_row, columns = header
    markup = _sample_setting_value(value_ws, {"mat. mark up", "mat mark up", "material mark up"}, 0.15)
    labor_rate = _sample_setting_value(value_ws, {"labour rate", "labor rate"}, 120.0)
    overhead = _sample_setting_value(value_ws, {"unit install cost", "unit installation cost"}, 0.10)
    unit_install_col = _sample_unit_install_col(value_ws, header_row, columns)
    grand_total_col = columns.get("grand_total")
    qty_col = columns["qty"]
    code_col = columns.get("code")
    description_col = columns["description"]
    scope_col = columns.get("scope", 1)

    blocks: list[tuple[int, int, int, str]] = []
    if unit_install_col and grand_total_col:
        for row in range(header_row + 1, value_ws.max_row + 1):
            formula = _formula_text(formula_ws.cell(row, unit_install_col).value)
            total_range = _sum_range_from_formula(formula, grand_total_col)
            if not total_range:
                continue
            range_start, range_end = total_range
            start = min(row, range_start)
            end = max(row, range_end)
            start = max(header_row + 1, start)
            end = min(value_ws.max_row, end)
            base_row = _base_row_from_formula(formula, qty_col, start, end) or row
            if not (start <= base_row <= end):
                base_row = row
            blocks.append((start, end, base_row, formula))

    # Normalize duplicate/overlapping formulas. Keep the wider block when the same start repeats.
    by_start: dict[int, tuple[int, int, int, str]] = {}
    for block in blocks:
        start, end, base_row, formula = block
        current = by_start.get(start)
        if current is None or end > current[1]:
            by_start[start] = block
    blocks = sorted(by_start.values())

    covered: set[int] = set()
    for start, end, base_row, formula in blocks:
        description = clean_text(value_ws.cell(start, description_col).value)
        if not description:
            # A rare merged/display row can be blank. Use the quantity row's description.
            description = clean_text(value_ws.cell(base_row, description_col).value)
        if not description:
            continue
        source_base_qty = as_float(value_ws.cell(base_row, qty_col).value, 0.0)
        takeoff_col = columns.get("takeoff_qty")
        if source_base_qty <= 0 and takeoff_col:
            source_base_qty = as_float(value_ws.cell(start, takeoff_col).value, 0.0)
        effective_base_qty = source_base_qty if source_base_qty > 0 else 1.0
        raw_code = clean_text(value_ws.cell(start, code_col).value) if code_col else ""
        code = raw_code or _extract_code(description)
        scope = _sample_scope(value_ws, start, scope_col)
        assembly = ParsedAssembly(
            code=code,
            name=description,
            description=description,
            phase=scope,
            subphase="",
            scope=scope,
            base_quantity=effective_base_qty,
            base_uom="EA",
            source_row=start,
            source_type="sample",
            material_markup_hint=markup,
            labor_rate_hint=labor_rate,
            overhead_markup_hint=overhead,
        )
        for row in range(start, end + 1):
            covered.add(row)
            component = _sample_component(
                value_ws,
                formula_ws,
                row,
                columns,
                base_row,
                source_base_qty,
                len(assembly.components),
                row == base_row,
                code_col,
            )
            if component:
                assembly.components.append(component)
        if not assembly.components:
            continue
        if source_base_qty > 0:
            assembly.source_material_total = sum(c.source_quantity * c.unit_material for c in assembly.components)
            assembly.source_labor_hours = sum(c.source_quantity * c.unit_hours for c in assembly.components)
        else:
            # The source line is currently zero, but its formulas still define a reusable one-unit assembly.
            assembly.source_material_total = sum(c.quantity_per_unit * c.unit_material for c in assembly.components)
            assembly.source_labor_hours = sum(c.quantity_per_unit * c.unit_hours for c in assembly.components)
        parsed.assemblies.append(assembly)

    # Import meaningful direct cost rows that are not part of a connected assembly block.
    # This captures demolition/allowance rows and other one-line catalog items that older
    # sample parsing skipped entirely.
    for row in range(header_row + 1, value_ws.max_row + 1):
        if row in covered:
            continue
        raw_description = value_ws.cell(row, description_col).value
        description = clean_text(raw_description)
        if not description:
            continue
        normalized_description = normalize_text(description)
        if normalized_description.startswith(("total", "estimate total", "phase totals")):
            continue
        # Skip indented orphan detail rows; they should belong to a formula-connected block.
        if isinstance(raw_description, str) and raw_description[:1].isspace():
            continue
        qty = as_float(value_ws.cell(row, qty_col).value, 0.0)
        takeoff_col = columns.get("takeoff_qty")
        takeoff_qty = as_float(value_ws.cell(row, takeoff_col).value, 0.0) if takeoff_col else 0.0
        unit_material = as_float(value_ws.cell(row, columns["unit_material"]).value, 0.0)
        unit_hours = as_float(value_ws.cell(row, columns.get("unit_hours", 0)).value, 0.0) if columns.get("unit_hours") else 0.0
        material_total = as_float(value_ws.cell(row, columns.get("material_total", 0)).value, 0.0) if columns.get("material_total") else 0.0
        hours_total = as_float(value_ws.cell(row, columns.get("hours_total", 0)).value, 0.0) if columns.get("hours_total") else 0.0
        grand_total = as_float(value_ws.cell(row, columns.get("grand_total", 0)).value, 0.0) if columns.get("grand_total") else 0.0
        # Standalone rows must carry an actual material/labour/cost signal.
        # Quantity-only rows are commonly headings, legends, or auxiliary BOQ schedules
        # outside the connected estimating table and should not become zero-cost catalog items.
        cost_signal = any(
            abs(value) > 1e-12
            for value in (unit_material, unit_hours, material_total, hours_total, grand_total)
        )
        if not cost_signal:
            continue
        # A scope heading usually has a value in the scope column and no measurable cost/quantity.
        scope_here = clean_text(value_ws.cell(row, scope_col).value)
        if scope_here and not any(abs(value) > 1e-12 for value in (qty, takeoff_qty, unit_material, unit_hours, grand_total)):
            continue
        base_qty = takeoff_qty if takeoff_qty > 0 else qty if qty > 0 else 1.0
        raw_code = clean_text(value_ws.cell(row, code_col).value) if code_col else ""
        code = raw_code or _extract_code(description)
        scope = _sample_scope(value_ws, row, scope_col)
        assembly = ParsedAssembly(
            code=code,
            name=description,
            description=description,
            phase=scope,
            subphase="",
            scope=scope,
            base_quantity=base_qty,
            base_uom="EA",
            source_row=row,
            source_type="sample",
            material_markup_hint=markup,
            labor_rate_hint=labor_rate,
            overhead_markup_hint=overhead,
        )
        component = ParsedComponent(
            line_order=0,
            item_no=raw_code,
            description=description,
            quantity_per_unit=1.0,
            source_quantity=qty if qty > 0 else base_qty,
            uom="EA",
            unit_material=max(0.0, unit_material),
            material_factor=1.0,
            unit_hours=max(0.0, unit_hours),
            labor_factor=1.0,
            is_parent=1,
            source_row=row,
            quantity_formula=_formula_text(formula_ws.cell(row, qty_col).value),
            notes=clean_text(value_ws.cell(row, columns.get("comments", 0)).value) if columns.get("comments") else "",
        )
        assembly.components.append(component)
        basis = component.source_quantity if component.source_quantity > 0 else 1.0
        assembly.source_material_total = basis * component.unit_material
        assembly.source_labor_hours = basis * component.unit_hours
        parsed.assemblies.append(assembly)

    parsed.assemblies.sort(key=lambda item: item.source_row)
    return parsed

def parse_workbook(path: str | Path) -> list[ParsedSheet]:
    from .component_master import is_component_master, parse_component_master
    from .range_importer import budget_range_header, parse_budget_range
    source = Path(path)
    formula_book = openpyxl.load_workbook(source, data_only=False, read_only=False)
    value_book = openpyxl.load_workbook(source, data_only=True, read_only=False)
    results: list[ParsedSheet] = []
    try:
        for formula_ws in formula_book.worksheets:
            value_ws = value_book[formula_ws.title]
            name = formula_ws.title.upper()
            if is_component_master(value_ws):
                results.append(parse_component_master(value_ws, formula_ws))
                continue
            if any(_header_map(value_ws,row) for row in range(1,min(value_ws.max_row,80)+1)):
                results.append(parse_conest_sheet(value_ws, formula_ws))
                continue
            if budget_range_header(value_ws):
                results.append(parse_budget_range(value_ws))
                continue
            if _sample_header_row(value_ws):
                results.append(parse_sample_sheet(value_ws, formula_ws))
    finally:
        formula_book.close()
        value_book.close()
    if not results:
        results.append(
            ParsedSheet(
                sheet_name="",
                project_name="",
                import_format="unknown",
                row_count=0,
                header_count=0,
                issues=[ParsedIssue("error", "NO_SUPPORTED_SHEET", "No supported ConEst or Sample sheet was found.")],
            )
        )
    return results


def _fingerprint(assembly: ParsedAssembly) -> str:
    payload = {
        "code": normalize_text(assembly.code),
        "name": normalize_text(assembly.name),
        "components": [
            {
                "item": normalize_text(component.item_no),
                "description": normalize_text(component.description),
                "ratio": round(component.quantity_per_unit, 8),
                "uom": normalize_text(component.uom),
                "size": normalize_text(component.size),
                "material": round(component.unit_material, 8),
                "material_factor": round(component.material_factor, 8),
                "hours": round(component.unit_hours, 8),
                "labor_factor": round(component.labor_factor, 8),
            }
            for component in assembly.components
        ],
    }
    if assembly.source_type in ("component_master",'budget_range'):
        payload.update(parent_unit=normalize_text(assembly.base_uom), work_type=assembly.work_type,
                       notes=[c.notes for c in assembly.components])
    if assembly.source_type=='budget_range':
        payload['fixed_labor']=[(c.labor_pricing,c.unit_labor_price) for c in assembly.components]
    return stable_hash(payload)


def _search_text(assembly: ParsedAssembly) -> str:
    fields = [assembly.code, assembly.name, assembly.description, assembly.phase, assembly.subphase, assembly.scope]
    for component in assembly.components:
        fields.extend([component.item_no, component.description, component.size, component.uom])
    return normalize_text(" ".join(filter(None, fields)))


def _insert_assembly(conn: sqlite3.Connection, assembly: ParsedAssembly, source_id: int) -> tuple[int, bool]:
    now = utc_now()
    fingerprint = _fingerprint(assembly)
    existing = conn.execute("SELECT id FROM catalog_items WHERE fingerprint = ?", (fingerprint,)).fetchone()
    created = existing is None
    if existing:
        catalog_id = int(existing["id"])
    else:
        cursor = conn.execute(
            """INSERT INTO catalog_items
               (code, name, description, phase, subphase, scope, base_quantity, base_uom, source_type,
                fingerprint, search_text, is_manual, material_markup_hint, labor_rate_hint,
                overhead_markup_hint, created_at, updated_at, work_type, category_override)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, ?, ?, ?, ?, ?, ?, ?)""",
            (
                assembly.code,
                assembly.name,
                assembly.description,
                assembly.phase,
                assembly.subphase,
                assembly.scope,
                assembly.base_quantity,
                assembly.base_uom,
                assembly.source_type,
                fingerprint,
                _search_text(assembly),
                assembly.material_markup_hint,
                assembly.labor_rate_hint,
                assembly.overhead_markup_hint,
                now,
                now,
                assembly.work_type,
                assembly.category_override,
            ),
        )
        catalog_id = int(cursor.lastrowid)
        for component in assembly.components:
            conn.execute(
                """INSERT INTO catalog_components
                   (catalog_item_id, line_order, item_no, description, quantity_per_unit, source_quantity,
                    uom, q_m, size, unit_material, material_factor, unit_hours, labor_factor, is_parent,
                    source_row, quantity_formula, notes,labor_pricing,unit_labor_price)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    catalog_id,
                    component.line_order,
                    component.item_no,
                    component.description,
                    component.quantity_per_unit,
                    component.source_quantity,
                    component.uom,
                    component.q_m,
                    component.size,
                    component.unit_material,
                    component.material_factor,
                    component.unit_hours,
                    component.labor_factor,
                    component.is_parent,
                    component.source_row,
                    component.quantity_formula,
                    component.notes,
                    component.labor_pricing,
                    component.unit_labor_price,
                ),
            )
        aliases = {assembly.code, assembly.name, assembly.description, assembly.scope}
        for alias in filter(None, aliases):
            normalized = normalize_text(alias)
            if normalized:
                conn.execute(
                    """INSERT OR IGNORE INTO catalog_aliases
                       (catalog_item_id, alias, normalized_alias, weight, source, created_at)
                       VALUES (?, ?, ?, 1, 'import', ?)""",
                    (catalog_id, clean_text(alias), normalized, now),
                )
    conn.execute(
        """INSERT OR IGNORE INTO source_occurrences
           (catalog_item_id, source_id, source_row, base_quantity, source_scope, source_phase,
            source_subphase, source_material_total, source_labor_hours)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            catalog_id,
            source_id,
            assembly.source_row,
            assembly.base_quantity,
            assembly.scope,
            assembly.phase,
            assembly.subphase,
            assembly.source_material_total,
            assembly.source_labor_hours,
        ),
    )
    return catalog_id, created


def import_workbook(path: str | Path, db_path: str | Path | None = None, display_name: str | None = None,
                    *, initialize_database: bool = True, actor_id: int | None = None, authorize=None) -> dict[str, Any]:
    if initialize_database:initialize(db_path)
    source_path = Path(path)
    digest = sha256_file(source_path)
    parsed_sheets = parse_workbook(source_path)
    report: dict[str, Any] = {
        "file_name": display_name or source_path.name,
        "file_sha256": digest,
        "status": "imported",
        "sheets": [],
        "assemblies_seen": 0,
        "catalog_items_created": 0,
        "components_seen": 0,
        "issues": 0,
    }
    with write_lock:
        conn = connect(db_path)
        try:
            conn.execute("BEGIN IMMEDIATE")
            if authorize:authorize(conn)
            for parsed in parsed_sheets:
                existing = conn.execute(
                    "SELECT id,assembly_count FROM sources WHERE file_sha256 = ? AND sheet_name = ?",
                    (digest, parsed.sheet_name),
                ).fetchone()
                if existing and (existing['assembly_count'] or not parsed.assemblies):
                    report["sheets"].append(
                        {"sheet_name": parsed.sheet_name, "status": "already_imported", "source_id": existing["id"]}
                    )
                    continue
                component_count = sum(len(item.components) for item in parsed.assemblies)
                if existing:
                    # An earlier unsupported-sheet record may now be recoverable.
                    # Keep its identity and mark the historical parsing error resolved.
                    source_id=existing['id']
                    conn.execute('''UPDATE sources SET project_name=?,import_format=?,row_count=?,header_count=?,
                        assembly_count=?,component_count=?,warning_count=?,imported_by=? WHERE id=?''',
                        (parsed.project_name,parsed.import_format,parsed.row_count,parsed.header_count,
                         len(parsed.assemblies),component_count,len(parsed.issues),actor_id,source_id))
                    conn.execute('UPDATE import_issues SET resolved=1 WHERE source_id=?',(source_id,))
                else:
                    cursor = conn.execute(
                        """INSERT INTO sources
                           (file_name, file_sha256, sheet_name, project_name, import_format, row_count,
                            header_count, assembly_count, component_count, warning_count, imported_at)
                           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        (display_name or source_path.name, digest, parsed.sheet_name, parsed.project_name,
                         parsed.import_format, parsed.row_count, parsed.header_count, len(parsed.assemblies),
                         component_count, len(parsed.issues), utc_now()),
                    )
                    source_id = int(cursor.lastrowid)
                    conn.execute('UPDATE sources SET imported_by=? WHERE id=?',(actor_id,source_id))
                created_count = 0
                for assembly in parsed.assemblies:
                    _, created = _insert_assembly(conn, assembly, source_id)
                    created_count += int(created)
                for issue in parsed.issues:
                    conn.execute(
                        """INSERT INTO import_issues
                           (source_id, severity, issue_code, message, row_no, payload_json, created_at)
                           VALUES (?, ?, ?, ?, ?, ?, ?)""",
                        (
                            source_id,
                            issue.severity,
                            issue.issue_code,
                            issue.message,
                            issue.row_no,
                            json_dumps(issue.payload),
                            utc_now(),
                        ),
                    )
                report["assemblies_seen"] += len(parsed.assemblies)
                report["catalog_items_created"] += created_count
                report["components_seen"] += component_count
                report["issues"] += len(parsed.issues)
                report["sheets"].append(
                    {
                        "sheet_name": parsed.sheet_name,
                        "format": parsed.import_format,
                        "status": "imported",
                        "source_id": source_id,
                        "rows_read": parsed.row_count,
                        "headers_found": parsed.header_count,
                        "assemblies_seen": len(parsed.assemblies),
                        "catalog_items_created": created_count,
                        "components_seen": component_count,
                        "issues": len(parsed.issues),
                    }
                )
            if report["catalog_items_created"]:
                refresh_catalog_index(conn)
            if report['assemblies_seen']:
                conn.execute('''UPDATE import_issues SET resolved=1 WHERE severity='error' AND
                    source_id IN (SELECT id FROM sources WHERE file_sha256=? AND assembly_count=0)''',(digest,))
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()
    if report["sheets"] and all(sheet.get("status") == "already_imported" for sheet in report["sheets"]):
        report["status"] = "already_imported"
    return report


def import_many(paths: Iterable[str | Path], db_path: str | Path | None = None) -> dict[str, Any]:
    reports = []
    for path in paths:
        reports.append(import_workbook(path, db_path))
    return {
        "files": reports,
        "totals": {
            "files": len(reports),
            "sheets": sum(len(report["sheets"]) for report in reports),
            "assemblies_seen": sum(report["assemblies_seen"] for report in reports),
            "catalog_items_created": sum(report["catalog_items_created"] for report in reports),
            "components_seen": sum(report["components_seen"] for report in reports),
            "issues": sum(report["issues"] for report in reports),
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Import ConEst Excel workbooks into the catalog database.")
    parser.add_argument("paths", nargs="+", help="Excel workbook paths")
    parser.add_argument("--database", help="SQLite database path")
    parser.add_argument("--report", help="Optional JSON report path")
    args = parser.parse_args()
    report = import_many(args.paths, args.database)
    text = json.dumps(report, indent=2, ensure_ascii=False)
    if args.report:
        Path(args.report).write_text(text + "\n", encoding="utf-8")
    print(text)


if __name__ == "__main__":
    main()
