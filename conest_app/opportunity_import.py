"""Bounded, previewed imports for upcoming work and the packaged v3.5 seed.

Imports add missing opportunities only. Re-importing a workbook never overwrites
an approval, assignment, deletion, or a change made in the application. Original
spreadsheet values and source decisions remain attached to every imported record.
"""
from __future__ import annotations

import hashlib
import io
import json
import math
import posixpath
import re
import zlib
from datetime import date, datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import quoteattr
from zipfile import BadZipFile, ZipFile, ZIP_DEFLATED

from .tenders import ValidationError, project_key
from .opportunity_retention import active_sql

ROOT = Path(__file__).resolve().parents[1]
MAX_BYTES = 15 * 1024 * 1024
MAX_ROWS = 10000
MAX_PREVIEW = 200
ORIGINS = {
    'NSW': '63-79 Parramatta Rd Silverwater, NSW 2128',
    'VIC': 'CBD, Melbourne',
}
ALIASES = {
    'opportunity id': 'sourceRecordId', 'project': 'project',
    'project name': 'project', 'opportunity': 'project', 'opportunity name': 'project',
    'region': 'region', 'state': 'region', 'location': 'location',
    'project location': 'location', 'address': 'location',
    'contact': 'contact', 'contact person': 'contact', 'contact name': 'contact',
    'email': 'email', 'e-mail': 'email', 'email address': 'email',
    'link': 'link', 'url': 'link', 'project link': 'link',
    'opportunity link': 'link', 'website': 'link',
    'working scope': 'workingScope', 'work scope': 'workingScope', 'scope': 'workingScope',
    'scope of works': 'workingScope', 'category': 'category',
    'budget': 'budget', 'project budget': 'budget', 'budget range': 'budget',
    'value': 'value', 'estimated value': 'value', 'estimated value (aud)': 'value',
    'comments': 'remarks', 'remarks': 'remarks', 'notes': 'remarks',
    'added date': 'addedDate', 'date added': 'addedDate',
    'closing date': 'due', 'close date': 'due', 'due': 'due', 'due date': 'due',
    'submission deadline': 'due', 'deadline': 'due',
    'distance': 'distanceKm', 'distance km': 'distanceKm', 'distance (km)': 'distanceKm',
    'distance origin': 'distanceOrigin', 'distance from': 'distanceOrigin',
    'origin': 'distanceOrigin', 'approved': 'approved',
    'not interested': 'rejected', 'rejected': 'rejected',
    'status': 'exportStatus', 'source status': 'sourceStatus',
    'source decision': 'sourceStatus',
}
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
REL_NS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
INACTIVE_SOURCE_FIELDS = {'approved', 'rejected', 'sourceStatus', 'exportStatus', 'contact', 'email'}


def norm(value):
    return ' '.join(str('' if value is None else value).split()).casefold()


def _text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _json_cell(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def _header(value):
    label = norm(value)
    if re.fullmatch(r'builders?(?:\s*\d+)?', label):
        return 'builders'
    if label.startswith('distance from ') and label not in ALIASES:
        return 'distanceKm'
    if label in ALIASES:
        return ALIASES[label]
    # The supplied workbooks include their date format on the header's next line.
    label = re.sub(r'\s*\(?[dmy]{1,4}[/.-][dmy]{1,4}[/.-][dmy]{1,4}\)?$', '', label)
    return ALIASES.get(label)


def _region(value):
    value = _text(value).upper()
    value = {'NEW SOUTH WALES': 'NSW', 'VICTORIA': 'VIC'}.get(value, value)
    if value and not re.fullmatch(r'[A-Z][A-Z0-9 _-]{1,19}', value):
        raise ValidationError('Region must be a short name or code, for example NSW or VIC.')
    return value


def _inferred_region(filename, sheet, headers, hint):
    if hint:
        return _region(hint)
    context = ' '.join([filename, sheet] + [_text(h) for h in headers])
    matches = set(re.findall(r'\b(NSW|VIC)\b', context.upper()))
    if len(matches) == 1:
        return matches.pop()
    # Both supplied sheets are called E1; their distance heading identifies VIC.
    origins = [norm(h) for h in headers if norm(h).startswith('distance from')]
    if any('melbourne' in h for h in origins):
        return 'VIC'
    if any('silverwater' in h or 'sydney' in h for h in origins):
        return 'NSW'
    return ''


def _iso(value, epoch, header=''):
    if value in (None, '') or norm(value) in ('n/a', 'na', 'tbc', 'tba', '-', 'not specified'):
        return ''
    if isinstance(value, bool):
        raise ValidationError('Choose a valid date; a checkbox is not a date.')
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value < 1 or value > 2958465:
            raise ValidationError('Date serial number is outside the supported Excel range.')
        from openpyxl.utils.datetime import from_excel
        value = from_excel(value, epoch=epoch)
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    value = _text(value)
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError:
        pass
    numeric = ('%m/%d/%Y', '%m/%d/%y') if 'mm/dd' in norm(header) else ('%d/%m/%Y', '%d/%m/%y')
    for fmt in (*numeric, '%d-%m-%Y', '%d %b %Y', '%d %B %Y', '%Y/%m/%d'):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValidationError(f'Invalid date: {value}. Use a real Excel date or YYYY-MM-DD.')


def _number(value, label, distance=False):
    if value in (None, '') or norm(value) in ('n/a', 'na', '-', 'tbc', 'tba'):
        return None
    if isinstance(value, bool):
        raise ValidationError(label + ' must be a number, not a checkbox.')
    raw = _text(value).replace(',', '')
    multiplier = 1
    if distance:
        match = re.fullmatch(r'\s*([+]?(?:\d+(?:\.\d*)?|\.\d+))\s*(km|kilometres?|kilometers?|m|metres?|meters?)?\s*', raw, re.I)
        if not match:
            raise ValidationError('Distance must be a non-negative number, optionally followed by km or m.')
        raw = match.group(1)
        if (match.group(2) or '').lower() in ('m', 'metre', 'metres', 'meter', 'meters'):
            multiplier = 0.001
    else:
        raw = re.sub(r'^(?:AUD\s*)?\$?\s*', '', raw, flags=re.I)
    try:
        result = float(raw) * multiplier
    except (ValueError, TypeError):
        raise ValidationError(label + ' must be numeric. Keep budget ranges in the Budget column.')
    if not math.isfinite(result) or result < 0:
        raise ValidationError(label + ' must be zero or greater.')
    return result


def _bool(value, label):
    if value in (None, '', False, 0):
        return False
    if value in (True, 1):
        return True
    word = norm(value)
    if word in ('false', 'no', 'n', '0', 'unchecked'):
        return False
    if word in ('true', 'yes', 'y', '1', 'checked', 'x'):
        return True
    raise ValidationError(f'{label} must be Yes/No or a checked/unchecked checkbox.')


def _status(value):
    word = norm(value)
    if not word:
        return 'pending'
    choices = {
        'awaiting approval': 'pending', 'awaiting for approval': 'pending',
        'pending': 'pending', 'new': 'pending', 'approved': 'approved',
        'rejected': 'rejected', 'not interested': 'rejected', 'expired': 'expired',
    }
    if word not in choices:
        raise ValidationError('Status must be Awaiting approval, Approved, Rejected, or Expired.')
    return choices[word]


def _builders(values):
    result = []
    seen = set()
    for value in values:
        # Punctuation can be part of a company's name. Match the application's
        # builder contract: separate builders by rows, numbered columns, or lines.
        for builder in _text(value).splitlines():
            builder = builder.strip()
            key = project_key(builder)
            if not key or key in seen:
                continue
            if len(builder) > 500:
                raise ValidationError('Each builder name must be 500 characters or shorter.')
            seen.add(key)
            result.append(builder)
    if len(result) > 100:
        raise ValidationError('Use at most 100 builders per opportunity.')
    if len('\n'.join(result)) > 20000:
        raise ValidationError('Builder names must total 20,000 characters or fewer.')
    return result


def _archive_metadata(content):
    """Read only formula/link metadata; never follow a workbook's external links."""
    with ZipFile(io.BytesIO(content)) as archive:
        infos = archive.infolist()
        if len(infos) > 5000 or sum(i.file_size for i in infos) > 100 * 1024 * 1024:
            raise ValidationError('Workbook is too large after decompression.')
        if any(i.flag_bits & 0x1 for i in infos):
            raise ValidationError('Remove workbook encryption before importing.')
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        relationships = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {r.attrib['Id']: r.attrib['Target'] for r in relationships}
        metadata = {}
        for sheet in workbook.findall('s:sheets/s:sheet', NS):
            target = targets[sheet.attrib[REL_NS + 'id']]
            path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
            if not path.startswith('xl/') or '..' in path.split('/'):
                raise ValidationError('Workbook contains an invalid worksheet path.')
            xml = ET.fromstring(archive.read(path))
            formulae = {}
            occupied_rows = {}
            for cell in xml.findall('.//s:sheetData/s:row/s:c', NS):
                address = cell.attrib.get('r', '')
                coordinate = re.fullmatch(r'([A-Z]+)([1-9]\d*)', address)
                value = cell.find('s:v', NS)
                inline = ''.join(node.text or '' for node in cell.findall('s:is//s:t', NS))
                if coordinate and ((value is not None and value.text is not None) or inline):
                    column = 0
                    for letter in coordinate.group(1):
                        column = column * 26 + ord(letter) - 64
                    occupied_rows[column] = max(occupied_rows.get(column, 0), int(coordinate.group(2)))
                formula = cell.find('s:f', NS)
                if formula is not None:
                    cached = cell.find('s:v', NS)
                    empty_text = cached is not None and cached.text is None and cell.attrib.get('t') == 'str'
                    valid = cached is not None and (cached.text is not None or empty_text) and cell.attrib.get('t') != 'e'
                    formulae[address] = {'formula': '=' + (formula.text or ''), 'valid': valid}
            link_targets = {}
            relpath = posixpath.join(posixpath.dirname(path), '_rels', posixpath.basename(path) + '.rels')
            if relpath in archive.namelist():
                rels = ET.fromstring(archive.read(relpath))
                link_targets = {r.attrib['Id']: r.attrib.get('Target', '') for r in rels}
            links = {}
            for link in xml.findall('s:hyperlinks/s:hyperlink', NS):
                destination = link_targets.get(link.attrib.get(REL_NS + 'id'), link.attrib.get('location', ''))
                links[link.attrib.get('ref', '')] = destination
            metadata[sheet.attrib['name']] = {'formulae': formulae, 'links': links, 'occupiedRows': occupied_rows}
        return metadata


def parse_workbook(filename, content, region_hint=''):
    """Parse both the supplied E1 worksheets and an Opportunities export.

    Preview collects row issues. Apply refuses an import containing invalid rows,
    so malformed records are never silently dropped.
    """
    filename = Path(str(filename)).name
    if Path(filename).suffix.lower() not in ('.xlsx', '.xlsm'):
        raise ValidationError('Choose an Excel .xlsx or .xlsm workbook.')
    if not isinstance(content, bytes) or not content or len(content) > MAX_BYTES:
        raise ValidationError('Choose a workbook smaller than 15 MB.')
    region_hint = _region(region_hint)
    try:
        metadata = _archive_metadata(content)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except ValidationError:
        raise
    except (BadZipFile, ValueError, TypeError, KeyError, IndexError, OSError, EOFError,
            ET.ParseError, RuntimeError, NotImplementedError, zlib.error) as exc:
        raise ValidationError('Cannot read this workbook. Save it as .xlsx and try again.') from exc
    records, issues, warnings, sheets = [], [], [], []
    scanned = 0
    try:
        for ws in wb:
            ws.reset_dimensions()
            mapping = None
            headers = []
            inferred = ''
            info = metadata.get(ws.title, {'formulae': {}, 'links': {}})
            for rn, cells in enumerate(ws.iter_rows(max_row=MAX_ROWS + 51, max_col=100), 1):
                values = [c.value for c in cells]
                if mapping is None:
                    candidate = {i: _header(v) for i, v in enumerate(values) if _header(v)}
                    fields = set(candidate.values())
                    if 'project' in fields and fields.intersection({'due', 'builders', 'region', 'category', 'budget', 'workingScope'}):
                        mapping = candidate
                        headers = values
                        inferred = _inferred_region(filename, ws.title, headers, region_hint)
                        occupied = info.get('occupiedRows', {})
                        if any(column > 100 for column in occupied):
                            raise ValidationError(f'{ws.title}: use at most 100 populated columns per opportunity worksheet.')
                        active_columns = [i + 1 for i, key in mapping.items() if key not in INACTIVE_SOURCE_FIELDS]
                        if any(occupied.get(column, 0) > MAX_ROWS + 50 for column in active_columns):
                            raise ValidationError(f'{ws.title}: opportunity values occur below row {MAX_ROWS + 50:,}. Remove unused rows or split the workbook before importing; no rows were added.')
                        sheets.append(ws.title)
                    elif rn >= 50:
                        break
                    continue
                # Ignore source template rows containing only checkboxes and a
                # Remaining Time formula, even if its cached text says expired.
                active = [i for i, key in mapping.items() if key not in INACTIVE_SOURCE_FIELDS]
                if not any(values[i] not in (None, '') for i in active):
                    continue
                if _header(values[next(i for i, key in mapping.items() if key == 'project')]) == 'project':
                    continue
                scanned += 1
                if scanned > MAX_ROWS:
                    raise ValidationError(f'Import at most {MAX_ROWS:,} opportunities per workbook.')
                def val(field, default=None):
                    return next((values[i] for i, name in mapping.items() if name == field), default)
                def heading(field):
                    return next((_text(headers[i]) for i, name in mapping.items() if name == field), '')
                try:
                    from openpyxl.utils import get_column_letter
                    for i, key in mapping.items():
                        if key in ('contact', 'email'):
                            # Preserve old shared contacts exactly as source
                            # cells, without treating them as current fields.
                            continue
                        address = f'{get_column_letter(i + 1)}{rn}'
                        formula = info['formulae'].get(address)
                        if key == 'link' and formula:
                            raise ValidationError(f'{address}: paste the web address in Link or use Excel Insert > Link; Link formulas are not imported.')
                        if formula and not formula['valid']:
                            raise ValidationError(f'{address}: formula has no valid saved result. Recalculate and save the workbook in Excel first.')
                        if cells[i].data_type == 'e':
                            raise ValidationError(f'{address}: Excel error {values[i]}. Correct the cell before importing.')
                    project = _text(val('project'))
                    if not project:
                        raise ValidationError('Project Name is required.')
                    if len(project) > 500:
                        raise ValidationError('Project Name must be 500 characters or shorter.')
                    region = _region(val('region')) or inferred
                    if not region:
                        raise ValidationError('Region is missing. Add a Region column or choose the workbook region before previewing.')
                    approved = _bool(val('approved'), 'Approved')
                    rejected = _bool(val('rejected'), 'Not interested')
                    if approved and rejected:
                        raise ValidationError('Approved and Not interested are both checked. Choose one source decision.')
                    source_status = _status(val('sourceStatus') or val('exportStatus'))
                    if approved:
                        source_status = 'approved'
                    elif rejected:
                        source_status = 'rejected'
                    builders = _builders([values[i] for i, key in mapping.items() if key == 'builders'])
                    origin = _text(val('distanceOrigin')) or ORIGINS.get(region, '')
                    if not origin:
                        distance_header = heading('distanceKm')
                        if norm(distance_header).startswith('distance from '):
                            origin = re.sub(r'^distance\s+from\s+', '', distance_header, flags=re.I).strip()
                    source_cells = []
                    for i, (header, cell) in enumerate(zip(headers, cells)):
                        column = get_column_letter(i + 1)
                        address = f'{column}{rn}'
                        formula = info['formulae'].get(address)
                        link = info['links'].get(address)
                        if header is None or (cell.value is None and not formula and not link):
                            continue
                        saved = {'column': column, 'header': _text(header), 'value': _json_cell(cell.value)}
                        if formula:
                            saved['formula'] = formula['formula']
                        if link:
                            saved['hyperlink'] = link
                        source_cells.append(saved)
                    from .opportunities import normalise_link
                    link_column = next((i for i, key in mapping.items() if key == 'link'), None)
                    link_value = val('link')
                    if link_column is not None:
                        address = f'{get_column_letter(link_column + 1)}{rn}'
                        if address in info['links']:
                            # Excel may display "Open documents" while storing
                            # the actual URL in the worksheet relationships.
                            link_value = info['links'][address]
                            if not link_value:
                                raise ValidationError(f'{address}: Link has no destination. Paste the web address or correct its Excel hyperlink.')
                    item = {
                        'project': project, 'region': region, 'builders': builders,
                        'due': _iso(val('due'), wb.epoch, heading('due')),
                        'addedDate': _iso(val('addedDate'), wb.epoch, heading('addedDate')),
                        'location': _text(val('location')), 'workingScope': _text(val('workingScope')),
                        'category': _text(val('category')), 'budget': _text(val('budget')),
                        'value': _number(val('value'), 'Estimated value'),
                        'distanceKm': _number(val('distanceKm'), 'Distance', True),
                        'distanceOrigin': origin, 'remarks': _text(val('remarks')),
                        'link': normalise_link(link_value),
                        'sourceRecordId': _text(val('sourceRecordId')),
                        'sourceStatus': source_status,
                        '_providedFields': sorted(set(mapping.values()).intersection({
                            'project', 'builders', 'location', 'workingScope', 'due', 'addedDate',
                            'category', 'budget', 'value', 'distanceKm', 'distanceOrigin', 'remarks',
                            'link',
                        }) | {'region'} | ({'distanceOrigin'} if 'distanceKm' in mapping.values() else set())),
                        'sourceData': {'filename': filename, 'sheet': ws.title, 'row': rn, 'cells': source_cells},
                    }
                    for field, label, limit in (
                        ('location', 'Location', 1000), ('workingScope', 'Working Scope', 20000),
                        ('remarks', 'Comments', 20000), ('category', 'Category', 500),
                        ('budget', 'Budget', 500), ('distanceOrigin', 'Distance Origin', 1000),
                        ('sourceRecordId', 'Opportunity ID', 100),
                    ):
                        if len(item[field]) > limit:
                            raise ValidationError(f'{label} must be {limit:,} characters or shorter.')
                    row_warnings = []
                    if not builders:
                        row_warnings.append('No builder is named in the source; left unassigned.')
                    declared = re.search(r'\b(?:total|all)\s+(\d+)\s+builders?\b', item['remarks'], re.I)
                    if declared and int(declared.group(1)) > len(builders):
                        row_warnings.append(f'The source mentions {declared.group(1)} builders but names only {len(builders)}; only named builders were imported.')
                    if row_warnings:
                        item['sourceData']['warnings'] = row_warnings
                        warnings.extend({'sheet': ws.title, 'row': rn, 'message': message} for message in row_warnings)
                    records.append(item)
                except (ValidationError, ValueError, TypeError, OverflowError) as exc:
                    issues.append({'sheet': ws.title, 'row': rn, 'project': _text(val('project')), 'message': str(exc)})
        if not sheets:
            raise ValidationError('No opportunity table found. Include Project Name and a Region, Closing Date, Builder, or Working Scope column.')
        if not records and not issues:
            raise ValidationError('No opportunity rows found. Add at least one project below the column headings.')
        return {'filename': filename, 'records': records, 'issues': issues, 'warnings': warnings, 'sheets': sheets, 'totalRows': scanned}
    except ValidationError:
        raise
    except (ValueError, TypeError, KeyError, IndexError, OSError, EOFError,
            ET.ParseError, RuntimeError, NotImplementedError, zlib.error) as exc:
        raise ValidationError('Workbook contains invalid cells. Check dates and numbers, then save it again in Excel.') from exc
    finally:
        wb.close()


def _key(record):
    return record['region'] + '|' + project_key(record['project'])


def _comparable(record):
    return {key: value for key, value in record.items() if key not in ('sourceData', 'sourceRecordId', '_providedFields')}


def _different_fields(incoming, current):
    differences = []
    for key in incoming.get('_providedFields', incoming.keys()):
        if key not in ('project', 'builders', 'region', 'location', 'workingScope', 'due', 'addedDate', 'category', 'budget', 'value', 'distanceKm', 'distanceOrigin', 'remarks', 'link'):
            continue
        before, after = current.get(key), incoming.get(key)
        if key == 'builders':
            before = sorted(project_key(v) for v in before or [])
            after = sorted(project_key(v) for v in after or [])
        elif key not in ('value', 'distanceKm'):
            before, after = before or '', after or ''
        if before != after:
            differences.append(key)
    return differences


def import_workbook(store, filename, content, actor, expected_token=None, region_hint=''):
    from . import opportunities
    opportunities.initialize(store)
    parsed = parse_workbook(filename, content, region_hint)
    with store.connect() as db:
        db.execute('BEGIN')
        generation = db.execute('SELECT value FROM opportunity_generation').fetchone()[0]
        current = list(db.execute('SELECT id,region,project_key,body,deleted_at FROM opportunities WHERE ' +
                                  active_sql('opportunity', 'opportunities.id')))
        known_regions = {row['code'] for row in db.execute('SELECT code FROM opportunity_regions')}
    existing = {row['region'] + '|' + row['project_key']: row['id'] for row in current}
    existing_ids = {row['id'] for row in current}
    existing_bodies = {row['id']: json.loads(row['body']) for row in current}
    deleted_ids = {row['id'] for row in current if row['deleted_at']}
    digest = hashlib.sha256(content).hexdigest()
    token = hashlib.sha256(f'{digest}:{generation}:{actor["id"]}:{_region(region_hint)}'.encode()).hexdigest()
    if expected_token is not None and expected_token != token:
        raise ValidationError('Opportunity data changed after preview. Preview the workbook again before importing.', 409)
    rows, new_records, seen = [], [], {}
    issues = list(parsed['issues'])
    counts = {'new': 0, 'added': 0, 'updated': 0, 'unchanged': 0, 'duplicates': 0, 'skippedChanged': 0, 'errors': len(issues)}
    for record in parsed['records']:
        key = _key(record)
        action = 'new'
        message = ''
        differences = []
        if record.get('sourceRecordId') in existing_ids or key in existing:
            action = 'duplicate'
            identifier = record.get('sourceRecordId') if record.get('sourceRecordId') in existing_ids else existing[key]
            if identifier in deleted_ids:
                message = 'Skipped: this opportunity was deleted in the app and will not be restored.'
            else:
                differences = _different_fields(record, existing_bodies[identifier])
                if differences:
                    action = 'skipped'
                    message = 'Skipped: the existing project differs. Edit the opportunity in the app; this import preserves its current data.'
        elif key in seen:
            if _comparable(record) != _comparable(seen[key]):
                action = 'invalid'
                issues.append({'sheet': record['sourceData']['sheet'], 'row': record['sourceData']['row'], 'project': record['project'], 'message': 'Conflicting rows have the same project and region. Combine the builders and keep one row per opportunity.'})
                counts['errors'] += 1
            else:
                action = 'duplicate'
        if action == 'new' and record['region'] not in known_regions and not record['distanceOrigin']:
            action = 'invalid'
            issues.append({'sheet': record['sourceData']['sheet'], 'row': record['sourceData']['row'], 'project': record['project'], 'message': 'Add Distance Origin for this new region, or configure the region origin in Opportunities first.'})
            counts['errors'] += 1
        if action == 'new':
            seen[key] = record
            new_records.append(record)
            if record['distanceOrigin']:
                known_regions.add(record['region'])
            counts['new'] += 1
            counts['added'] += 1
        elif action in ('duplicate', 'skipped'):
            if action == 'duplicate':
                counts['unchanged'] += 1
            else:
                counts['skippedChanged'] += 1
            counts['duplicates'] += 1
        rows.append({'project': record['project'], 'region': record['region'], 'builders': record['builders'], 'due': record['due'], 'action': action, 'message': message, 'changedFields': differences, 'sourceRow': record['sourceData']['row'], 'sheet': record['sourceData']['sheet']})
    for issue in parsed['issues']:
        rows.append({'project': issue.get('project', ''), 'region': '', 'builders': [], 'due': '', 'action': 'invalid', 'sourceRow': issue['row'], 'sheet': issue['sheet']})
    result = {
        'filename': parsed['filename'], 'token': token, 'applied': False,
        'counts': counts, 'totalRows': parsed['totalRows'], 'rows': rows[:MAX_PREVIEW],
        'issues': issues[:MAX_PREVIEW], 'warnings': parsed['warnings'][:MAX_PREVIEW],
        'previewTruncated': len(rows) > MAX_PREVIEW,
        'issuesTruncated': len(issues) > MAX_PREVIEW,
        'notes': ['Imports add missing opportunities only. Existing projects, assignments, decisions, and deleted records are preserved.', 'New imported projects start Awaiting approval; imported historical decisions are kept as source information. Closing dates that have passed are marked Expired.', 'Distances and budgets are copied from the workbook; route distances are not recalculated.'],
    }
    if expected_token is not None:
        if issues:
            raise ValidationError('Correct the highlighted workbook rows and preview again before importing.')
        committed = opportunities.import_records(store, parsed['records'], actor, parsed['filename'], expected_generation=generation)
        result['applied'] = True
        result['counts']['added'] = committed.get('added', len(new_records))
        result['counts']['new'] = result['counts']['added']
        result['importedIds'] = [item['id'] if isinstance(item, dict) else item for item in committed.get('items', [])]
    return result


def initialize_seed(store):
    """Legacy explicit seed import; never called automatically at startup.

    Release 3.5.1 ships an empty sources list. New records are added through the
    reviewed Excel import or the New opportunity form instead.
    """
    from . import opportunities
    path = ROOT / 'opportunities_seed' / 'seed.json'
    if not path.is_file():
        return {'added': 0, 'seeded': False}
    data = json.loads(path.read_text(encoding='utf-8'))
    actor = {'id': 0, 'name': 'Workbook import', 'role': 'system'}
    added = 0
    for source in data.get('sources', []):
        records = source['records']
        result = opportunities.import_records(store, records, actor, source['filename'], seed_key=source['seedKey'])
        added += result.get('added', 0)
    return {'added': added, 'seeded': bool(data.get('sources'))}


EXPORT_HEADERS = [
    'Opportunity ID', 'Project Name', 'Region', 'Location', 'Working Scope',
    'Added Date', 'Closing Date', 'Category', 'Budget', 'Estimated value (AUD)',
    'Distance (km)', 'Distance Origin', 'Builders', 'Comments', 'Status',
    'Source decision', 'Assigned teams', 'Assigned people',
    'Link',
]


def export_workbook(store, ids=None, template=False):
    """Generate the application's Excel export using its existing XML helpers."""
    if template:
        path = ROOT / 'templates' / 'opportunities_import.xlsx'
        if not path.is_file():
            raise ValidationError('The import template is unavailable. Ask an administrator to restore the template from the update package.', 503)
        return path.read_bytes()
    from .tender_export import cell, col, styles, serial
    from .opportunities import directory, normalise_link
    teams, users = directory() if not template else ([], [])
    team_names = {item['id']: item['name'] for item in teams}
    user_names = {item['id']: item['name'] for item in users}
    wanted = set(ids) if ids is not None else None
    records = []
    if not template:
        with store.connect() as db:
            for row in db.execute("SELECT body,revision FROM opportunities WHERE (deleted_at IS NULL OR deleted_at='') AND " +
                                  active_sql('opportunity', 'opportunities.id') + ' ORDER BY project_key,region,id'):
                item = json.loads(row['body'])
                if wanted is None or item['id'] in wanted:
                    records.append(item)
    labels = {'pending': 'Awaiting approval', 'approved': 'Approved', 'rejected': 'Rejected', 'expired': 'Expired'}
    rows = []
    for item in records:
        rows.append([
            item['id'], item.get('project', ''), item.get('region', ''), item.get('location', ''),
            item.get('workingScope', ''), serial(item.get('addedDate')), serial(item.get('due')),
            item.get('category', ''), item.get('budget', ''), item.get('value'), item.get('distanceKm'),
            item.get('distanceOrigin', ''), '\n'.join(item.get('builders', [])), item.get('remarks', ''),
            labels.get(item.get('status'), item.get('status', '')),
            labels.get(item.get('sourceStatus'), item.get('sourceStatus', '')),
            '\n'.join(team_names.get(tid, 'Removed team') for tid in item.get('assignedTeamIds', [])),
            '\n'.join(user_names.get(uid, 'Inactive or removed user') for uid in item.get('assignedUserIds', [])),
            item.get('link', ''),
        ])
    widths = [38, 54, 13, 38, 58, 19, 19, 23, 20, 24, 17, 48, 42, 58, 23, 24, 28, 28, 54]
    last = max(6, 5 + len(rows))
    content = []
    notes = [
        'TEC Tenders — Opportunities' + (' import template' if template else ''),
        'Exported ' + datetime.now(timezone.utc).strftime('%d %b %Y %H:%M UTC') + ' · Values in AUD; distances in km',
        'Add missing projects in this table, save as .xlsx, then import under Opportunities. Existing IDs/projects are skipped.',
        'Use Excel dates or YYYY-MM-DD; builders on separate lines. Link: paste a web address or use Insert > Link. Manage approval/assignment in the app.',
    ]
    for rn, note in enumerate(notes, 1):
        content.append(f'<row r="{rn}" ht="{36 if rn == 1 else 28}" customHeight="1">{cell("A"+str(rn), note, 1 if rn == 1 else 2)}</row>')
    content.append('<row r="5" ht="34" customHeight="1">' + ''.join(cell(f'{col(i)}5', label, 3) for i, label in enumerate(EXPORT_HEADERS, 1)) + '</row>')
    hyperlink_cells, hyperlink_relationships = [], []
    link_column = EXPORT_HEADERS.index('Link') + 1
    export_styles = ET.fromstring(styles())
    fonts, formats = export_styles.find('s:fonts', NS), export_styles.find('s:cellXfs', NS)
    font_id = len(fonts)
    fonts.append(ET.fromstring(f'<font xmlns="{NS["s"]}"><sz val="10"/><color rgb="FF195D86"/><name val="Aptos"/><u/></font>'))
    fonts.set('count', str(len(fonts)))
    link_styles = []
    for original in (5, 6):
        style = ET.fromstring(ET.tostring(formats[original]))
        style.set('fontId', str(font_id))
        link_styles.append(len(formats))
        formats.append(style)
    formats.set('count', str(len(formats)))
    for rn, row in enumerate(rows or [[None] * len(EXPORT_HEADERS)], 6):
        out = []
        for cn, value in enumerate(row, 1):
            style = (9 if rn % 2 == 0 else 10) if cn in (6, 7) else (7 if rn % 2 == 0 else 8) if cn == 10 else (5 if rn % 2 == 0 else 6)
            if cn == link_column and value:
                try:
                    target = normalise_link(value)
                except ValidationError:
                    # Unknown historical values remain visible as literal text,
                    # but must never become unsafe executable workbook links.
                    target = ''
                if target:
                    relation = 'rId' + str(len(hyperlink_relationships) + 1)
                    hyperlink_cells.append(f'<hyperlink ref="{col(cn)}{rn}" r:id="{relation}"/>')
                    hyperlink_relationships.append(f'<Relationship Id="{relation}" Type="{REL_NS[1:-1]}/hyperlink" Target={quoteattr(target)} TargetMode="External"/>')
                    style = link_styles[rn % 2]
            out.append(cell(f'{col(cn)}{rn}', value, style))
        line_count = max((len(str(value or '').split('\n')) for value in row), default=1)
        content.append(f'<row r="{rn}" ht="{min(180, max(46, 17 * line_count + 14))}" customHeight="1">' + ''.join(out) + '</row>')
    columns = ''.join(f'<col min="{i}" max="{i}" width="{width}" customWidth="1"/>' for i, width in enumerate(widths, 1))
    merges = ''.join(f'<mergeCell ref="A{i}:H{i}"/>' for i in range(1, 5))
    hyperlinks = '<hyperlinks>' + ''.join(hyperlink_cells) + '</hyperlinks>' if hyperlink_cells else ''
    sheet = f'''<worksheet xmlns="{NS['s']}" xmlns:r="{REL_NS[1:-1]}"><sheetViews><sheetView workbookViewId="0" showGridLines="0"><pane xSplit="2" ySplit="5" topLeftCell="C6" activePane="bottomRight" state="frozen"/></sheetView></sheetViews><sheetFormatPr defaultRowHeight="24"/><cols>{columns}</cols><sheetData>{''.join(content)}</sheetData><autoFilter ref="A5:{col(len(EXPORT_HEADERS))}{last}"/><mergeCells count="4">{merges}</mergeCells>{hyperlinks}<pageMargins left="0.25" right="0.25" top="0.4" bottom="0.4" header="0.2" footer="0.2"/><pageSetup orientation="landscape" paperSize="9" fitToWidth="2" fitToHeight="0"/></worksheet>'''
    mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml'
    package = io.BytesIO()
    with ZipFile(package, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', f'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="{mime}.sheet.main+xml"/><Override PartName="/xl/styles.xml" ContentType="{mime}.styles+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="{mime}.worksheet+xml"/></Types>')
        archive.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        archive.writestr('xl/workbook.xml', f'<workbook xmlns="{NS["s"]}" xmlns:r="{REL_NS[1:-1]}"><sheets><sheet name="Opportunities" sheetId="1" r:id="rId1"/></sheets><definedNames><definedName name="_xlnm.Print_Titles" localSheetId="0">\'Opportunities\'!$1:$5</definedName></definedNames></workbook>')
        archive.writestr('xl/_rels/workbook.xml.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>')
        archive.writestr('xl/styles.xml', ET.tostring(export_styles, encoding='utf-8'))
        archive.writestr('xl/worksheets/sheet1.xml', sheet)
        if hyperlink_relationships:
            archive.writestr('xl/worksheets/_rels/sheet1.xml.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + ''.join(hyperlink_relationships) + '</Relationships>')
    return package.getvalue()
