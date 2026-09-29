"""Supplier quotes in the main application database; types stay internal."""
from pathlib import Path
from datetime import datetime
from collections import OrderedDict
from contextlib import contextmanager
from decimal import Decimal,InvalidOperation
import hashlib,io,json,os,re,sqlite3,threading
from .database import database_path, write_lock

LOCK=write_lock
def db_path():
    return database_path()

def ensure_schema(conn):
    conn.executescript('''CREATE TABLE IF NOT EXISTS quotes(id INTEGER PRIMARY KEY,reference TEXT NOT NULL,project TEXT,quote_date TEXT,source TEXT,content_hash TEXT NOT NULL,imported_at TEXT NOT NULL,total_ex REAL,UNIQUE(reference,quote_date,content_hash));
        CREATE TABLE IF NOT EXISTS groups(id INTEGER PRIMARY KEY,quote_id INTEGER NOT NULL REFERENCES quotes(id),type_code TEXT,base_qty REAL,description TEXT,supplier TEXT,cat_no TEXT,search_text TEXT,source_group TEXT);
        CREATE TABLE IF NOT EXISTS lines(id INTEGER PRIMARY KEY,group_id INTEGER NOT NULL REFERENCES groups(id),line_order INTEGER,item_no TEXT,type_raw TEXT,type_effective TEXT,supplier TEXT,quantity REAL,cat_no TEXT,description TEXT,raw_description TEXT,unit_ex REAL,total_ex REAL,status TEXT,page INTEGER,source_row INTEGER,connection_basis TEXT,source_values TEXT);
        CREATE INDEX IF NOT EXISTS idx_lighting_lines_group ON lines(group_id);CREATE INDEX IF NOT EXISTS idx_lighting_groups_quote ON groups(quote_id);''')

    _lighting_extensions(conn)

@contextmanager
def connect():
    # Use the same automatic migration path as every other application route.
    from .database import connect as main_connect
    with main_connect() as conn:
        yield conn

def clean(value):return re.sub(r'\s+',' ',str(value or '')).strip()
def number(value):
    if value is None or str(value).strip().upper() in ('','-','\u2014','\u2013','N/A','N.A.','NA','TBA','TBC','POA','NOT QUOTED','INCLUDED','$ -') :return None
    text=str(value).strip().replace(',','').replace('$','')
    try:
        d=Decimal(text)
        if not d.is_finite() or abs(d)>Decimal('1e12'):raise ValueError('Price or quantity is outside the supported range')
        return float(d)
    except InvalidOperation:raise ValueError('Invalid price or quantity: '+text[:30])
def date_text(value):
    if hasattr(value,'strftime'):return value.strftime('%Y-%m-%d')
    value=clean(value)
    for fmt in ('%Y-%m-%d','%Y-%m-%d %H:%M:%S','%A, %d %B %Y','%d/%m/%Y','%d %B %Y'):
        try:return datetime.strptime(value,fmt).strftime('%Y-%m-%d')
        except ValueError:pass
    return value

def effective_type(raw,description):
    if clean(raw):return clean(raw).upper()
    match=re.search(r'\bType\s+([A-Z][A-Z0-9_./-]*(?:\s+OPT)?)\b',description,re.I)
    return clean(match[1]).upper() if match else ''
def display_description(description):
    return clean(re.sub(r'(?:\|\s*)?\bType\s+[A-Z][A-Z0-9_./-]*(?:\s+OPT)?\b','',description,flags=re.I))

def normal_line(values,index):
    desc=str(values.get('Description') or '').strip();cat=clean(values.get('Cat No'));qty=number(values.get('Qty'));unit=number(values.get('Unit Ex GST'));total=number(values.get('Total Ex GST'))
    # Commercial notes without pricing are excluded; priced freight is retained even if Cat No says NOTE.
    if not desc or (qty is None and unit is None and total is None):return None
    if cat.upper() in ('NOTE','NOTES') and unit is None and total is None:return None
    if not cat and not clean(values.get('Supplier')):return None
    if any(v is not None and v<0 for v in (qty,unit,total)):raise ValueError('Negative quote amounts need review before importing')
    if total is None and qty is not None and unit is not None:total=round(qty*unit,2)
    typ=effective_type(values.get('Type'),desc)
    return dict(item_no=clean(values.get('Item')),type_raw=clean(values.get('Type')),type_effective=typ,
        supplier=clean(values.get('Supplier')),quantity=qty,cat_no=cat,description=display_description(desc),raw_description=desc,
        unit_ex=unit,total_ex=total,status=clean(values.get('Status')) or ('Not quoted' if unit is None else 'Quoted'),
        source_values=json.dumps(values,default=str,ensure_ascii=False),page=int(number(values.get('PDF Page')) or 0),source_row=index,source_group=clean(values.get('Connection Group')))

def parse_xlsx(content):
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(content),data_only=True,read_only=True)
    try:
        if 'All Items' not in wb.sheetnames:raise ValueError('Use a connected price workbook with an All Items sheet')
        if wb['All Items'].max_row and wb['All Items'].max_row > 50000:
            raise ValueError('Split the workbook into sheets with at most 50,000 rows')
        rows=wb['All Items'].iter_rows(values_only=True);headers=next(rows, None)
        if not headers:raise ValueError('The All Items sheet is empty')
        needed={'Project','Quote Ref','Quote Date','Supplier','Qty','Cat No','Description','Unit Ex GST','Total Ex GST'}
        if not needed.issubset(set(headers)):raise ValueError('The All Items sheet is missing required quote columns')
        quotes=OrderedDict();excluded=0
        for i,row in enumerate(rows,2):
            values=dict(zip(headers,row));ref=clean(values.get('Quote Ref'))
            if not ref:continue
            line=normal_line(values,i)
            if line is None:excluded+=1;continue
            key=(ref,date_text(values.get('Quote Date')))
            quote=quotes.setdefault(key,dict(reference=ref,quote_date=key[1],project=clean(values.get('Project')),source=clean(values.get('Source PDF')),total_ex=None,lines=[]))
            quote['lines'].append(line)
        if 'Project Index' in wb.sheetnames:
            for row in wb['Project Index'].iter_rows(min_row=4,values_only=True):
                if len(row)>6:
                    key=(clean(row[3]),date_text(row[2]))
                    if key in quotes:quotes[key]['total_ex']=number(row[6])
        if not quotes:raise ValueError('No supported price rows found')
        return list(quotes.values()),excluded
    finally:wb.close()

def parse_pdf(content,filename):
    from .quotation_parser import parse_pdf as parse
    return parse(content, filename)


def group_lines(lines):
    groups=OrderedDict();previous=None
    for index,line in enumerate(lines):
        explicit=line['type_effective']
        if explicit:key=explicit;basis='Same type' if key in groups else 'Primary'
        elif previous is not None:key=previous;basis='Blank type → previous'
        else:key=f'__untyped_{index}';basis='Standalone; no preceding type'
        # Description-embedded type is authoritative when the separate Type column is empty.
        if not line['type_raw'] and explicit:basis+=' (type in description)'
        if explicit:previous=key
        groups.setdefault(key,[]).append({**line,'type_effective':explicit or ('' if key.startswith('__untyped_') else key),'connection_basis':basis})
    return groups

def quote_digest(quote):
    # Same quoted line re-uploaded through XLSX or PDF must not create another record.
    fields=('item_no','type_effective','supplier','quantity','cat_no','description','unit_ex','total_ex')
    data=[{k:clean(l[k]).casefold() if isinstance(l[k],str) else l[k] for k in fields} for l in quote['lines']]
    return hashlib.sha256(json.dumps(data,sort_keys=True,ensure_ascii=False).encode()).hexdigest()

def import_content(filename,content):
    # Programmatic helper; browser imports always use preview -> confirmation.
    from .quotation_parser import parse_content
    quotes, excluded = parse_content(filename, content)
    return save_quotes(quotes, excluded)


def insert_quote(conn, quote, digest, canonical_hash=None, file_hash=None):
    qid=conn.execute('INSERT INTO quotes(reference,project,quote_date,source,content_hash,imported_at,total_ex,canonical_hash,file_sha256) VALUES(?,?,?,?,?,?,?,?,?)',
        (quote['reference'],quote['project'],quote['quote_date'],quote['source'],digest,datetime.now().isoformat(),quote['total_ex'],canonical_hash,file_hash)).lastrowid
    group_count = line_count = 0
    for typ,members in group_lines(quote['lines']).items():
        primary=members[0];search=' '.join([quote['project'],quote['reference']]+[str(l.get(k) or '') for l in members for k in ('supplier','cat_no','description')]).casefold()
        gid=conn.execute('INSERT INTO groups(quote_id,type_code,base_qty,description,supplier,cat_no,search_text,source_group) VALUES(?,?,?,?,?,?,?,?)',
            (qid,typ,primary['quantity'],primary['description'],primary['supplier'],primary['cat_no'],search,primary.get('source_group',''))).lastrowid
        for index,line in enumerate(members):
            keys=('item_no','type_raw','type_effective','supplier','quantity','cat_no','description','raw_description','unit_ex','total_ex','status','page','source_row','connection_basis','source_values')
            conn.execute('INSERT INTO lines(group_id,line_order,'+','.join(keys)+') VALUES('+','.join('?' for _ in range(len(keys)+2))+')',(gid,index,*(line.get(k) for k in keys)))
        group_count+=1;line_count+=len(members)
    return qid, group_count, line_count


def save_quotes(quotes,excluded=0):
    from .lighting_imports import canonical_digest, classify_quotes
    report={'quotes_added':0,'quotes_skipped':0,'groups_added':0,'lines_added':0,'notes_excluded':excluded}
    with LOCK:
        # Backup outside a write transaction: the backup connection must never
        # wait on a writer that is itself waiting for the backup to finish.
        with connect() as check:
            preview = classify_quotes(check, quotes)
            pending = [q for q, p in zip(quotes, preview) if p['status'] != 'identical']
            have_data = check.execute('SELECT COUNT(*) FROM quotes').fetchone()[0]
        if pending and have_data:
            from .backups import create_backup
            create_backup('before-import')
        with connect() as conn:
            conn.execute('BEGIN IMMEDIATE')
            for quote in quotes:
                # Recheck inside the writer transaction for concurrent imports.
                if classify_quotes(conn, [quote])[0]['status'] == 'identical':
                    report['quotes_skipped'] += 1; continue
                _, groups, lines = insert_quote(conn, quote, quote_digest(quote), canonical_digest(quote))
                report['quotes_added'] += 1; report['groups_added'] += groups; report['lines_added'] += lines
    return report


def _lighting_extensions(conn):
    from .release_v23 import add_columns
    add_columns(conn, 'groups', {
        'nickname': "TEXT NOT NULL DEFAULT ''", 'notes': "TEXT NOT NULL DEFAULT ''",
        'primary_unit': "TEXT NOT NULL DEFAULT 'primary unit'",
        'project_override': 'TEXT', 'reference_override': 'TEXT', 'quote_date_override': 'TEXT',
        'revision': 'INTEGER NOT NULL DEFAULT 0', 'updated_at': 'TEXT',
        'updated_by': 'INTEGER REFERENCES users(id) ON DELETE SET NULL',
    })
    add_columns(conn, 'lines', {
        'included': 'INTEGER NOT NULL DEFAULT 1',
        'total_mode': "TEXT NOT NULL DEFAULT 'quoted'",
    })
    conn.execute('''CREATE TABLE IF NOT EXISTS lighting_photos(
        group_id INTEGER PRIMARY KEY REFERENCES groups(id) ON DELETE CASCADE,
        mime TEXT NOT NULL, image BLOB NOT NULL, updated_at TEXT NOT NULL)''')
    conn.execute('''CREATE TABLE IF NOT EXISTS lighting_edit_history(
        id INTEGER PRIMARY KEY, group_id INTEGER NOT NULL REFERENCES groups(id) ON DELETE CASCADE,
        user_id INTEGER REFERENCES users(id) ON DELETE SET NULL, actor_name TEXT NOT NULL,
        before_json TEXT NOT NULL, after_json TEXT NOT NULL, created_at TEXT NOT NULL)''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_lighting_history ON lighting_edit_history(group_id,id)')


def _money_product(quantity, unit):
    if quantity is None or unit is None:
        return None
    from decimal import ROUND_HALF_UP
    return float((Decimal(str(quantity)) * Decimal(str(unit))).quantize(Decimal('.01'), rounding=ROUND_HALF_UP))


def _money_sum(values):
    return float(sum((Decimal(str(v)) for v in values if v is not None), Decimal(0)))


def _price_details(item, components):
    """One calculation path for search cards, detail, editor and regression tests."""
    base = item['base_qty']
    for line in components:
        calculated = _money_product(line['quantity'], line['unit_ex'])
        line['calculated_total_ex'] = calculated
        line['resolved_total_ex'] = calculated if line['total_mode'] == 'calculated' else line['total_ex']
        line['difference_ex'] = (float(Decimal(str(line['total_ex'])) - Decimal(str(calculated)))
                                 if line['total_ex'] is not None and calculated is not None else None)
        line['has_difference'] = line['difference_ex'] is not None and abs(line['difference_ex']) >= .009
        line['quantity_per_primary'] = line['quantity'] / base if base and base > 0 and line['quantity'] is not None else None
        line['contribution_unit_ex'] = (line['resolved_total_ex'] / base
            if base and base > 0 and line['included'] and line['resolved_total_ex'] is not None else None)
        # Original import values are retained separately from editable working values.
        try:
            values = json.loads(line.pop('source_values', '{}') or '{}')
            line['original'] = {k: values.get(k) for k in ('Supplier','Cat No','Description','Qty','Unit Ex GST','Total Ex GST')}
        except (ValueError, TypeError):
            line['original'] = {}
    selected = [line for line in components if line['included']]
    missing = sum(line['resolved_total_ex'] is None for line in selected)
    item.update({
        'component_count': len(components), 'included_count': len(selected),
        'missing_prices': missing, 'components': components,
        'connected_total_ex': _money_sum(line['resolved_total_ex'] for line in selected),
        'calculated_total_ex': _money_sum(line['calculated_total_ex'] for line in selected),
        'missing_calculated': sum(line['calculated_total_ex'] is None for line in selected),
        'difference_count': sum(line['has_difference'] and line['total_mode'] == 'quoted' for line in selected),
        'implicit_links': sum('previous' in (line['connection_basis'] or '').lower() for line in selected),
        'main_row_unit_ex': components[0]['unit_ex'] if components else None,
    })
    item['connected_unit_ex'] = item['connected_total_ex'] / base if base and base > 0 and not missing and selected else None
    item['base_differs'] = bool(components and base != components[0]['quantity'])
    return item


def _get_items(conn, identifiers):
    if not identifiers:
        return []
    marks = ','.join('?' for _ in identifiers)
    rows = conn.execute(f'''SELECT g.*,
        COALESCE(g.reference_override,q.reference) AS reference,
        COALESCE(g.project_override,q.project) AS project,
        COALESCE(g.quote_date_override,q.quote_date) AS quote_date,
        q.source,q.reference AS original_reference,q.project AS original_project,q.quote_date AS original_quote_date,
        u.name AS edited_by,EXISTS(SELECT 1 FROM lighting_photos p WHERE p.group_id=g.id) AS has_photo
        FROM groups g JOIN quotes q ON q.id=g.quote_id LEFT JOIN users u ON u.id=g.updated_by
        WHERE g.id IN ({marks})''', identifiers).fetchall()
    groups = {r['id']: dict(r) for r in rows}
    members = {ident: [] for ident in identifiers}
    for row in conn.execute(f'''SELECT * FROM lines WHERE group_id IN ({marks}) ORDER BY line_order,id''', identifiers):
        members[row['group_id']].append(dict(row))
    return [_price_details(groups[ident], members[ident]) for ident in identifiers if ident in groups]


def search(query='', limit=40, offset=0, supplier=''):
    terms = clean(query).casefold().split()[:12]
    offset = max(0, int(offset)); limit = min(100, max(1, int(limit)))
    where = ' AND '.join('instr(lower(g.search_text || \' \' || g.nickname || \' \' || g.description),?)>0' for _ in terms) or '1=1'
    suppliers = [supplier] if isinstance(supplier, str) else supplier
    if not isinstance(suppliers, (list, tuple)) or len(suppliers)>200:raise ValueError('Select up to 200 suppliers')
    suppliers = list(dict.fromkeys(clean(s) for s in suppliers if clean(s)))
    if suppliers:
        where += ' AND (' + ' OR '.join('(g.supplier=? COLLATE NOCASE OR EXISTS (SELECT 1 FROM lines sl WHERE sl.group_id=g.id AND sl.supplier=? COLLATE NOCASE))' for _ in suppliers) + ')'
        terms.extend(value for s in suppliers for value in (s,s))
    with connect() as conn:
        total = conn.execute(f'SELECT COUNT(*) FROM groups g WHERE {where}', terms).fetchone()[0]
        ids = [r[0] for r in conn.execute(f'''SELECT g.id FROM groups g JOIN quotes q ON q.id=g.quote_id
            WHERE {where} ORDER BY COALESCE(g.quote_date_override,q.quote_date) DESC,q.id DESC,g.id
            LIMIT ? OFFSET ?''', (*terms, limit, offset))]
        results = _get_items(conn, ids)
        for item in results:
            item.pop('components', None)
            item.pop('search_text', None)
        stats = {name: conn.execute(f'SELECT COUNT(*) FROM {table}').fetchone()[0]
                 for name, table in [('quotes','quotes'),('groups','groups'),('lines','lines')]}
        suppliers = [r[0] for r in conn.execute("SELECT supplier FROM lines WHERE supplier<>'' UNION SELECT supplier FROM groups WHERE supplier<>'' ORDER BY supplier COLLATE NOCASE")]
        return {'results': results, 'total': total, 'offset': offset, 'stats': stats, 'suppliers': suppliers}


def detail(ident, conn=None):
    if conn is None:
        with connect() as opened:
            return detail(ident, opened)
    items = _get_items(conn, [ident])
    if not items:
        raise LookupError('Lighting price group not found')
    item = items[0]
    item['edit_count'] = conn.execute('SELECT COUNT(*) FROM lighting_edit_history WHERE group_id=?', (ident,)).fetchone()[0]
    return item


def _field(value, label, maximum=1000, required=False):
    if not isinstance(value, str):
        raise ValueError(label + ' must be text')
    value = value.strip()
    if len(value) > maximum or (required and not value):
        raise ValueError(f'{label} must contain {"1 to " if required else "at most "}{maximum} characters')
    return value


def _positive(value, label):
    if isinstance(value, bool):
        raise ValueError(label + ' must be a number')
    value = number(value)
    if value is not None and not 0 <= value <= 1e12:
        raise ValueError(label + ' must be between zero and 1,000,000,000,000')
    return value


def _prepare_photo(data_url):
    from .images import prepare_image
    return prepare_image(data_url)


def update(conn, ident, payload, user):
    from .accounts import AccessError
    from .utils import utc_now
    before = detail(ident, conn)
    if type(payload.get('expected_revision')) is not int or payload['expected_revision'] != before['revision']:
        raise AccessError(409, 'This light was changed by someone else. Close the editor, reopen it and apply your changes again.')
    allowed = {'expected_revision','description','nickname','notes','base_qty','primary_unit',
               'project','reference','quote_date','components','photo_data_url','remove_photo'}
    if set(payload) - allowed:
        raise ValueError('Unknown lighting edit field')
    values = {key: _field(payload.get(key, before[key]) or '', label, length, required)
              for key,label,length,required in [
                  ('description','Light name',1000,True),('nickname','Nickname',120,False),
                  ('notes','Notes',10000,False),('primary_unit','Primary unit label',40,True),
                  ('project','Project',500,False),('reference','Quote reference',250,True),
                  ('quote_date','Quote date',10,False)]}
    if values['quote_date']:
        from datetime import date
        try:
            date.fromisoformat(values['quote_date'])
        except ValueError:
            raise ValueError('Quote date must be YYYY-MM-DD') from None
    values['base_qty'] = _positive(payload.get('base_qty', before['base_qty']), 'Base quantity')
    members = payload.get('components')
    if not isinstance(members, list) or not 1 <= len(members) <= 2000:
        raise ValueError('Provide between 1 and 2,000 component rows')
    existing = {r['id'] for r in before['components']}
    seen = set(); normalized = []
    for row in members:
        if not isinstance(row, dict):
            raise ValueError('Invalid component row')
        rid = row.get('id')
        if rid is not None:
            if type(rid) is not int or rid not in existing or rid in seen:
                raise ValueError('Component rows must belong to this light and must not repeat')
            seen.add(rid)
        included = row.get('included', True)
        if included not in (True,False,0,1):
            raise ValueError('Choose whether to include the component')
        mode = row.get('total_mode','quoted')
        if mode not in ('quoted','calculated'):
            raise ValueError('Choose Quoted total or Qty x rate')
        normalized.append({
            'id': rid,
            'description': _field(row.get('description',''), 'Component description', 4000, True),
            'supplier': _field(row.get('supplier',''), 'Supplier', 250),
            'cat_no': _field(row.get('cat_no',''), 'Catalogue number', 250),
            'quantity': _positive(row.get('quantity'), 'Component quantity'),
            'unit_ex': _positive(row.get('unit_ex'), 'Unit price'),
            'total_ex': _positive(row.get('total_ex'), 'Quoted row total'),
            'status': _field(row.get('status',''), 'Pricing status', 100),
            'included': int(included), 'total_mode': mode,
        })
    if seen != existing:
        raise ValueError('Existing rows must be retained. Untick Include to exclude a row without losing its source record.')
    if payload.get('photo_data_url') and payload.get('remove_photo'):
        raise ValueError('Choose either a new photo or Remove photo')
    photo = _prepare_photo(payload['photo_data_url']) if payload.get('photo_data_url') else None
    now = utc_now()
    for order, row in enumerate(normalized):
        fields = [k for k in row if k != 'id']
        if row['id'] is None:
            row['id'] = conn.execute('INSERT INTO lines(group_id,line_order,connection_basis,source_values,'+
                ','.join(fields)+') VALUES('+','.join('?' for _ in range(len(fields)+4))+')',
                (ident,order,'Manually added','{}',*(row[k] for k in fields))).lastrowid
        else:
            conn.execute('UPDATE lines SET '+','.join(k+'=?' for k in fields)+' WHERE id=? AND group_id=?',
                         (*(row[k] for k in fields),row['id'],ident))
    first = normalized[0]
    conn.execute('''UPDATE groups SET description=?,nickname=?,notes=?,base_qty=?,primary_unit=?,
        supplier=?,cat_no=?,project_override=?,reference_override=?,quote_date_override=?,revision=revision+1,
        updated_at=?,updated_by=? WHERE id=?''',
        (values['description'],values['nickname'],values['notes'],values['base_qty'],values['primary_unit'],
         first['supplier'],first['cat_no'],values['project'],values['reference'],values['quote_date'],now,user['id'],ident))
    searchable = ' '.join([values['description'],values['nickname'],values['notes'],values['project'],values['reference'],
        before['original_project'] or '',before['original_reference'] or ''] +
        [str(row.get(k) or '') for row in normalized for k in ('supplier','cat_no','description')]).casefold()
    conn.execute('UPDATE groups SET search_text=? WHERE id=?', (searchable, ident))
    if photo:
        conn.execute('''INSERT INTO lighting_photos VALUES (?,?,?,?) ON CONFLICT(group_id)
            DO UPDATE SET mime=excluded.mime,image=excluded.image,updated_at=excluded.updated_at''', (ident,*photo,now))
    elif payload.get('remove_photo'):
        conn.execute('DELETE FROM lighting_photos WHERE group_id=?', (ident,))
    after = detail(ident, conn)
    # Save working snapshots, never duplicate binary photos or credentials in history.
    conn.execute('''INSERT INTO lighting_edit_history(group_id,user_id,actor_name,before_json,after_json,created_at)
        VALUES (?,?,?,?,?,?)''', (ident,user['id'],user['name'],json.dumps(before),json.dumps(after),now))
    after['edit_count'] += 1
    return after
