"""Incremental, transactional workbook import. Never infer identity from row position."""
import hashlib
import io
import json
import math
import re
import uuid
import posixpath
from xml.etree import ElementTree as ET
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from .tenders import ValidationError, TEXT_FIELDS, date_value, text, timestamp

MAX_BYTES = 15 * 1024 * 1024
MAX_ROWS = 20000
ALIASES = {
    'submission id':'id', 'project':'project', 'project name':'project',
    'builder':'builder', 'builders':'builder', 'tec status':'status', 'tec quote':'status',
    'builder outcome':'builderStatus', 'project status (builders)':'builderStatus',
    'pricing feedback':'pricing', 'comment on value':'pricing',
    'quote value (aud)':'value', 'tender value':'value',
    'submission deadline':'due', 'due date':'due', 'region':'region',
    'location':'location', 'project location':'location', 'contact':'contact',
    'contact person':'contact', 'email':'email', 'submission month':'month',
    'month of submission':'month', 'status updated':'statusUpdated',
    'latest update':'latest', 'update chase up':'latest', 'remarks':'remarks',
    'chase up no.':'chaseNumber',
}

def norm(v): return ' '.join(str(v or '').split()).casefold()
def pair(v): return (norm(v.get('project')), norm(v.get('builder')))
def string(v):
    if isinstance(v, (date, datetime)): return v.isoformat()[:10]
    return text('' if v is None else str(v), 'Cell')
def iso(v, timed=False):
    if v in (None, ''): return ''
    if isinstance(v, (int, float)):
        from openpyxl.utils.datetime import from_excel
        v = from_excel(v)
    if isinstance(v, datetime): return v.strftime('%Y-%m-%dT%H:%M' if timed else '%Y-%m-%d')
    if isinstance(v, date): return v.isoformat() + ('T09:00' if timed else '')
    v = string(v)
    try: return date_value(v, timed=timed)
    except ValidationError:
        for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%d %b %Y', '%d %B %Y'):
            try: return datetime.strptime(v, fmt).strftime('%Y-%m-%dT09:00' if timed else '%Y-%m-%d')
            except ValueError: pass
    raise ValidationError('Invalid date: ' + v)

def parse_workbook(filename, content):
    if Path(filename).suffix.lower() not in ('.xlsx', '.xlsm'):
        raise ValidationError('Choose an Excel .xlsx or .xlsm workbook.')
    if not content or len(content) > MAX_BYTES:
        raise ValidationError('Choose a workbook smaller than 15 MB.')
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 5000 or sum(i.file_size for i in archive.infolist()) > 100*1024*1024:
                raise ValidationError('Workbook is too large after decompression.')
            ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            relationships=ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
            targets={r.attrib['Id']:r.attrib['Target'] for r in relationships}
            workbook_xml=ET.fromstring(archive.read('xl/workbook.xml'))
            sheets=workbook_xml.findall('s:sheets/s:sheet',ns)
            selected={'raw data'} if any(norm(s.attrib['name'])=='raw data' for s in sheets) else {'tenders','follow-ups','scheduled'}
            for s in sheets:
                if norm(s.attrib['name']) not in selected:continue
                target=targets[s.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
                path=target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
                root=ET.fromstring(archive.read(path))
                for c in root.findall('.//s:sheetData/s:row/s:c',ns):
                    if c.find('s:f',ns) is not None:
                        value=c.find('s:v',ns)
                        # Excel writes a calculated empty string as t="str" with <v/>.
                        # This is a valid result, unlike a missing numeric cache.
                        empty_string=value is not None and c.attrib.get('t')=='str' and value.text is None
                        if value is None or (value.text is None and not empty_string) or c.attrib.get('t')=='e':
                            raise ValidationError(f'{s.attrib["name"]}!{c.attrib.get("r","")}: formula has no valid saved result. Recalculate and save the workbook in Excel first.')
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except ValidationError: raise
    except (BadZipFile, ValueError, KeyError, OSError, EOFError, ET.ParseError) as exc:
        raise ValidationError('Cannot read this workbook. Save it as .xlsx and try again.') from exc
    try:
        sheet = next((s for s in wb if norm(s.title)=='raw data'), None)
        raw = sheet is not None
        if sheet is None: sheet = next((s for s in wb if norm(s.title)=='tenders'), None)
        if sheet is None: raise ValidationError('Workbook needs a Raw Data or Tenders sheet.')
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_row=MAX_ROWS+51, max_col=100, values_only=True)
        fields = None; records=[]; history=[]
        for rn, row in enumerate(rows, 1):
            if fields is None:
                mapping={i:ALIASES[norm(v)] for i,v in enumerate(row) if norm(v) in ALIASES}
                if {'project','builder','status','value'} <= set(mapping.values()):
                    fields=mapping; headers=row
                elif rn>=50: break
                continue
            if not any(v is not None for v in row): continue
            if len(records)>=MAX_ROWS: raise ValidationError('Import at most 20,000 submissions per workbook.')
            try:
                item={key:string(row[i]) for i,key in fields.items() if key not in ('value','due')}
                for key in ('project','builder'):
                    if not item.get(key): raise ValidationError(key.title()+' is required.')
                item['id']=text(item.get('id'),'Submission ID',limit=100)
                item['status']={'closed':'Close', **{s.lower():s for s in ('Open','Won','Lost','Close','Hold')}}.get(norm(item['status']))
                if not item['status']: raise ValidationError('TEC status must be Open, Won, Lost, Closed or Hold.')
                for field, choices in [('builderStatus', ('Awarded','Unsuccessful','In Progress','Hold','Builder Withdrawn','Suspended')), ('pricing', ('Nil','Competitive','High','Mid-Range','Lower','Mix'))]:
                    if field in item:
                        item[field]={norm(v):v for v in choices}.get(norm(item[field]),item[field])
                if item.get('pricing') in ('4','4.0'): item['pricing']='High'
                if norm(item.get('pricing'))=='no feedback': item['pricing']='Nil'
                vi=next(i for i,k in fields.items() if k=='value'); value=row[vi]
                item['value']=None if value in (None,'') else float(str(value).replace('$','').replace(',','').strip())
                if item['value'] is not None and (not math.isfinite(item['value']) or item['value']<0): raise ValidationError('Value must be zero or greater.')
                if 'due' in fields.values(): item['due']=iso(row[next(i for i,k in fields.items() if k=='due')])
                item['_row']=rn; records.append(item)
                if raw:
                    for i,h in enumerate(headers):
                        match=re.match(r'(\d+)(?:st|nd|rd|th)\s+chased\s+up\s*-\s*date', norm(h))
                        if match and row[i] is not None and i+1<len(row) and row[i+1]:
                            history.append({'_record':len(records)-1,'_slot':int(match[1]),'date':iso(row[i]),'note':string(row[i+1]),'method':'Imported','author':'Yunus'})
            except (ValueError, TypeError, ValidationError) as exc:
                raise ValidationError(f'{sheet.title}, row {rn}: {exc}') from exc
        if fields is None or not records: raise ValidationError('No tender records found under the expected column headings.')
        # Exported history uses stable IDs. Absent sheets never clear existing history.
        for name,kind in [('Follow-ups','chases'),('Scheduled','reminders')]:
            if raw or name not in wb.sheetnames: continue
            ws=wb[name];ws.reset_dimensions(); mapping=None
            for rn,row in enumerate(ws.iter_rows(max_row=MAX_ROWS+51,max_col=30,values_only=True),1):
                if mapping is None:
                    keys={norm(v):i for i,v in enumerate(row) if v is not None}
                    if 'submission id' in keys: mapping=keys
                    elif rn>=50: break
                    continue
                if not any(v is not None for v in row): continue
                if len(history)>=MAX_ROWS*4: raise ValidationError('Too many follow-up records.')
                def get(k): return row[mapping[k]] if k in mapping else None
                sid=string(get('submission id'))
                if not sid: raise ValidationError(f'{name}, row {rn}: Submission ID is required.')
                entry={'_kind':kind,'submissionId':sid,'id':string(get('follow-up id' if kind=='chases' else 'reminder id'))}
                if kind=='chases': entry.update(date=iso(get('recorded date')),note=string(get('update')),method=string(get('method')),author=string(get('recorded by')))
                else:
                    state=norm(get('status'))
                    if state not in ('pending','completed'): raise ValidationError(f'{name}, row {rn}: invalid reminder status.')
                    entry.update(due=iso(get('scheduled time (sydney)'),True),note=string(get('reminder note')),done=state=='completed')
                if not entry.get('date',entry.get('due')) or not entry['note']: raise ValidationError(f'{name}, row {rn}: date and note are required.')
                history.append(entry)
        return {'records':records,'history':history,'raw':raw}
    except ValidationError: raise
    except (ValueError, TypeError, KeyError, ET.ParseError) as exc:
        raise ValidationError('Workbook contains invalid cells. Check dates and numbers, then save it again in Excel.') from exc
    finally: wb.close()

def import_workbook(store, filename, content, actor, expected=None):
    parsed=parse_workbook(filename,content)
    digest=hashlib.sha256(content).hexdigest()
    with store.connect(write=expected is not None) as db:
        if expected is None: db.execute('BEGIN')
        generation=db.execute('SELECT value FROM tender_generation').fetchone()[0]
        token=hashlib.sha256(f'{digest}:{generation}'.encode()).hexdigest()
        if expected is not None and expected!=token:
            raise ValidationError('Tender data changed after preview. Preview the workbook again before importing.',409)
        current={r['id']:{**json.loads(r['body']),'revision':r['revision']} for r in db.execute('SELECT * FROM submissions')}
        by_pair={}
        for s in current.values(): by_pair.setdefault(pair(s),[]).append(s['id'])
        saved=db.execute("SELECT body FROM settings WHERE id='excel-import-baselines'").fetchone()
        baselines=json.loads(saved['body']) if saved else {}
        # Seed is the baseline for the original workbook, so unchanged cells do not undo app edits.
        seed={s['id']:s for s in store.seed['submissions']}
        changes=[]; seen={}; resolved=[]; counts={'added':0,'updated':0,'unchanged':0,'duplicates':0,'followupsAdded':0,'followupsUpdated':0,'remindersAdded':0,'remindersUpdated':0}
        for incoming in parsed['records']:
            sid=incoming.get('id'); key=pair(incoming)
            if not sid:
                candidates=by_pair.get(key,[])
                if len(candidates)>1: raise ValidationError(f'Row {incoming["_row"]}: multiple submissions match this project and builder. Import an export with Submission IDs.')
                sid=candidates[0] if candidates else str(uuid.uuid5(uuid.NAMESPACE_URL,'tec-import:'+json.dumps(key)))
            values={k:v for k,v in incoming.items() if k not in ('id','_row')}
            if sid in seen:
                if seen[sid]!=values: raise ValidationError(f'Row {incoming["_row"]}: conflicting duplicate submission. Keep one row per submission.')
                counts['duplicates']+=1;resolved.append(sid);continue
            seen[sid]=values;resolved.append(sid)
            old=current.get(sid)
            base=baselines.get(sid,seed.get(sid,{}) if parsed['raw'] else {})
            updates={k:v for k,v in values.items() if old is None or (v!=old.get(k) and (k not in base or v!=base[k]))}
            if old and values.get('statusUpdated')=='' and '1900' in old.get('statusUpdated',''):
                updates.pop('statusUpdated',None)  # Legacy missing-date sentinel exports as blank.
            action='added' if old is None else 'updated' if updates else 'unchanged'
            counts[action]+=1
            if updates:
                changes.append({'action':action,'project':values['project'],'builder':values['builder'],'fields':list(updates)})
            baselines[sid]={**base,**values}
            if old is None:
                item={**{k:'' for k in TEXT_FIELDS},'latest':'','chaseNumber':'',**values,'id':sid,'revision':0,'createdBy':actor['id'],'createdAt':timestamp()}
            else: item={**old,**updates,'revision':old['revision']+(1 if updates else 0)}
            if old is None:
                item.update(origin='Excel import',importedBy=actor['id'],importedByName=actor['name'],importedAt=timestamp())
            if old is None or updates:
                item.update(lastImportedBy=actor['id'],lastImportedByName=actor['name'],lastImportedAt=timestamp())
            current[sid]=item
            if sid not in by_pair.setdefault(key,[]): by_pair[key].append(sid)
            if expected is not None and (old is None or updates):
                db.execute('INSERT INTO submissions VALUES(?,?,?) ON CONFLICT(id) DO UPDATE SET body=excluded.body,revision=excluded.revision',(sid,json.dumps(item,ensure_ascii=False),item['revision']))
        known_history={kind:{r['id']:json.loads(r['body']) for r in db.execute('SELECT id,body FROM '+kind)} for kind in ('chases','reminders')}
        seen_history={}
        for entry in parsed['history']:
            kind=entry.get('_kind','chases');sid=resolved[entry['_record']] if '_record' in entry else entry['submissionId']
            if sid not in resolved: raise ValidationError('Follow-up references a submission missing from the Tenders sheet.')
            identifier=entry.get('id') or (f'{sid}-chase-{entry["_slot"]}' if '_slot' in entry else str(uuid.uuid5(uuid.NAMESPACE_URL,json.dumps(entry,sort_keys=True))))
            fields={k:v for k,v in entry.items() if not k.startswith('_') and k!='id'};fields['submissionId']=sid
            if kind=='chases' and fields.get('method')=='Imported' and fields.get('author') in ('Workbook',''):
                fields['author']='Yunus'
            hk=kind+':'+identifier
            if hk in seen_history:
                if seen_history[hk]!=fields: raise ValidationError('Conflicting duplicate follow-up IDs in workbook.')
                continue
            seen_history[hk]=fields
            old=known_history[kind].get(identifier)
            if old and old['submissionId']!=sid: raise ValidationError('A follow-up ID belongs to a different submission.')
            base=baselines.get(hk,{})
            if parsed['raw'] and not base: base=next((c for c in store.seed['chases'] if c['id']==identifier),{})
            if kind=='chases' and base.get('author')=='Workbook':
                base={**base,'author':'Yunus'}
            # Existing authorship is immutable; importing edits records the importer separately.
            if old and kind=='chases':
                fields['author']=old.get('author','Yunus')
            updates={k:v for k,v in fields.items() if old is None or (old.get(k)!=v and (k not in base or base[k]!=v))}
            baselines[hk]={**base,**fields}
            if not updates: continue
            counts[('followups' if kind=='chases' else 'reminders')+('Added' if old is None else 'Updated')]+=1
            item={**(old or {'id':identifier,'createdAt':timestamp()}),**updates}
            if kind=='chases':
                item['revision']=(old or {}).get('revision',0)+(1 if old else 0)
                if old:
                    item.setdefault('edits',[]).append({k:old.get(k) for k in ('note','date','method','updatedByName','updatedAt')})
                    item.update(updatedBy=actor['id'],updatedByName=actor['name'],updatedAt=timestamp())
                elif parsed['raw']:
                    item.update(legacyChase=True)
                item.update(lastImportedBy=actor['id'],lastImportedByName=actor['name'])
            if expected is not None:
                db.execute(f'INSERT INTO {kind} VALUES(?,?,?) ON CONFLICT(id) DO UPDATE SET body=excluded.body',(identifier,sid,json.dumps(item,ensure_ascii=False)))
        if expected is not None:
            if any(counts[k] for k in ('added','updated','followupsAdded','followupsUpdated','remindersAdded','remindersUpdated')):
                store.activity(db,actor['id'],actor['name'],'Imported Excel',details={'filename':filename,**counts})
            db.execute("INSERT INTO settings VALUES('excel-import-baselines',?) ON CONFLICT(id) DO UPDATE SET body=excluded.body",(json.dumps(baselines,ensure_ascii=False),))
        return {'ok':True,'token':token,'counts':counts,'changes':changes[:100],'totalChanges':len(changes),'applied':expected is not None}
