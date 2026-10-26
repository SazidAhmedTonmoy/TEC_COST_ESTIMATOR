"""Section inclusion preserves estimate data while consistently changing totals."""
import ast
import copy
import json
import re
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

import test_hours_markup as helpers
from conest_app import auto_upgrade, estimate_tools, productivity, project_json, project_sharing, projects
from conest_app.accounts import AccessError
from conest_app.database import _raw_connect, connect, database_path
from conest_app.estimate_sections import excluded_categories
from conest_app.exporter import export_project
from conest_app.workspace import project_contributors
from legacy_schema_fixture import drop_fixture_columns


class EstimateSectionTests(unittest.TestCase):
    setUp = helpers.HoursMarkupTests.setUp
    tearDown = helpers.HoursMarkupTests.tearDown

    def prepare(self, conn):
        pid = self.data['project_id']
        projects.update_project_line(conn, pid, self.data['second_line_id'], {'category_snapshot': 'Power'})
        projects.update_project(conn, pid, {'labor_hours_markup': .2})
        return pid, projects.calculate_project(conn, pid)

    def toggle(self, conn, category='Lighting', included=False, user=None, pid=None):
        pid = pid or self.data['project_id']
        revision = conn.execute('SELECT revision FROM projects WHERE id=?', (pid,)).fetchone()[0]
        estimate_tools.set_section_inclusion(conn, pid, user or self.data['owner'],
            {'category': category, 'included': included, 'expected_revision': revision})
        return projects.calculate_project(conn, pid)

    def test_totals_raw_values_cache_and_reenable(self):
        with connect() as conn:
            pid, before = self.prepare(conn)
            self.assertEqual(excluded_categories(before['project']['excluded_categories']), set())
            self.assertTrue(all(row['included_in_total'] for row in before['lines']))
            self.assertEqual(projects.project_totals(conn, pid), before['totals'])
            after = self.toggle(conn)
            remaining = before['lines'][1]
            for key, field in [('material', 'total_material'), ('hours', 'total_hours'),
                               ('labor', 'total_labor'), ('direct', 'direct_total'), ('quoted', 'quoted_total')]:
                self.assertAlmostEqual(after['totals'][key], remaining[field])
                self.assertAlmostEqual(sum(row[key] for row in after['work_summary']) if key in ('material','hours','labor','quoted')
                                       else after['totals'][key], after['totals'][key])
            self.assertAlmostEqual(after['totals']['overhead'], remaining['quoted_total']-remaining['direct_total'])
            self.assertEqual(after['totals']['line_count'], 2)
            self.assertEqual(after['totals']['included_line_count'], 1)
            self.assertAlmostEqual(after['crew_plan']['working_days'], remaining['total_hours']/16)
            self.assertEqual(after['lines'][0], {**before['lines'][0], 'included_in_total': False})
            self.assertEqual(after['lines'][1], before['lines'][1])
            self.assertFalse(next(row for row in after['scope_summary'] if row['category']=='Lighting')['included'])
            self.assertEqual(projects.project_totals(conn, pid), after['totals'])
            self.assertEqual(self.toggle(conn, included=True)['totals'], before['totals'])

    def test_all_off_persists_across_connection_without_changing_catalog(self):
        with connect() as conn:
            pid, before = self.prepare(conn)
            catalog = list(conn.execute('SELECT * FROM catalog_items ORDER BY id'))
            self.toggle(conn)
            after = self.toggle(conn, 'Power')
            self.assertTrue(all(after['totals'][key] == 0 for key in ('material','hours','labor','direct','quoted','overhead')))
            self.assertEqual(after['work_summary'], [])
            self.assertEqual(after['crew_plan']['working_days'], 0)
            self.assertEqual(list(conn.execute('SELECT * FROM catalog_items ORDER BY id')), catalog)
        with connect() as conn:
            persisted = projects.calculate_project(conn, pid)
            self.assertEqual(persisted['totals'], after['totals'])
            self.assertEqual(projects.list_projects(conn)[0]['totals'], after['totals'])
            self.assertEqual(len(persisted['lines']), len(before['lines']))

    def test_permission_validation_revision_noop_and_contribution(self):
        with connect() as conn:
            pid, before = self.prepare(conn)
            rev = before['project']['revision']
            body = {'category': 'Lighting', 'included': False, 'expected_revision': rev}
            for user in (self.data['viewer'],):
                with self.assertRaises(AccessError):estimate_tools.set_section_inclusion(conn,pid,user,body)
            for bad in (None,0,1,'false',[],{}):
                with self.assertRaises(ValueError):estimate_tools.set_section_inclusion(conn,pid,self.data['owner'],{**body,'included':bad})
            with self.assertRaises(ValueError):estimate_tools.set_section_inclusion(conn,pid,self.data['owner'],{'category':'Lighting','included':False})
            with self.assertRaises(LookupError):estimate_tools.set_section_inclusion(conn,pid,self.data['owner'],{**body,'category':'Missing'})
            self.assertEqual(projects.calculate_project(conn,pid),before)
            after=self.toggle(conn,user=self.data['editor'])
            with self.assertRaises(AccessError):estimate_tools.set_section_inclusion(conn,pid,self.data['owner'],body)
            count=conn.execute('SELECT COUNT(*) FROM project_activity').fetchone()[0]
            unchanged=self.toggle(conn)
            self.assertEqual(unchanged,after)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM project_activity').fetchone()[0],count)
            self.assertTrue(any(row['id']==self.data['editor']['id'] for row in project_contributors(conn,pid)))
            self.assertEqual(conn.execute("SELECT details FROM project_activity WHERE action='Excluded estimate section'").fetchone()[0],'Lighting')

    def test_copy_json_and_revisions_preserve_inclusion_and_legacy_defaults(self):
        with connect() as conn:
            pid, original=self.prepare(conn)
            rid_before=productivity.save_revision(conn,pid,self.data['owner'],{'label':'Before exclusion'})
            excluded=self.toggle(conn)
            rid=productivity.save_revision(conn,pid,self.data['owner'],{'label':'Excluded lighting'})
            copied=estimate_tools.copy_project(conn,pid,self.data['owner'],{'name':'Section copy'})
            self.assertEqual(projects.calculate_project(conn,copied)['totals'],excluded['totals'])
            exported=project_json.export_json(conn,pid,self.data['owner'])
            imported=project_json.import_json(conn,exported,self.data['owner'])
            self.assertEqual(projects.calculate_project(conn,imported)['totals'],excluded['totals'])
            invalid=json.loads(json.dumps(exported));invalid['project']['excluded_categories']='["Unknown"]'
            count=conn.execute('SELECT COUNT(*) FROM projects').fetchone()[0]
            with self.assertRaises(ValueError):project_json.import_json(conn,invalid,self.data['owner'])
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM projects').fetchone()[0],count)
            delta=productivity.compare_snapshots(original,excluded)
            self.assertEqual(delta['settings']['excluded_categories'],{'before':[],'after':['Lighting']})
            revision=conn.execute('SELECT revision FROM projects WHERE id=?',(pid,)).fetchone()[0]
            productivity.redeploy_revision(conn,pid,rid_before,self.data['owner'],{'expected_revision':revision})
            self.assertEqual(projects.calculate_project(conn,pid)['totals'],original['totals'])
            revision=conn.execute('SELECT revision FROM projects WHERE id=?',(pid,)).fetchone()[0]
            productivity.redeploy_revision(conn,pid,rid,self.data['owner'],{'expected_revision':revision})
            self.assertEqual(projects.calculate_project(conn,pid)['totals'],excluded['totals'])
            # A snapshot predating the feature restores the original all-on behavior.
            saved=json.loads(conn.execute('SELECT snapshot_json FROM tender_revisions WHERE id=?',(rid_before,)).fetchone()[0])
            saved['project'].pop('excluded_categories')
            conn.execute('UPDATE tender_revisions SET snapshot_json=? WHERE id=?',(json.dumps(saved),rid_before))
            revision=conn.execute('SELECT revision FROM projects WHERE id=?',(pid,)).fetchone()[0]
            productivity.redeploy_revision(conn,pid,rid_before,self.data['owner'],{'expected_revision':revision})
            self.assertEqual(projects.calculate_project(conn,pid)['totals'],original['totals'])

    def test_removed_or_renamed_section_does_not_disable_a_later_new_section(self):
        with connect() as conn:
            pid,_=self.prepare(conn)
            self.toggle(conn)
            projects.update_project_line(conn,pid,self.data['line_id'],{'category_snapshot':'General'})
            self.assertEqual(excluded_categories(projects.calculate_project(conn,pid)['project']['excluded_categories']),set())
            line=projects.add_catalog_line(conn,pid,self.data['catalog_id'],1)['lines']
            self.assertTrue(next(row for row in line if row['category_snapshot']=='Lighting')['included_in_total'])
            self.toggle(conn,'Lighting')
            project=projects.calculate_project(conn,pid)['project']
            estimate_tools.delete_section(conn,pid,self.data['owner'],{'category':'Lighting','expected_revision':project['revision']})
            self.assertEqual(excluded_categories(projects.calculate_project(conn,pid)['project']['excluded_categories']),set())

    def test_imported_category_whitespace_uses_one_key_without_rewriting_rows(self):
        variants = ['Cable  Tray', '  Cable Tray  ', 'Cable\tTray', 'Cable\u00a0Tray']
        with connect() as conn:
            portable = project_json.export_json(conn, self.data['project_id'], self.data['owner'])
            sample = portable['lines'][0]
            portable['lines'] = []
            for index, category in enumerate(variants + ['Power']):
                line = copy.deepcopy(sample)
                line.update(category_snapshot=category, name_snapshot=f'Imported fixture {index}')
                portable['lines'].append(line)
            portable['project']['scope_order'] = json.dumps(variants + ['Power'])
            pid = project_json.import_json(conn, portable, self.data['owner'])
            before = projects.calculate_project(conn, pid)
            self.assertEqual(json.loads(before['project']['scope_order']), ['Cable Tray', 'Power'])
            self.assertEqual([row['category'] for row in before['scope_summary']], ['Cable Tray', 'Power'])
            self.assertEqual([row['category_key'] for row in before['lines']], ['Cable Tray']*4 + ['Power'])
            ids = [line['id'] for line in before['lines']]
            estimate_tools.reorder_lines(conn, pid, {'ids': [ids[-1], *ids[:-1]],
                'scope_order': ['Power', ' Cable  Tray '], 'expected_revision': before['project']['revision']})
            before = projects.calculate_project(conn, pid)
            self.assertEqual([row['category'] for row in before['scope_summary']], ['Power', 'Cable Tray'])
            original_rows = [tuple(r) for r in conn.execute('SELECT * FROM project_lines WHERE project_id=? ORDER BY id', (pid,))]
            original_parts = [tuple(r) for r in conn.execute('SELECT c.* FROM project_line_components c JOIN project_lines l ON l.id=c.project_line_id WHERE l.project_id=? ORDER BY c.id', (pid,))]
            off = self.toggle(conn, 'Cable  Tray', pid=pid)
            self.assertEqual(excluded_categories(off['project']['excluded_categories']), {'Cable Tray'})
            cable = [row for row in off['lines'] if row['category_key']=='Cable Tray']
            self.assertEqual(len(cable), 4)
            self.assertTrue(all(row['included_in_total'] is False for row in cable))
            self.assertEqual(sorted(row['category_snapshot'] for row in cable), sorted(variants))
            self.assertAlmostEqual(off['totals']['quoted'], before['lines'][0]['quoted_total'])
            self.assertEqual(projects.project_totals(conn, pid), off['totals'])
            self.assertEqual(original_rows, [tuple(r) for r in conn.execute('SELECT * FROM project_lines WHERE project_id=? ORDER BY id', (pid,))])
            self.assertEqual(original_parts, [tuple(r) for r in conn.execute('SELECT c.* FROM project_line_components c JOIN project_lines l ON l.id=c.project_line_id WHERE l.project_id=? ORDER BY c.id', (pid,))])
            exported = project_json.export_json(conn, pid, self.data['owner'])
            self.assertEqual(sorted(line['category_snapshot'] for line in exported['lines']), sorted(variants+['Power']))
            self.assertEqual(json.loads(exported['project']['excluded_categories']), ['Cable Tray'])
            copied = estimate_tools.copy_project(conn, pid, self.data['owner'], {'name': 'Whitespace copy'})
            imported = project_json.import_json(conn, exported, self.data['owner'])
            for other in (copied, imported):
                data = projects.calculate_project(conn, other)
                self.assertEqual(data['totals'], off['totals'])
                self.assertEqual(sorted(line['category_snapshot'] for line in data['lines']), sorted(variants+['Power']))
            path = export_project(conn, pid, Path(self.tmp.name)/'whitespace.xlsx')
            # Deleting a displayed canonical section removes all its raw variants.
            revision = projects.calculate_project(conn, copied)['project']['revision']
            estimate_tools.delete_section(conn, copied, self.data['owner'],
                {'category': ' Cable\tTray ', 'expected_revision': revision})
            remaining = projects.calculate_project(conn, copied)
            self.assertEqual([line['category_snapshot'] for line in remaining['lines']], ['Power'])
            self.assertEqual(excluded_categories(remaining['project']['excluded_categories']), set())
            self.assertEqual(conn.execute('PRAGMA foreign_key_check').fetchall(), [])
        with connect() as conn:
            self.assertEqual(projects.calculate_project(conn, pid)['totals'], off['totals'])
            self.assertEqual(self.toggle(conn, 'Power', pid=pid)['totals']['quoted'], 0)
            self.toggle(conn, 'Cable\u00a0Tray', included=True, pid=pid)
            self.assertEqual(self.toggle(conn, 'Power', included=True, pid=pid)['totals'], before['totals'])
        book = load_workbook(path)
        try:
            detail = book['CONNECTED COST DETAIL']
            rows = [row for row in range(1,detail.max_row+1) if detail.cell(row,1).value=='Cable Tray']
            self.assertEqual(len(rows), 4)
            self.assertTrue(all(detail.cell(row,16).value==0 for row in rows))
            for parent in rows:
                self.assertTrue(any(isinstance(cell.value,str) and f'IF(P{parent}=0," [Optional cost]","")' in cell.value
                                    for row in detail for cell in row[:1]))
        finally:
            book.close()

    def test_noncanonical_exclusion_import_and_pruning_keep_remaining_alias_off(self):
        with connect() as conn:
            portable = project_json.export_json(conn, self.data['project_id'], self.data['owner'])
            portable['lines'][0]['category_snapshot'] = 'Cable  Tray'
            portable['lines'][1]['category_snapshot'] = ' Cable\tTray '
            portable['project']['excluded_categories'] = json.dumps([' Cable\u00a0 Tray '])
            pid = project_json.import_json(conn, portable, self.data['owner'])
            data = projects.calculate_project(conn, pid)
            self.assertEqual(data['totals']['quoted'], 0)
            self.assertEqual(json.loads(data['project']['excluded_categories']), ['Cable Tray'])
            projects.delete_project_line(conn, pid, data['lines'][0]['id'])
            data = projects.calculate_project(conn, pid)
            self.assertEqual(data['totals']['quoted'], 0)
            self.assertEqual(excluded_categories(data['project']['excluded_categories']), {'Cable Tray'})
            projects.update_project_line(conn, pid, data['lines'][0]['id'], {'category_snapshot': 'Power'})
            data = projects.calculate_project(conn, pid)
            self.assertEqual(excluded_categories(data['project']['excluded_categories']), set())
            self.assertGreater(data['totals']['quoted'], 0)

    def test_shared_estimate_identifies_excluded_raw_costs_and_included_total(self):
        with connect() as conn:
            pid, before = self.prepare(conn)
            off = self.toggle(conn)
            html = project_sharing.shared_html(off)
            self.assertIn('<h2>Included estimate total</h2>', html)
            self.assertEqual(html.count('Optional cost</'), 2)  # Category and its one item.
            self.assertIn('<details class="category section-excluded">', html)
            self.assertIn('<details class="estimate-item section-excluded">', html)
            self.assertIn(f'{before["lines"][0]["quoted_total"]:,.2f}', html)
            self.assertIn(f'{off["totals"]["quoted"]:,.2f}', html)
            all_on = project_sharing.shared_html(self.toggle(conn, included=True))
            self.assertNotIn('Optional cost', all_on)
            self.assertNotIn('class="category section-excluded"', all_on)

    def test_complete_schema31_upgrades_without_historical_backfills_and_once(self):
        with connect() as conn:
            pid,before=self.prepare(conn)
            drop_fixture_columns(conn,'projects',{'excluded_categories'})
            conn.execute('UPDATE schema_version SET version=31')
            conn.execute('DELETE FROM app_migrations WHERE rowid IN (SELECT rowid FROM app_migrations WHERE '+
                         conn.execute('PRAGMA table_info(app_migrations)').fetchone()[1]+'=?)',(auto_upgrade.MIGRATION_KEY,))
        auto_upgrade._ready.clear()
        with patch('conest_app.database._initialize_schema',side_effect=AssertionError('Historical migrations replayed')):
            with connect() as conn:
                after=projects.calculate_project(conn,pid)
                self.assertEqual(after['totals'],before['totals'])
                self.assertEqual(conn.execute('SELECT version FROM schema_version').fetchone()[0],auto_upgrade.CURRENT_SCHEMA)
                self.assertEqual(conn.execute('PRAGMA foreign_key_check').fetchall(),[])
        backup=database_path().parent/'upgrade_backups'/f'{database_path().stem}-before-schema-{auto_upgrade.CURRENT_SCHEMA}.db'
        self.assertTrue(backup.is_file())
        with _raw_connect(backup) as saved:
            self.assertEqual(saved.execute('SELECT version FROM schema_version').fetchone()[0],31)
            self.assertEqual(saved.execute('SELECT name FROM projects WHERE id=?',(pid,)).fetchone()[0],before['project']['name'])
        stamp=backup.stat().st_mtime_ns
        auto_upgrade._ready.clear()
        with patch('conest_app.auto_upgrade._safety_backup',side_effect=AssertionError('Unnecessary second backup')):
            with connect() as conn:self.assertEqual(projects.calculate_project(conn,pid)['totals'],before['totals'])
        self.assertEqual(backup.stat().st_mtime_ns,stamp)
        self.assertTrue(auto_upgrade._status(database_path())[0])

    def test_invalid_existing_column_and_corrupt_settings_are_not_silently_accepted(self):
        with connect() as conn:
            pid,_=self.prepare(conn)
            conn.execute('UPDATE projects SET excluded_categories=? WHERE id=?',('{"bad":true}',pid))
            with self.assertRaises(ValueError):projects.calculate_project(conn,pid)
            conn.execute("UPDATE projects SET excluded_categories='[]' WHERE id=?",(pid,))
            drop_fixture_columns(conn,'projects',{'excluded_categories'})
            conn.execute('ALTER TABLE projects ADD COLUMN excluded_categories INTEGER DEFAULT 0')
            conn.execute('UPDATE schema_version SET version=31')
            first=conn.execute('PRAGMA table_info(app_migrations)').fetchone()[1]
            conn.execute(f'DELETE FROM app_migrations WHERE "{first}"=?',(auto_upgrade.MIGRATION_KEY,))
        auto_upgrade._ready.clear()
        with self.assertRaises(auto_upgrade.UpgradeError):auto_upgrade.ensure_database()
        with _raw_connect(database_path()) as conn:
            self.assertEqual(conn.execute('SELECT version FROM schema_version').fetchone()[0],31)
            self.assertEqual(conn.execute(f'SELECT COUNT(*) FROM app_migrations WHERE "{first}"=?',(auto_upgrade.MIGRATION_KEY,)).fetchone()[0],0)

    def test_export_retains_raw_costs_but_totals_follow_editable_inclusion(self):
        with connect() as conn:
            pid,before=self.prepare(conn)
            after=self.toggle(conn)
            path=export_project(conn,pid,Path(self.tmp.name)/'sections.xlsx')
        book=load_workbook(path)
        detail=book['CONNECTED COST DETAIL']
        parent=next(row for row in range(1,detail.max_row+1) if detail.cell(row,1).value=='Lighting')
        self.assertEqual(detail.cell(parent,16).value,0)
        self.assertEqual(detail['P5'].value,'Cost treatment')
        self.assertIn('[Optional cost]',str(next(cell.value for row in detail for cell in row[:1] if isinstance(cell.value,str) and '[Optional cost]' in cell.value)))

        def evaluate(sheet,cell):
            raw=book[sheet][cell].value
            if not isinstance(raw,str) or not raw.startswith('='):return raw or 0
            expr=raw[1:].replace('$','')
            def rng(match):
                a,b,x,y=range_boundaries(match[1])
                return str(sum(evaluate(sheet,f'{get_column_letter(i)}{j}') for i in range(a,x+1) for j in range(b,y+1)))
            expr=re.sub(r'SUM\(([A-Z]+\d+:[A-Z]+\d+)\)',rng,expr)
            expr=re.sub(r"'([^']+)'!([A-Z]+\d+)",lambda m:str(evaluate(m[1],m[2])),expr)
            expr=re.sub(r'\b([A-Z]+\d+)\b',lambda m:str(evaluate(sheet,m[1])),expr)
            expr=re.sub(r'SUM\(([^()]*)\)',lambda m:'('+m[1].replace(',','+')+')',expr)
            tree=ast.parse(expr,mode='eval')
            self.assertTrue(all(isinstance(node,(ast.Expression,ast.Constant,ast.BinOp,ast.UnaryOp,ast.Add,ast.Sub,ast.Mult,ast.Div,ast.USub,ast.UAdd)) for node in ast.walk(tree)))
            return eval(compile(tree,'formula','eval'),{'__builtins__':{}})
        for cell,key in [('M5','material'),('M6','labor'),('M7','hours'),('M8','direct'),('M9','quoted')]:
            self.assertAlmostEqual(evaluate('COST ESTIMATE',cell),after['totals'][key])
        self.assertAlmostEqual(evaluate('CONNECTED COST DETAIL',f'L{parent}'),before['lines'][0]['quoted_total'])
        detail.cell(parent,16,1)
        self.assertAlmostEqual(evaluate('COST ESTIMATE','M9'),before['totals']['quoted'])
        book.close()


if __name__ == '__main__':
    unittest.main()
