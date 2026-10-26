import unittest,json
from pathlib import Path
import test_estimate_sections as helpers
from conest_app import projects,project_sharing,project_json,estimate_tools,productivity
from conest_app.estimate_areas import allocations
from conest_app.database import connect
from conest_app.exporter import export_project
from openpyxl import load_workbook
class AreaTests(unittest.TestCase):
    setUp=helpers.EstimateSectionTests.setUp
    tearDown=helpers.EstimateSectionTests.tearDown
    prepare=helpers.EstimateSectionTests.prepare
    toggle=helpers.EstimateSectionTests.toggle
    def test_allocation_totals_optional_fixed_and_roundtrip(self):
        with connect() as conn:
            pid,before=self.prepare(conn);lid=self.data['line_id']
            conn.execute('UPDATE project_line_components SET fixed_quantity=1 WHERE project_line_id=?',(lid,))
            data=projects.update_project_line(conn,pid,lid,{'area_allocations':[{'area':'Level 1','quantity':3},{'area':'Level 2','quantity':7}]})
            self.assertEqual(next(l for l in data['lines'] if l['id']==lid)['quantity'],10)
            for k in ('material','hours','labor','quoted'):self.assertAlmostEqual(sum(a['included'][k] for a in data['area_summary']),data['totals'][k])
            a=next(a for a in data['area_summary'] if a['area']=='Level 1');b=next(a for a in data['area_summary'] if a['area']=='Level 2')
            self.assertAlmostEqual(a['included']['quoted']/3,b['included']['quoted']/7)
            with self.assertRaises(ValueError):projects.update_project_line(conn,pid,lid,{'quantity':11})
            data=self.toggle(conn)
            for k in ('material','hours','labor','quoted'):self.assertAlmostEqual(sum(a['included'][k] for a in data['area_summary']),data['totals'][k])
            self.assertGreater(sum(a['optional']['quoted'] for a in data['area_summary']),0)
            html=project_sharing.shared_html(data);self.assertIn('Level 1',html);self.assertIn('Optional cost',html)
            raw=project_json.export_json(conn,pid,self.data['owner']);self.assertIn('area_allocations',raw['lines'][0])
            result=project_json.import_json(conn,raw,self.data['owner'])
            path=export_project(conn,pid,Path(self.tmp.name)/'areas.xlsx')
        wb=load_workbook(path);ws=wb['AREA ALLOCATIONS'];detail=wb['CONNECTED COST DETAIL']
        self.assertEqual(ws['A2'].value,'Level 1');self.assertEqual(ws['D2'].value,3)
        self.assertTrue(any(isinstance(c.value,str) and "SUM('AREA ALLOCATIONS'!" in c.value for row in detail for c in row))
        self.assertIn('SUMPRODUCT',wb['AREA SUMMARY']['E2'].value)
        self.assertIn('!P',ws['I2'].value)
        for sheet in wb:
            for row in sheet:
                for c in row:
                    if c.data_type=='f':self.assertNotIn('#REF!',c.value)
        wb.close()
    def test_invalid_and_clear(self):
        for value in ([{'area':'A','quantity':float('nan')}],[{'area':'A','quantity':-1}],[{'area':'A','quantity':0}],[{'area':'A','quantity':1},{'area':'a','quantity':2}]):
            with self.assertRaises(ValueError):allocations(value)
        with connect() as conn:
            pid,_=self.prepare(conn);lid=self.data['line_id']
            projects.update_project_line(conn,pid,lid,{'area_allocations':[{'area':'Room','quantity':4}]})
            data=projects.update_project_line(conn,pid,lid,{'area_allocations':[]})
            self.assertEqual(next(l for l in data['lines'] if l['id']==lid)['quantity'],4)
            self.assertEqual(len(data['area_summary']),1)
