"""Optional sections remain priced separately in connected Excel and public views."""
import ast
import copy
import re
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

import test_estimate_sections as section_helpers
from conest_app import project_sharing, projects
from conest_app.database import connect
from conest_app.exporter import export_project


class OptionalCostExportTests(unittest.TestCase):
    setUp = section_helpers.EstimateSectionTests.setUp
    tearDown = section_helpers.EstimateSectionTests.tearDown
    prepare = section_helpers.EstimateSectionTests.prepare
    toggle = section_helpers.EstimateSectionTests.toggle

    def numeric_cell(self, book, sheet, cell):
        """Evaluate only the exporter arithmetic subset, independently of app totals."""
        raw = book[sheet][cell].value
        if not isinstance(raw, str) or not raw.startswith('='):
            return raw or 0
        expression = raw[1:].replace('$', '')

        def range_total(match):
            left, top, right, bottom = range_boundaries(match[1])
            return str(sum(self.numeric_cell(book, sheet, f'{get_column_letter(column)}{row}')
                           for column in range(left, right + 1) for row in range(top, bottom + 1)))

        expression = re.sub(r'SUM\(([A-Z]+\d+:[A-Z]+\d+)\)', range_total, expression)
        expression = re.sub(r"'([^']+)'!([A-Z]+\d+)",
                            lambda match: str(self.numeric_cell(book, match[1], match[2])), expression)
        expression = re.sub(r'\b([A-Z]+\d+)\b',
                            lambda match: str(self.numeric_cell(book, sheet, match[1])), expression)
        expression = re.sub(r'SUM\(([^()]*)\)',
                            lambda match: '(' + match[1].replace(',', '+') + ')', expression)
        tree = ast.parse(expression, mode='eval')
        allowed = (ast.Expression, ast.Constant, ast.BinOp, ast.UnaryOp,
                   ast.Add, ast.Sub, ast.Mult, ast.Div, ast.USub, ast.UAdd)
        self.assertTrue(all(isinstance(node, allowed) for node in ast.walk(tree)), expression)
        return eval(compile(tree, 'export formula', 'eval'), {'__builtins__': {}})

    def test_workbook_separates_optional_costs_and_recalculates_after_flag_and_price_edits(self):
        with connect() as conn:
            pid, _ = self.prepare(conn)
            # A per-item final mark-up exemption must stay part of both cost bases.
            projects.update_project_line(conn, pid, self.data['line_id'], {'exclude_final_markup': True})
            before = projects.calculate_project(conn, pid)
            after = self.toggle(conn)
            path = export_project(conn, pid, Path(self.tmp.name) / 'optional.xlsx')
        book = load_workbook(path)
        try:
            detail = book['CONNECTED COST DETAIL']
            summary = book['COST ESTIMATE']
            parent_rows = [row for row in range(1, detail.max_row + 1)
                           if detail.cell(row, 1).value in ('Lighting', 'Power')]
            included_row = next(row for row in range(1, detail.max_row + 1)
                                if detail.cell(row, 1).value == 'INCLUDED ESTIMATE TOTAL')
            optional_row = included_row + 1
            self.assertEqual(detail.cell(optional_row, 1).value, 'OPTIONAL COST TOTAL (excluded)')
            self.assertEqual(summary['J10'].value, 'OPTIONAL QUOTED COST')
            self.assertEqual(summary['M10'].value, f"='CONNECTED COST DETAIL'!L{optional_row}")
            self.assertEqual(detail['P5'].value, 'Cost treatment')
            self.assertIn(f'$P${optional_row}', str(detail.print_area))
            self.assertTrue(detail.row_dimensions[optional_row].height >= 30)
            self.assertTrue(detail.column_dimensions['P'].width >= 20)
            for row in parent_rows:
                self.assertEqual(detail.cell(row, 16).number_format,
                                 '[=1]"Included";[=0]"Optional cost";General')
                self.assertNotRegex(detail.cell(row, 16).comment.text, r'\b(?:On|Off)\b')
            self.assertTrue(all(not re.search(r'\b(?:On|Off)\b', rule.error or '')
                                for rule in detail.data_validations.dataValidation))
            self.assertAlmostEqual(self.numeric_cell(book, 'COST ESTIMATE', 'M9'), after['totals']['quoted'])
            self.assertAlmostEqual(self.numeric_cell(book, 'COST ESTIMATE', 'M10'), before['lines'][0]['quoted_total'])

            cases = ((0, 1), (0, 0), (1, 1), (1, 0))
            mappings = (('G', 'material', 'total_material'), ('I', 'hours', 'total_hours'),
                        ('J', 'labor', 'total_labor'), ('K', 'direct', 'direct_total'),
                        ('L', 'quoted', 'quoted_total'))
            for flags in cases:
                for row, flag in zip(parent_rows, flags):
                    detail.cell(row, 16, flag)
                for column, key, field in mappings:
                    expected_included = sum(line[field] for line, flag in zip(before['lines'], flags) if flag)
                    expected_optional = sum(line[field] for line, flag in zip(before['lines'], flags) if not flag)
                    included = self.numeric_cell(book, detail.title, f'{column}{included_row}')
                    optional = self.numeric_cell(book, detail.title, f'{column}{optional_row}')
                    self.assertAlmostEqual(included, expected_included, msg=f'{flags}: included {key}')
                    self.assertAlmostEqual(optional, expected_optional, msg=f'{flags}: optional {key}')
                    self.assertAlmostEqual(included + optional, before['totals'][key])

            # The optional amount follows the same component price inputs, rather
            # than being a snapshot written only at export time.
            first = parent_rows[0]
            detail.cell(first, 16, 0)
            detail.cell(parent_rows[1], 16, 1)
            included_before = self.numeric_cell(book, summary.title, 'M9')
            optional_before = self.numeric_cell(book, summary.title, 'M10')
            component_price = detail.cell(first + 1, 6)
            component_price.value += 25
            self.assertAlmostEqual(self.numeric_cell(book, summary.title, 'M9'), included_before)
            self.assertGreater(self.numeric_cell(book, summary.title, 'M10'), optional_before)
            self.assertAlmostEqual(self.numeric_cell(book, summary.title, 'M10'),
                                   self.numeric_cell(book, detail.title, f'L{first}'))
        finally:
            book.close()

    def test_shared_optional_summary_is_separate_and_preserves_item_values(self):
        with connect() as conn:
            pid, before = self.prepare(conn)
            after = self.toggle(conn)
            original = copy.deepcopy(after)
            html = project_sharing.shared_html(after)
            included = re.search(r'<section class="project-totals">(.*?)</section>', html).group(1)
            optional = re.search(r'<section class="optional-totals">(.*?)</section>', html).group(1)
            self.assertIn('<h2>Included estimate total</h2>', included)
            self.assertIn('<h2>Optional costs</h2>', optional)
            self.assertIn('These items are not included in the estimate total.', optional)
            self.assertIn(f'AUD {after["totals"]["quoted"]:,.2f}', included)
            self.assertIn(f'AUD {before["lines"][0]["quoted_total"]:,.2f}', optional)
            self.assertEqual(html.count('Optional cost</'), 2)
            self.assertNotIn('Excluded from total</', html)
            self.assertEqual(after, original)

            all_optional = self.toggle(conn, 'Power')
            html = project_sharing.shared_html(all_optional)
            included = re.search(r'<section class="project-totals">(.*?)</section>', html).group(1)
            optional = re.search(r'<section class="optional-totals">(.*?)</section>', html).group(1)
            self.assertIn('AUD 0.00', included)
            self.assertIn(f'AUD {before["totals"]["quoted"]:,.2f}', optional)
            self.assertEqual(html.count('Optional cost</'), 4)

            self.toggle(conn, included=True)
            all_included = self.toggle(conn, 'Power', included=True)
            html = project_sharing.shared_html(all_included)
            self.assertNotIn('optional-totals', html)
            self.assertNotIn('Optional cost</', html)
            self.assertIn(f'AUD {before["totals"]["quoted"]:,.2f}', html)

    def test_empty_estimate_has_zero_optional_cost_without_invalid_references(self):
        with connect() as conn:
            pid = projects.create_project(conn, {'name': 'Empty optional costs'})['project']['id']
            path = export_project(conn, pid, Path(self.tmp.name) / 'empty.xlsx')
            self.assertNotIn('optional-totals', project_sharing.shared_html(projects.calculate_project(conn, pid)))
        book = load_workbook(path)
        try:
            self.assertEqual(self.numeric_cell(book, 'COST ESTIMATE', 'M9'), 0)
            self.assertEqual(self.numeric_cell(book, 'COST ESTIMATE', 'M10'), 0)
        finally:
            book.close()


if __name__ == '__main__':
    unittest.main()
