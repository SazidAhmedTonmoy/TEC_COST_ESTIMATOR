"""Workbook identity, idempotency, atomicity, concurrency and export regression checks."""
import copy
import io
import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from conest_app.tenders import Store, ValidationError
from conest_app.tender_import import import_workbook, parse_workbook
from conest_app.tender_export import export_workbook, sheet_xml

ACTOR={'id':1,'name':'Import test'}

def rewrite(blob, replacements):
    output=io.BytesIO()
    with ZipFile(io.BytesIO(blob)) as source,ZipFile(output,'w',ZIP_DEFLATED) as target:
        for name in source.namelist():target.writestr(name,replacements.get(name,source.read(name)))
    return output.getvalue()

class TenderImportTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.store=Store(Path(self.temp.name)/'tenders.db');self.data=self.store.all_data()
    def run_import(self,blob,apply=False):
        result=import_workbook(self.store,'update.xlsx',blob,ACTOR)
        if apply:result=import_workbook(self.store,'update.xlsx',blob,ACTOR,result['token'])
        return result
    def export(self,data=None):
        data=data or self.data
        return export_workbook(data,[s['id'] for s in data['submissions']])
    def raw(self, rows):
        source=json.loads((Path(__file__).resolve().parents[1]/'tenders_seed/workbook-source.json').read_text())['Raw Data']
        xml=sheet_xml('Raw Data',source[0],rows,[25]*len(source[0]),'Original workbook','','')
        blob=self.export();z=ZipFile(io.BytesIO(blob))
        return rewrite(blob,{'xl/workbook.xml':z.read('xl/workbook.xml').replace(b'name="Tenders"',b'name="Raw Data"'),'xl/worksheets/sheet4.xml':xml})
    def test_roundtrip_original_export_and_repeat(self):
        result=self.run_import(self.export(),True)
        self.assertEqual(result['counts']['unchanged'],291)
        self.assertEqual(result['counts']['updated'],0)
        self.assertEqual(result['counts']['followupsUpdated'],0)
        again=self.run_import(self.export(),True)
        self.assertEqual(again['counts'],result['counts'])
        self.assertEqual(len(self.store.all_data()['chases']),332)
    def test_new_changed_duplicate_and_preserve_absent(self):
        d=copy.deepcopy(self.data);s=d['submissions'][0];sid=s['id'];s['value']+=10;s['status']='Won'
        new={**s,'id':'new-tender','project':'New tender project','builder':'New builder'}
        d['submissions']=[s,new,dict(new)];d['chases']=[];d['reminders']=[]
        blob=self.export(d);p=self.run_import(blob)
        self.assertEqual((p['counts']['added'],p['counts']['updated'],p['counts']['duplicates']),(1,sum(x['project']==s['project'] for x in self.data['submissions']),1))
        self.assertEqual(len(self.store.all_data()['submissions']),291)
        self.run_import(blob,True)
        current=self.store.all_data();self.assertEqual(len(current['submissions']),292)
        self.assertEqual(len(current['chases']),332)
        changed=next(s for s in current['submissions'] if s['id']==sid)
        self.assertEqual(changed['status'],'Won');self.assertEqual(changed['revision'],1)
        self.assertEqual(self.run_import(blob)['counts']['unchanged'],2)
    def test_raw_reordered_rows_and_website_edit_survives(self):
        raw=json.loads((Path(__file__).resolve().parents[1]/'tenders_seed/workbook-source.json').read_text())['Raw Data'][1:]
        blob=self.raw(list(reversed(raw)))
        result=self.run_import(blob,True)
        self.assertEqual(result['counts']['added'],0)
        self.assertEqual(result['counts']['updated'],0)
        sid=self.data['submissions'][0]['id'];item=next(s for s in self.store.all_data()['submissions'] if s['id']==sid)
        item['remarks']='New app note';self.store.save_submission(item)
        self.assertEqual(self.run_import(blob,True)['counts']['updated'],0)
        self.assertEqual(next(s for s in self.store.all_data()['submissions'] if s['id']==sid)['remarks'],'New app note')
        raw[0][0]+=100
        changed=self.run_import(self.raw(raw),True)
        self.assertEqual(changed['counts']['updated'],sum(x['project']==self.data['submissions'][0]['project'] for x in self.data['submissions']))
        self.assertEqual(len(self.store.all_data()['submissions']),291)
    def test_conflicting_duplicate_is_rejected_atomically(self):
        d=copy.deepcopy(self.data);a=d['submissions'][0];b={**a,'value':a['value']+1};d['submissions']=[a,b]
        before=self.store.all_data()
        with self.assertRaisesRegex(ValidationError,'conflicting duplicate'):self.run_import(self.export(d),True)
        self.assertEqual(self.store.all_data(),before)
    def test_concurrent_change_requires_new_preview(self):
        blob=self.export();preview=self.run_import(blob)
        item=copy.deepcopy(self.data['submissions'][0]);item['remarks']='Another user';self.store.save_submission(item)
        with self.assertRaisesRegex(ValidationError,'changed after preview'):
            import_workbook(self.store,'update.xlsx',blob,ACTOR,preview['token'])
    def test_history_updates_and_new_reminders_are_idempotent(self):
        d=copy.deepcopy(self.data);d['chases'][0]['note']='Updated historical note'
        d['reminders']=[{'id':'new-reminder','submissionId':d['submissions'][0]['id'],'due':'2026-10-10T09:00','note':'Call builder','done':False,'createdAt':'2026-10-05'}]
        blob=self.export(d);result=self.run_import(blob,True)
        self.assertEqual(result['counts']['followupsUpdated'],1);self.assertEqual(result['counts']['remindersAdded'],1)
        again=self.run_import(blob,True)
        self.assertEqual(again['counts']['followupsUpdated'],0);self.assertEqual(again['counts']['remindersAdded'],0)
    def test_bad_dates_and_negative_values(self):
        for value in (-1,float('nan')):
            d=copy.deepcopy(self.data);d['submissions'][0]['value']=value
            with self.assertRaises(ValidationError):self.run_import(self.export(d))
        with self.assertRaises(ValidationError):self.run_import(b'not an xlsx')
    def test_calculated_empty_string_is_valid_but_missing_or_error_cache_is_not(self):
        blob=self.export();ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
        with ZipFile(io.BytesIO(blob)) as z:original=z.read('xl/worksheets/sheet4.xml')
        for kind in ('empty-string','missing-number','missing-string','error'):
            root=ET.fromstring(original)
            c=next(c for c in root.iter('{'+ns+'}c') if c.get('r')=='O6')
            c.clear();c.set('r','O6')
            if kind!='missing-number':c.set('t','e' if kind=='error' else 'str')
            ET.SubElement(c,'{'+ns+'}f',{'t':'shared','si':'5'}).text='IF(1=1,"","Follow up")'
            if kind!='missing-string':
                value=ET.SubElement(c,'{'+ns+'}v')
                if kind=='error':value.text='#VALUE!'
            updated=rewrite(blob,{'xl/worksheets/sheet4.xml':ET.tostring(root)})
            if kind=='empty-string':
                parsed=parse_workbook('test.xlsx',updated)
                self.assertEqual(parsed['records'][0]['latest'],'')
                self.assertEqual(len(parsed['records']),291)
            else:
                with self.assertRaisesRegex(ValidationError,'formula has no valid saved result'):
                    parse_workbook('test.xlsx',updated)
    def test_visuals_formula_cache_empty_export_and_literal_text(self):
        d=copy.deepcopy(self.data);d['submissions'][0]['remarks']='=HYPERLINK("bad")'
        blob=self.export(d);w=load_workbook(io.BytesIO(blob),data_only=False)
        self.assertEqual(len(w['Summary']._charts),3)
        self.assertEqual(w['Tenders'].freeze_panes,'D6');self.assertTrue(w['Tenders'].column_dimensions['A'].hidden)
        notes=[c for row in w['Tenders'].iter_rows(min_row=6,min_col=16,max_col=16) for c in row]
        note=next(c for c in notes if c.value=='=HYPERLINK("bad")');self.assertEqual(note.data_type,'s')
        for zname in ZipFile(io.BytesIO(blob)).namelist():
            if zname.endswith(('.xml','.rels')):ET.fromstring(ZipFile(io.BytesIO(blob)).read(zname))
        cached=load_workbook(io.BytesIO(blob),data_only=True)
        self.assertEqual(cached['Summary']['A6'].value,291);self.assertEqual(cached['Summary']['E6'].value,106)
        self.assertEqual(cached['Summary']['I6'].value,121)
        empty=load_workbook(io.BytesIO(export_workbook(d,[])),data_only=True)
        self.assertEqual(empty['Summary']['A6'].value,0);self.assertEqual(len(empty['Summary']._charts),0)

if __name__=='__main__':unittest.main()
