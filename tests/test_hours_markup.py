import unittest,tempfile,os,math,re,ast
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries,get_column_letter
from conest_app.database import initialize,connect
from conest_app.projects import create_project,update_project,calculate_project
from conest_app.exporter import export_project
from conest_app.estimate_tools import copy_project
from fixtures_v232 import seed

class HoursMarkupTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.env=patch.dict(os.environ,{'CONEST_DB_PATH':self.tmp.name+'/test.db','CONEST_LIGHT_PRICES_DB':self.tmp.name+'/none.db'});self.env.start();initialize()
  with connect() as c:self.data=seed(c)
 def tearDown(self):self.env.stop();self.tmp.cleanup()
 def test_markup_defaults_costs_copy_and_excel(self):
  with connect() as c:
   pid=self.data['project_id'];original=calculate_project(c,pid)
   self.assertEqual(original['project']['labor_hours_markup'],0)
   marked=update_project(c,pid,{'labor_hours_markup':.1})
   fixed=sum(x['total_labor'] for line in original['lines'] for x in line['components'] if x['labor_pricing']=='fixed')
   self.assertAlmostEqual(marked['totals']['hours'],original['totals']['hours']*1.1)
   self.assertAlmostEqual(marked['totals']['labor'],(original['totals']['labor']-fixed)*1.1+fixed)
   self.assertEqual(marked['totals']['material'],original['totals']['material'])
   self.assertAlmostEqual(marked['crew_plan']['working_days'],original['crew_plan']['working_days']*1.1)
   copied=copy_project(c,pid,self.data['owner'],{'name':'Copy'})
   self.assertEqual(calculate_project(c,copied)['totals'],marked['totals'])
   for bad in [-1,float('inf'),'bad']:
    with self.assertRaises(ValueError):update_project(c,pid,{'labor_hours_markup':bad})
   c.execute("UPDATE settings SET value='.2' WHERE key='default_labor_hours_markup'")
   self.assertEqual(create_project(c,{'name':'Default'})['project']['labor_hours_markup'],.2)
   file=export_project(c,pid,Path(self.tmp.name)/'export.xlsx')
  w=load_workbook(file);self.assertEqual(w['COST ESTIMATE']['F7'].value,.1)
  # Evaluate the exported arithmetic chain independently of application totals.
  def value(sheet,cell):
   raw=w[sheet][cell].value
   if not isinstance(raw,str) or not raw.startswith('='):return raw or 0
   expr=raw[1:].replace('$','')
   def rng(m):
    a,b,x,y=range_boundaries(m[1]);return str(sum(value(sheet,f'{get_column_letter(i)}{j}') for i in range(a,x+1) for j in range(b,y+1)))
   expr=re.sub(r'SUM\(([A-Z]+\d+:[A-Z]+\d+)\)',rng,expr)
   expr=re.sub(r"'([^']+)'!([A-Z]+\d+)",lambda m:str(value(m[1],m[2])),expr)
   expr=re.sub(r'\b([A-Z]+\d+)\b',lambda m:str(value(sheet,m[1])),expr)
   expr=re.sub(r'SUM\(([^()]*)\)',lambda m:'('+m[1].replace(',','+')+')',expr)
   tree=ast.parse(expr,mode='eval')
   self.assertTrue(all(isinstance(n,(ast.Expression,ast.Constant,ast.BinOp,ast.UnaryOp,ast.Add,ast.Sub,ast.Mult,ast.Div,ast.USub,ast.UAdd)) for n in ast.walk(tree)),expr)
   return eval(compile(tree,'formula','eval'),{'__builtins__':{}})
  for markup in [0,.1,.25]:
   w['COST ESTIMATE']['F7']=markup
   with connect() as c:expected=update_project(c,pid,{'labor_hours_markup':markup})['totals']
   for cell,key in [('M5','material'),('M6','labor'),('M7','hours'),('M8','direct'),('M9','quoted')]:
    self.assertAlmostEqual(value('COST ESTIMATE',cell),expected[key],places=6)
  w.close()

 def test_recent_matches_before_limit_and_after_new_usage(self):
  from conest_app.catalog import search_catalog,create_manual_catalog_item
  from conest_app.projects import add_catalog_line
  with connect() as c:
   a=create_manual_catalog_item(c,{'name':'Recent Test A','components':[{'description':'A','quantity_per_unit':1,'unit_material':1}]})
   b=create_manual_catalog_item(c,{'name':'Recent Test B','components':[{'description':'B','quantity_per_unit':1,'unit_material':2}]})
   pid=self.data['project_id'];add_catalog_line(c,pid,b['id'],1)
  with connect() as c:
   self.assertEqual(search_catalog(c,'Recent Test',limit=1,recent_first=True)[0]['id'],b['id'])
   self.assertFalse(search_catalog(c,'not matching any fixture',recent_first=True))
   add_catalog_line(c,pid,a['id'],1)
   c.execute("UPDATE project_lines SET created_at='2099-01-01T00:00:00' WHERE project_id=? AND catalog_item_id=?",(pid,a['id']))
  with connect() as c:
   self.assertEqual(search_catalog(c,'Recent Test',limit=1,recent_first=True)[0]['id'],a['id'])
