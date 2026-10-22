"""Active-workspace behavior after a reset that retains legacy FK parents."""
import io
import os
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
import uuid
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from conest_app import opportunities as op, opportunity_http, opportunity_reset as reset, tenders
from conest_app.database import connect
from conest_app.opportunity_import import export_workbook, import_workbook, parse_workbook


def workbook(records):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Opportunities'
    headers = ['Opportunity ID', 'Project Name', 'Region', 'Closing Date',
               'Builders', 'Working Scope', 'Historical tag']
    sheet.append(headers)
    for record in records:
        sheet.append([record.get(header, '') for header in headers])
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class OpportunityRetentionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        environment = patch.dict(os.environ, {
            'CONEST_DB_PATH': str(root / 'catalog.db'),
            'CONEST_TENDERS_DB_PATH': str(root / 'tenders.db'),
        })
        environment.start()
        self.addCleanup(environment.stop)
        with connect() as db:
            self.actor = dict(db.execute("SELECT id,name,role,team_id FROM users WHERE role='admin' LIMIT 1").fetchone())
            stamp = tenders.timestamp()
            self.team = db.execute('INSERT INTO teams(name,created_at,updated_at) VALUES(?,?,?)',
                                   ('Retention team', stamp, stamp)).lastrowid
            self.other_team = db.execute('INSERT INTO teams(name,created_at,updated_at) VALUES(?,?,?)',
                                         ('Other retention team', stamp, stamp)).lastrowid
            self.member = self.make_user(db, 'Retention member', self.team)
            self.other_member = self.make_user(db, 'Other retention member', self.other_team)
        # Prepare the old register before the constructor's one-time reset hook.
        with patch.object(reset, 'reset_once', return_value={'applied': False, 'removed': 0}):
            self.store = tenders.Store(root / 'tenders.db').keep_open()
        self.addCleanup(self.store.close)

    @staticmethod
    def make_user(db, name, team):
        stamp = tenders.timestamp()
        identifier = db.execute('''INSERT INTO users
            (username,name,password_hash,role,team_id,is_active,must_change_password,created_at,updated_at)
            VALUES(?,?,?,'user',?,1,0,?,?)''',
            (name, name, 'not-a-login-hash', team, stamp, stamp)).lastrowid
        return {'id': identifier, 'name': name, 'role': 'user', 'team_id': team}

    def create(self, project='Old opportunity', **changes):
        return op.save(self.store, {
            'project': project, 'region': 'NSW', 'due': '2099-01-11',
            'builders': ['Old opportunity-only builder'], 'workingScope': 'Retained electrical scope',
            **changes,
        }, self.actor)['opportunity']

    def row(self, identifier):
        with self.store.connect() as db:
            return tuple(db.execute('SELECT * FROM opportunities WHERE id=?', (identifier,)).fetchone())

    def assert_missing(self, operation):
        with self.assertRaises(tenders.ValidationError) as caught:
            operation()
        self.assertEqual(caught.exception.status, 404, str(caught.exception))

    def assert_empty_export(self, content):
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            rows = list(book['Opportunities'].values)
            header = next(index for index, row in enumerate(rows) if 'Project Name' in row)
            self.assertFalse(any(value not in (None, '') for row in rows[header + 1:] for value in row))
        finally:
            book.close()

    def notice(self, identifier, target_kind, target_id, actor=None):
        actor = actor or self.actor
        with self.store.connect() as db:
            db.execute('''INSERT INTO tender_notifications
                (id,user_id,kind,target_kind,target_id,title,message,created_at,event_key)
                VALUES(?,?,?,?,?,?,?,?,?)''',
                (identifier, actor['id'], 'reminder' if target_kind == 'submission' else 'updated',
                 target_kind, target_id, identifier, 'Keep the original notification row',
                 tenders.timestamp(), 'retention-test:' + identifier))
        return identifier

    def test_reset_hides_every_register_view_and_export_without_removing_rows(self):
        pending = self.create(assignedTeamIds=[self.team], assignedUserIds=[self.member['id']])
        approved = self.create('Old approved opportunity', region='VIC')
        approved = op.decide(self.store, {**approved, 'decision': 'approved'}, self.actor)['opportunity']
        rejected = self.create('Old rejected opportunity')
        rejected = op.decide(self.store, {**rejected, 'decision': 'rejected', 'reason': 'Historical'}, self.actor)['opportunity']
        expired = self.create('Old expired opportunity', due='2001-01-01')
        old = [pending, approved, rejected, expired]
        before = {item['id']: self.row(item['id']) for item in old}
        self.assertEqual(op.list_opportunities(self.store, {}, self.actor)['counts']['all'], 4)
        self.assertEqual(reset.reset_once(self.store)['removed'], 4)

        filters = [{}, *({'status': status} for status in op.STATUSES), {'region': 'VIC'},
                   {'teamId': str(self.team)}, {'teamId': 'unassigned'},
                   {'userId': str(self.member['id'])}, {'q': 'Retained electrical scope'}]
        for selected in filters:
            with self.subTest(filters=selected):
                result = op.list_opportunities(self.store, selected, self.actor)
                self.assertEqual(result['items'], [])
                self.assertEqual(result['total'], 0)
                self.assertTrue(all(count == 0 for count in result['counts'].values()))
                self.assertTrue(all(team['count'] == 0 for team in result['teamCounts']))
                self.assertEqual(result['unassignedCount'], 0)
        self.assertEqual(op.existing_keys(self.store), {})
        for identifiers in (None, [pending['id']]):
            self.assert_empty_export(export_workbook(self.store, ids=identifiers))
        response = []
        path = '/api/tenders/opportunities/export.xlsx'
        handler = SimpleNamespace(path=path + '?teamId=' + str(self.team), command='GET', user=self.actor)
        self.assertTrue(opportunity_http.route(handler, path, self.store,
                        lambda _handler, body, *_args, **_kwargs: response.append(body)))
        self.assert_empty_export(response[0])
        self.assertEqual({item['id']: self.row(item['id']) for item in old}, before)
        self.assertEqual({r['code'] for r in op.regions(self.store, self.actor)['regions']}, {'NSW', 'VIC'})

    def test_retired_ids_cannot_be_opened_or_changed_even_by_creator_admin(self):
        old = self.create(assignedTeamIds=[self.team])
        before = self.row(old['id'])
        with self.store.connect() as db:
            history = [tuple(row) for row in db.execute('SELECT * FROM opportunity_activity ORDER BY id')]
        reset.reset_once(self.store)
        operations = {
            'detail': lambda: op.detail(self.store, old['id'], self.actor),
            'activity': lambda: op.activity_history(self.store, {'id': old['id']}, self.actor),
            'edit': lambda: op.save(self.store, {**old, 'remarks': 'Must not change history'}, self.actor),
            'assign': lambda: op.update_authority(self.store, {**old, 'assignedTeamIds': []}, self.actor, 'assign'),
            'access': lambda: op.update_authority(self.store, {**old, 'editorUserIds': [self.member['id']]}, self.actor, 'access'),
            'approve': lambda: op.decide(self.store, {**old, 'decision': 'approved'}, self.actor),
            'delete': lambda: op.delete(self.store, old, self.actor),
        }
        for name, operation in operations.items():
            with self.subTest(operation=name):
                self.assert_missing(operation)
        self.assertEqual(self.row(old['id']), before)
        with self.store.connect() as db:
            self.assertEqual([tuple(row) for row in db.execute('SELECT * FROM opportunity_activity ORDER BY id')], history)

    def test_notifications_hide_reset_ids_and_late_parent_notices_but_keep_normal_deletions(self):
        old = self.create()
        submission = self.store.all_data()['submissions'][0]
        self.notice('retained-reminder', 'submission', submission['id'])
        orphan = self.notice('old-orphan-notification', 'opportunity', 'unknown-historical-opportunity')
        before = op.notifications(self.store, {'pageSize': 100}, self.actor)
        reminder_ids = {row['id'] for row in before['items'] if 'submissionId' in row}
        reset.reset_once(self.store)
        late = self.notice('late-old-worker-notification', 'opportunity', old['id'])
        fresh = self.create('New opportunity after the reset')
        op.delete(self.store, fresh, self.actor)
        with self.store.connect() as db:
            new_ids = {row[0] for row in db.execute(
                'SELECT id FROM tender_notifications WHERE user_id=? AND target_id=?',
                (self.actor['id'], fresh['id']))}
        expected = reminder_ids | new_ids
        result = op.notifications(self.store, {'pageSize': 100}, self.actor)
        self.assertEqual({row['id'] for row in result['items']}, expected)
        self.assertEqual(result['total'], len(expected))
        self.assertEqual(result['unreadCount'], sum(not row['read'] for row in result['items']))
        self.assertTrue(any(row['kind'] == 'deleted' and row.get('opportunityId') == fresh['id'] for row in result['items']))
        paged = [op.notifications(self.store, {'pageSize': 1, 'page': number}, self.actor)
                 for number in range(1, len(expected) + 1)]
        self.assertEqual({page['items'][0]['id'] for page in paged}, expected)
        self.assertTrue(all(page['total'] == len(expected) for page in paged))
        with self.store.connect() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM tender_notifications WHERE id IN (?,?)', (orphan, late)).fetchone()[0], 2)

    def test_read_actions_only_change_visible_notifications_owned_by_the_reader(self):
        old = self.create()
        with self.store.connect() as db:
            old_ids = {row[0] for row in db.execute(
                'SELECT id FROM tender_notifications WHERE user_id=? AND target_id=?',
                (self.actor['id'], old['id']))}
        reset.reset_once(self.store)
        late = self.notice('late-hidden-read-test', 'opportunity', old['id'])
        submission = self.store.all_data()['submissions'][0]
        visible = self.notice('visible-read-test', 'submission', submission['id'])
        other = self.notice('other-user-read-test', 'submission', submission['id'], self.member)
        before = op.notifications(self.store, {'pageSize': 100}, self.actor)
        marked = op.mark_read(self.store, {'ids': [*old_ids, late, visible, other]}, self.actor)
        self.assertEqual(marked['unreadCount'], before['unreadCount'] - 1)
        self.assertEqual(op.mark_read(self.store, {'all': True}, self.actor)['unreadCount'], 0)
        self.assertEqual(op.notifications(self.store, {}, self.actor)['unreadCount'], 0)
        with self.store.connect() as db:
            for identifier in old_ids | {late, other}:
                self.assertIsNone(db.execute('SELECT read_at FROM tender_notifications WHERE id=?', (identifier,)).fetchone()[0])
            self.assertIsNotNone(db.execute('SELECT read_at FROM tender_notifications WHERE id=?', (visible,)).fetchone()[0])

    def test_builder_reuse_restores_suggestions_without_rewriting_historical_rows(self):
        shared = next(row['builder'] for row in self.store.all_data()['submissions'] if row.get('builder'))
        old_name = 'Builder only in the cleared register'
        self.create(builders=[old_name, shared])
        key = tenders.builder_key(old_name)
        with self.store.connect() as db:
            original = tuple(db.execute('SELECT * FROM opportunity_builders WHERE builder_key=?', (key,)).fetchone())
        reset.reset_once(self.store)
        builders = op.list_opportunities(self.store, {}, self.actor)['builders']
        self.assertNotIn(old_name, builders)
        self.assertIn(tenders.builder_key(shared), {tenders.builder_key(name) for name in builders})
        fresh = self.create('New use of a historical builder', builders=[old_name.upper()])
        builders = op.list_opportunities(self.store, {}, self.actor)['builders']
        self.assertIn(key, {tenders.builder_key(name) for name in builders})
        with self.store.connect() as db:
            self.assertEqual(tuple(db.execute('SELECT * FROM opportunity_builders WHERE builder_key=?', (key,)).fetchone()), original)
        reopened = tenders.Store(self.store.path)
        self.assertEqual(op.list_opportunities(reopened, {}, self.actor)['items'][0]['id'], fresh['id'])
        self.assertIn(key, {tenders.builder_key(name) for name in op.list_opportunities(reopened, {}, self.actor)['builders']})

    def test_failed_creation_does_not_restore_retired_builder_visibility(self):
        old_name = 'Builder whose new use failed'
        self.create(builders=[old_name])
        reset.reset_once(self.store)
        with patch.object(op, '_activity', side_effect=RuntimeError('Injected activity failure')):
            with self.assertRaisesRegex(RuntimeError, 'Injected activity failure'):
                self.create('Failed new project', builders=[old_name])
        result = op.list_opportunities(self.store, {}, self.actor)
        self.assertEqual(result['total'], 0)
        self.assertNotIn(old_name, result['builders'])

    def test_workbook_reimport_uses_new_ids_preserves_sources_and_keeps_normal_delete_protection(self):
        old = [self.create('Workbook project NSW'), self.create('Workbook project VIC', region='VIC')]
        before = {item['id']: self.row(item['id']) for item in old}
        reset.reset_once(self.store)
        deleted = self.create('Deleted individually after reset')
        op.delete(self.store, deleted, self.actor)
        records = [{'Opportunity ID': item['id'], 'Project Name': item['project'], 'Region': item['region'],
                    'Closing Date': item['due'], 'Builders': 'Workbook Builder', 'Working Scope': item['workingScope'],
                    'Historical tag': 'Untouched source value: ' + item['id']} for item in old]
        records.extend([
            {'Opportunity ID': deleted['id'], 'Project Name': 'Changed old export name', 'Region': 'VIC'},
            {'Project Name': deleted['project'], 'Region': deleted['region']},
        ])
        blob = workbook(records)
        parsed = parse_workbook('reimport.xlsx', blob)
        preview = import_workbook(self.store, 'reimport.xlsx', blob, self.actor)
        self.assertEqual((preview['counts']['new'], preview['counts']['duplicates'], preview['counts']['errors']), (2, 2, 0))
        self.assertTrue(all('deleted in the app' in row['message'] for row in preview['rows'][2:]))
        applied = import_workbook(self.store, 'reimport.xlsx', blob, self.actor, preview['token'])
        self.assertEqual(applied['counts']['added'], 2)
        new_ids = set(applied['importedIds'])
        self.assertFalse(new_ids & {item['id'] for item in [*old, deleted]})
        self.assertEqual({item['id']: self.row(item['id']) for item in old}, before)
        items = op.list_opportunities(self.store, {'pageSize': 100}, self.actor)['items']
        self.assertEqual({item['id'] for item in items}, new_ids)
        for item in items:
            source = next(record for record in parsed['records'] if record['project'] == item['project'])
            self.assertEqual(item['sourceData'], source['sourceData'])
        self.assertEqual(set(op.existing_keys(self.store).values()), new_ids | {deleted['id']})
        again = import_workbook(self.store, 'reimport.xlsx', blob, self.actor)
        self.assertEqual(again['counts']['new'], 0)
        restarted = tenders.Store(self.store.path)
        self.assertEqual({item['id'] for item in op.list_opportunities(restarted, {}, self.actor)['items']}, new_ids)

    def test_generated_id_collision_with_retained_parent_is_retried_without_mutation(self):
        old = self.create('Physical ID remains reserved')
        before = self.row(old['id'])
        reset.reset_once(self.store)
        fresh_id = uuid.uuid4()
        generate = uuid.uuid4
        identifiers = iter([uuid.UUID(old['id']), fresh_id])

        def next_identifier():
            try:
                return next(identifiers)
            except StopIteration:
                return generate()

        record = {'sourceRecordId': old['id'], 'project': old['project'], 'region': old['region'],
                  'sourceData': {'rawCells': {'Opportunity ID': old['id'], 'Unknown column': ['keep', 123]}}}
        original_record = deepcopy(record)
        with patch.object(op.uuid, 'uuid4', side_effect=next_identifier):
            result = op.import_records(self.store, [record], self.actor, 'collision.xlsx')
        self.assertEqual(result['items'], [str(fresh_id)])
        self.assertEqual(record, original_record)
        self.assertEqual(self.row(old['id']), before)
        self.assertEqual(op.detail(self.store, str(fresh_id), self.actor)['opportunity']['sourceData'], record['sourceData'])

    def test_deadline_sweeps_skip_retired_rows_and_process_new_rows_once(self):
        old = self.create('Old deadline', due='2099-01-11')
        before = self.row(old['id'])
        reset.reset_once(self.store)
        fresh = self.create('New deadline', due='2099-01-11')
        for day in (10, 10, 12, 12):
            op.process_due(self.store, force=True, now=datetime(2099, 1, day, tzinfo=timezone.utc))
        self.assertEqual(self.row(old['id']), before)
        with self.store.connect() as db:
            events = [row[0] for row in db.execute('SELECT event_key FROM opportunity_events')]
            self.assertFalse(any(old['id'] in key for key in events))
            self.assertEqual(sum(fresh['id'] in key for key in events), 2)
            self.assertEqual(db.execute('SELECT status FROM opportunities WHERE id=?', (fresh['id'],)).fetchone()[0], 'expired')
            self.assertEqual(db.execute("SELECT COUNT(*) FROM opportunity_activity WHERE opportunity_id=? AND action='Expired opportunity'", (fresh['id'],)).fetchone()[0], 1)
        visible = op.notifications(self.store, {'pageSize': 100}, self.actor)['items']
        self.assertTrue(any(row['kind'] == 'expired' and row.get('opportunityId') == fresh['id'] for row in visible))
        self.assertFalse(any(row.get('opportunityId') == old['id'] for row in visible))

    def test_reset_invalidates_a_pending_workbook_preview(self):
        old = self.create()
        blob = workbook([{'Project Name': 'New workbook project', 'Region': 'VIC'}])
        preview = import_workbook(self.store, 'pending.xlsx', blob, self.actor)
        reset.reset_once(self.store)
        with self.assertRaises(tenders.ValidationError) as caught:
            import_workbook(self.store, 'pending.xlsx', blob, self.actor, preview['token'])
        self.assertEqual(caught.exception.status, 409)
        self.assertEqual(op.list_opportunities(self.store, {}, self.actor)['total'], 0)
        self.assertIsNotNone(self.row(old['id']))

    def test_edits_decisions_and_deadline_sweep_preserve_unchanged_assignment_references(self):
        item = self.create(assignedTeamIds=[self.team], assignedUserIds=[self.member['id']])
        with self.store.connect() as db:
            db.execute('''CREATE TABLE historical_assignment_reference (
                opportunity_id TEXT, kind TEXT, person_id INTEGER,
                FOREIGN KEY(opportunity_id,kind,person_id)
                    REFERENCES opportunity_people(opportunity_id,kind,person_id))''')
            db.execute('INSERT INTO historical_assignment_reference VALUES(?,?,?)',
                       (item['id'], 'assignedTeamIds', self.team))
        item = op.save(self.store, {**item, 'remarks': 'Ordinary edit'}, self.actor)['opportunity']
        item = op.decide(self.store, {**item, 'decision': 'approved'}, self.actor)['opportunity']
        item = op.update_authority(self.store, {**item, 'assignedUserIds': [self.other_member['id']]}, self.actor, 'assign')['opportunity']
        self.assertEqual(op.list_opportunities(self.store, {'userId': str(self.other_member['id'])}, self.actor)['total'], 1)
        with self.store.connect() as db:
            assigned = {(row['kind'], row['person_id']) for row in db.execute(
                'SELECT kind,person_id FROM opportunity_people WHERE opportunity_id=?', (item['id'],))}
            self.assertEqual(assigned, {('assignedTeamIds', self.team), ('assignedUserIds', self.other_member['id'])})
        op.process_due(self.store, force=True, now=datetime(2100, 1, 1, tzinfo=timezone.utc))
        item = op.detail(self.store, item['id'], self.actor)['opportunity']
        self.assertEqual(item['status'], 'approved')
        op.delete(self.store, item, self.actor)
        with self.store.connect() as db:
            self.assertEqual(db.execute('PRAGMA foreign_key_check').fetchall(), [])
            self.assertEqual(db.execute('SELECT COUNT(*) FROM historical_assignment_reference').fetchone()[0], 1)
        self.assertEqual(op.list_opportunities(self.store, {}, self.actor)['total'], 0)


if __name__ == '__main__':
    unittest.main()
