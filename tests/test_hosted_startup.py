"""Run the real Passenger adapter with isolated databases and startup failures."""
import atexit
import base64
import io
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from urllib.parse import urlsplit
from zoneinfo import ZoneInfoNotFoundError

import passenger_wsgi as wsgi
from conest_app.auto_upgrade import CURRENT_SCHEMA
from conest_app import __version__, accounts, tenders
from conest_app.database import connect
from conest_app import runtime_health, startup_diagnostics


class HostedStartupTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.diagnostic_paths = patch.object(startup_diagnostics, '_candidate_directories',
                                             return_value=[root / 'diagnostics'])
        self.diagnostic_paths.start()
        self.environment = patch.dict(os.environ, {
            'CONEST_DB_PATH': str(root / 'catalog.db'),
            'CONEST_TENDERS_DB_PATH': str(root / 'tenders.db'),
        })
        self.environment.start()
        tenders.close_store()
        self.previous_initialized = wsgi._initialized
        self.previous_keeper = wsgi._database_keeper
        wsgi._initialized = False
        wsgi._database_keeper = None
        self.quiet = patch.object(wsgi.conest.ConEstHandler, 'log_message', lambda *_args: None)
        self.quiet.start()
        with connect() as db:
            identifier = db.execute("SELECT id FROM users WHERE role='admin' LIMIT 1").fetchone()[0]
            self.actor_id = identifier
            db.execute('UPDATE users SET must_change_password=0 WHERE id=?', (identifier,))
            self.token = accounts.start_session(db, identifier)

    def tearDown(self):
        tenders.close_store()
        if wsgi._database_keeper is not None:
            atexit.unregister(wsgi._database_keeper.close)
            wsgi._database_keeper.close()
        wsgi._database_keeper = self.previous_keeper
        wsgi._initialized = self.previous_initialized
        self.quiet.stop()
        self.environment.stop()
        self.diagnostic_paths.stop()
        self.temp.cleanup()

    def request(self, target, signed_in=False, method='GET', data=None):
        url = urlsplit(target)
        body = json.dumps(data).encode() if data is not None else b''
        environ = {
            'REQUEST_METHOD': method, 'PATH_INFO': url.path,
            'QUERY_STRING': url.query, 'wsgi.input': io.BytesIO(body),
            'HTTP_HOST': 'localhost', 'REMOTE_ADDR': '127.0.0.1',
            'HTTP_X_CONEST_REQUEST': '1',
            'CONTENT_LENGTH': str(len(body)), 'CONTENT_TYPE': 'application/json',
        }
        if signed_in:
            environ['HTTP_COOKIE'] = 'conest_session=' + self.token
            environ['HTTP_X_CSRF_TOKEN'] = accounts.csrf_token(self.token)
        response = {}

        def start(status, headers):
            response['status'] = int(status.split()[0])
            response['headers'] = dict(headers)

        raw = b''.join(wsgi.application(environ, start))
        response['raw'] = raw
        response['data'] = json.loads(raw) if raw and 'application/json' in response['headers'].get('Content-Type', '') else None
        return response

    def assert_safe_failure(self, response, code):
        self.assertEqual(response['status'], 503)
        data = response['data']
        self.assertFalse(data['success'])
        self.assertEqual(data['error_code'], code)
        self.assertRegex(data['error_reference'], r'^[a-f0-9]{12}$')
        self.assertIn(data['error_reference'], data['error'])
        self.assertNotIn('/private/', data['error'])
        self.assertNotIn('SECRET', data['error'])
        self.assertEqual(response['headers']['Cache-Control'], 'no-store')
        self.assertEqual(int(response['headers']['Content-Length']), len(response['raw']))

    def test_passenger_initialization_failure_returns_safe_json_and_retries(self):
        with patch.object(wsgi, '_ensure_initialized', side_effect=PermissionError('/private/SECRET/database.db')):
            with self.assertLogs(runtime_health.log, level='ERROR') as logged:
                response = self.request('/api/tenders/bootstrap', signed_in=True)
        self.assert_safe_failure(response, 'SESSION_STARTUP_FAILED')
        self.assertIn('/private/SECRET/database.db', '\n'.join(logged.output))
        self.assertEqual(self.request('/api/tenders/bootstrap', signed_in=True)['status'], 200)

    def test_passenger_import_failure_and_head_have_valid_responses(self):
        failure = ModuleNotFoundError('SECRET production import missing')
        with patch.object(wsgi, '_startup_import_error', failure):
            with self.assertLogs(runtime_health.log, level='ERROR'):
                response = self.request('/api/auth/session')
                head = self.request('/api/auth/session', method='HEAD')
        self.assert_safe_failure(response, 'SESSION_STARTUP_FAILED')
        self.assertEqual(head['status'], 503)
        self.assertEqual(head['raw'], b'')
        self.assertGreater(int(head['headers']['Content-Length']), 0)

    def test_tender_failure_keeps_session_and_authentication_rules(self):
        with patch.object(tenders, 'get_store', side_effect=PermissionError('/private/SECRET/tenders.db')):
            with self.assertLogs(runtime_health.log, level='ERROR'):
                failed = self.request('/api/tenders/bootstrap', signed_in=True)
            self.assert_safe_failure(failed, 'TENDERS_STARTUP_FAILED')
            self.assertEqual(self.request('/api/tenders/bootstrap')['status'], 401)
            self.assertEqual(self.request('/api/auth/session', signed_in=True)['status'], 200)
        recovered = self.request('/api/tenders/bootstrap', signed_in=True)
        self.assertEqual(recovered['status'], 200)
        self.assertTrue(recovered['data']['data']['submissions'])
        self.assertEqual(recovered['headers']['X-ConEst-Adapter'], __version__)

    def test_readiness_checks_tenders_when_catalog_alone_is_healthy(self):
        with patch.object(tenders, 'get_store', side_effect=PermissionError('/private/SECRET/tenders.db')):
            self.assertEqual(self.request('/api/health')['status'], 200)
            with self.assertLogs(runtime_health.log, level='ERROR'):
                deep = self.request('/api/health?check=tenders')
        self.assert_safe_failure(deep, 'TENDERS_STARTUP_FAILED')

    def test_readiness_is_private_and_does_not_change_generation(self):
        store = tenders.get_store()
        with store.connect() as db:
            before = db.execute('SELECT value FROM opportunity_generation').fetchone()[0]
        response = self.request('/api/health?check=tenders')
        self.assertEqual(response['status'], 200, response['data'])
        self.assertEqual(response['data']['checks'], {
            'tenders': 'ok', 'opportunities': 'ok', 'timezones': 'ok', 'sqlite': 'ok',
        })
        self.assertNotIn(str(store.path), response['raw'].decode())
        self.assertNotIn('submissions', response['data'])
        with store.connect() as db:
            self.assertEqual(db.execute('SELECT value FROM opportunity_generation').fetchone()[0], before)
            self.assertEqual(db.execute('PRAGMA quick_check').fetchone()[0], 'ok')

    def test_missing_timezone_or_sqlite_support_is_not_a_false_healthy_check(self):
        for target, failure in [
            ('ZoneInfo', ZoneInfoNotFoundError('SECRET unavailable timezone path')),
            ('check_sqlite_runtime', RuntimeError('SECRET missing SQLite JSON')),
        ]:
            with self.subTest(target=target), patch.object(runtime_health, target, side_effect=failure):
                with self.assertLogs(runtime_health.log, level='ERROR'):
                    response = self.request('/api/health?check=tenders')
                self.assert_safe_failure(response, 'TENDERS_STARTUP_FAILED')

    def test_legacy_builder_migration_recovers_health_and_accepts_new_builder(self):
        from conest_app.opportunity_reset import RESET_KEY
        from conest_app.opportunity_retention import active_sql

        path = Path(os.environ['CONEST_TENDERS_DB_PATH'])
        # Reproduce a retained hosted database, before the real Passenger
        # readiness request constructs the Store and applies the pending reset.
        with sqlite3.connect(path) as db:
            db.execute('CREATE TABLE opportunities(id TEXT PRIMARY KEY, body TEXT NOT NULL)')
            db.execute('INSERT INTO opportunities VALUES(?,?)', ('legacy-project', json.dumps({
                'id': 'legacy-project', 'project': 'Legacy VIC project', 'region': 'VIC',
                'builders': ['Old Builder'], 'distanceOrigin': 'CBD, Melbourne',
            })))
            db.execute('''CREATE TABLE opportunity_builders (
                id TEXT PRIMARY KEY NOT NULL, name TEXT NOT NULL,
                legacy_contact TEXT NOT NULL,
                opportunity_id TEXT NOT NULL REFERENCES opportunities(id))''')
            db.execute('INSERT INTO opportunity_builders VALUES(?,?,?,?)',
                       ('old-builder', 'Old Builder', 'Retain historical contact', 'legacy-project'))

        health = self.request('/api/health?check=tenders')
        self.assertEqual(health['status'], 200, health['data'])
        self.assertEqual(health['data']['release'], __version__)
        self.assertEqual(health['data']['schema_version'], CURRENT_SCHEMA)
        self.assertEqual(health['data']['checks'], {
            'tenders': 'ok', 'opportunities': 'ok', 'timezones': 'ok', 'sqlite': 'ok',
        })
        self.assertEqual(self.request('/api/tenders/bootstrap', signed_in=True)['status'], 200)

        store = tenders.get_store()
        with store.connect() as db:
            self.assertEqual(tuple(db.execute(
                'SELECT * FROM opportunity_builders_legacy_v354').fetchone()),
                ('old-builder', 'Old Builder', 'Retain historical contact', 'legacy-project'))
            self.assertEqual(db.execute('SELECT COUNT(*) FROM opportunities').fetchone()[0], 1)
            self.assertEqual(db.execute('SELECT COUNT(*) FROM opportunities WHERE ' +
                                       active_sql('opportunity', 'opportunities.id')).fetchone()[0], 0)
            self.assertEqual(db.execute('SELECT COUNT(*) FROM opportunity_builders WHERE ' +
                                       active_sql('builder', 'opportunity_builders.builder_key')).fetchone()[0], 0)
            marker = db.execute('SELECT body FROM settings WHERE id=?', (RESET_KEY,)).fetchone()[0]

        created = self.request('/api/tenders/opportunities/save', signed_in=True, method='POST', data={
            'project': 'Created after hosted migration', 'region': 'VIC',
            'builders': ['New Builder'], 'workingScope': 'Electrical services',
        })
        self.assertEqual(created['status'], 200, created['data'])
        with store.connect() as db:
            self.assertEqual(tuple(db.execute('SELECT builder_key,name FROM opportunity_builders WHERE ' +
                                             active_sql('builder', 'opportunity_builders.builder_key')).fetchone()),
                             ('newbuilder', 'New Builder'))
            self.assertEqual(db.execute('SELECT COUNT(*) FROM opportunities').fetchone()[0], 2)
            self.assertEqual(db.execute('PRAGMA foreign_key_check').fetchall(), [])
        self.assertEqual(len(list((path.parent / 'upgrade_backups').glob('*schema-v354-*.db'))), 1)

        tenders.close_store()
        self.assertEqual(self.request('/api/health?check=tenders')['status'], 200)
        with tenders.get_store().connect() as db:
            self.assertEqual(db.execute('SELECT body FROM settings WHERE id=?', (RESET_KEY,)).fetchone()[0], marker)
            self.assertEqual(db.execute('SELECT COUNT(*) FROM opportunities').fetchone()[0], 2)
            self.assertEqual(db.execute('SELECT COUNT(*) FROM opportunities WHERE ' +
                                       active_sql('opportunity', 'opportunities.id')).fetchone()[0], 1)
            self.assertEqual(db.execute('SELECT COUNT(*) FROM opportunity_builders WHERE ' +
                                       active_sql('builder', 'opportunity_builders.builder_key')).fetchone()[0], 1)
        self.assertEqual(len(list((path.parent / 'upgrade_backups').glob('*schema-v354-*.db'))), 1)

    def test_legacy_creator_column_accepts_current_form_after_completed_reset(self):
        """The hosted screenshot failure must recover through the real adapter."""
        from conest_app.opportunity_reset import RESET_KEY
        from conest_app.opportunity_retention import ensure_schema

        path = Path(os.environ['CONEST_TENDERS_DB_PATH'])
        original = {
            'id': 'retired-creator-project', 'project': 'Retired creator project',
            'region': 'NSW', 'status': 'pending', 'revision': 3,
            'createdBy': self.actor_id, 'assignedUserIds': [self.actor_id],
            'contact': 'Historical builder contact', 'email': 'legacy invalid email',
        }
        original_json = json.dumps(original)
        marker = json.dumps({'release': '3.5.5', 'completedAt': '2026-10-21T08:00:00Z'})
        with sqlite3.connect(path) as db:
            db.execute('PRAGMA foreign_keys=ON')
            db.execute('''CREATE TABLE opportunities (
                id TEXT PRIMARY KEY, body TEXT NOT NULL, creator_id INTEGER NOT NULL)''')
            db.execute('INSERT INTO opportunities VALUES(?,?,?)',
                       (original['id'], original_json, self.actor_id))
            db.execute('''CREATE TABLE opportunity_people (
                opportunity_id TEXT NOT NULL REFERENCES opportunities(id),
                kind TEXT NOT NULL, person_id INTEGER NOT NULL,
                PRIMARY KEY(opportunity_id,kind,person_id))''')
            db.execute('INSERT INTO opportunity_people VALUES(?,?,?)',
                       (original['id'], 'assignedUserIds', self.actor_id))
            db.execute('''CREATE TABLE historical_creator_note (
                opportunity_id TEXT REFERENCES opportunities(id), note TEXT NOT NULL)''')
            db.execute('INSERT INTO historical_creator_note VALUES(?,?)',
                       (original['id'], 'Keep the original creator relationship'))
            db.execute('CREATE TABLE settings(id TEXT PRIMARY KEY, body TEXT NOT NULL)')
            db.execute('INSERT INTO settings VALUES(?,?)', (RESET_KEY, marker))
            ensure_schema(db)
            db.execute('INSERT INTO opportunity_reset_entries VALUES(?,?,?,NULL)',
                       ('opportunity', original['id'], '2026-10-21T08:00:00Z'))

        health = self.request('/api/health?check=tenders')
        self.assertEqual(health['status'], 200, health['data'])
        self.assertEqual(health['data']['release'], __version__)
        listed = self.request('/api/tenders/opportunities', signed_in=True)
        self.assertEqual(listed['data']['total'], 0)

        created = self.request('/api/tenders/opportunities/save', signed_in=True,
                               method='POST', data={
            'project': 'Current form after creator migration', 'region': 'VIC',
            'builders': ['Builder One', 'Builder Two'], 'workingScope': 'Electrical works',
            'due': '2099-11-25', 'link': 'example.com/tender/documents',
            'assignedUserIds': [self.actor_id],
            # An older browser may still submit removed shared contact fields.
            'contact': 'Unused shared contact', 'email': 'not an email',
        })
        self.assertEqual(created['status'], 200, created['data'])
        item = created['data']['opportunity']
        self.assertEqual(item['createdBy'], self.actor_id)
        self.assertEqual(item['link'], 'https://example.com/tender/documents')
        self.assertEqual(item['builders'], ['Builder One', 'Builder Two'])
        self.assertEqual(item['status'], 'pending')
        self.assertNotIn('contact', item)
        self.assertNotIn('email', item)

        approved = self.request('/api/tenders/opportunities/decision', signed_in=True,
                                method='POST', data={
            'id': item['id'], 'revision': item['revision'], 'decision': 'approved',
        })
        self.assertEqual(approved['status'], 200, approved['data'])
        self.assertEqual(approved['data']['opportunity']['link'], item['link'])
        self.assertEqual(approved['data']['opportunity']['decisionBy'], self.actor_id)
        history = self.request('/api/tenders/opportunities/activity?id=' + item['id'], signed_in=True)
        self.assertTrue(any(row['action'] == 'Approved opportunity'
                            for row in history['data']['activity']))
        notices = self.request('/api/tenders/notifications', signed_in=True)
        self.assertTrue(any(row.get('opportunityId') == item['id'] for row in notices['data']['items']))
        contribution = self.request('/api/users/' + str(self.actor_id) + '/performance', signed_in=True)
        self.assertEqual(contribution['data']['opportunities']['projects'], 1)

        with tenders.get_store().connect() as db:
            self.assertEqual(db.execute('PRAGMA foreign_key_check').fetchall(), [])
            self.assertEqual(db.execute('SELECT body FROM settings WHERE id=?', (RESET_KEY,)).fetchone()[0], marker)
            self.assertEqual(tuple(db.execute('SELECT * FROM opportunities_legacy_v356').fetchone()),
                             (original['id'], original_json, self.actor_id))
            self.assertEqual(db.execute('SELECT body FROM opportunities WHERE id=?',
                                        (original['id'],)).fetchone()[0], original_json)
            self.assertEqual(db.execute('PRAGMA foreign_key_list(opportunity_people)').fetchone()['table'],
                             'opportunities')
            self.assertEqual(db.execute('PRAGMA foreign_key_list(historical_creator_note)').fetchone()['table'],
                             'opportunities_legacy_v356')
        tenders.close_store()
        self.assertEqual(self.request('/api/health?check=tenders')['status'], 200)
        listing = self.request('/api/tenders/opportunities', signed_in=True)['data']
        self.assertEqual(listing['total'], 1)
        self.assertEqual(listing['items'][0]['id'], item['id'])
        self.assertEqual(listing['items'][0]['status'], 'approved')

    def test_legacy_relationships_support_full_hosted_opportunity_workflow(self):
        from openpyxl import Workbook, load_workbook

        path = Path(os.environ['CONEST_TENDERS_DB_PATH'])
        with sqlite3.connect(path) as db:
            db.execute('PRAGMA foreign_keys=ON')
            db.execute('CREATE TABLE opportunities(id TEXT PRIMARY KEY, body TEXT NOT NULL)')
            db.execute('INSERT INTO opportunities VALUES(?,?)', ('retired-project', json.dumps({
                'id': 'retired-project', 'project': 'Historical linked project', 'region': 'NSW',
                'builders': ['Retired Builder'], 'assignedUserIds': [self.actor_id],
            })))
            db.execute('CREATE TABLE legacy_actors(id INTEGER PRIMARY KEY)')
            db.execute('INSERT INTO legacy_actors VALUES(?)', (self.actor_id,))
            db.execute('''CREATE TABLE opportunity_people (
                opportunity_id TEXT NOT NULL REFERENCES opportunities(id),
                kind TEXT NOT NULL, person_id INTEGER NOT NULL,
                PRIMARY KEY(opportunity_id,kind,person_id))''')
            db.execute('INSERT INTO opportunity_people VALUES(?,?,?)',
                       ('retired-project', 'assignedUserIds', self.actor_id))
            db.execute('CREATE INDEX opportunity_people_lookup ON opportunity_people(kind,person_id,opportunity_id)')
            db.execute('''CREATE TABLE historical_assignment_approval (
                opportunity_id TEXT, kind TEXT, person_id INTEGER, note TEXT,
                FOREIGN KEY(opportunity_id,kind,person_id)
                REFERENCES opportunity_people(opportunity_id,kind,person_id))''')
            db.execute('INSERT INTO historical_assignment_approval VALUES(?,?,?,?)',
                       ('retired-project', 'assignedUserIds', self.actor_id, 'Retain original approval'))
            db.execute('''CREATE TABLE opportunity_activity (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                opportunity_id TEXT NOT NULL REFERENCES opportunities(id),
                actor_id INTEGER NOT NULL REFERENCES legacy_actors(id),
                actor_name TEXT NOT NULL, project TEXT NOT NULL, action TEXT NOT NULL,
                details TEXT NOT NULL DEFAULT '{}', created_at TEXT NOT NULL,
                original_source TEXT NOT NULL)''')
            db.execute('INSERT INTO opportunity_activity VALUES(?,?,?,?,?,?,?,?,?)',
                       (41, 'retired-project', self.actor_id, 'Original actor',
                        'Historical linked project', 'Created opportunity', '{}',
                        '2026-01-01T00:00:00+00:00', 'Preserve this legacy field'))
            db.execute('CREATE INDEX opportunity_activity_record ON opportunity_activity(opportunity_id,id DESC)')

        def call(route, data=None, method=None):
            response = self.request('/api/tenders/' + route, signed_in=True,
                                    method=method or ('POST' if data is not None else 'GET'), data=data)
            self.assertEqual(response['status'], 200, response['data'] or response['raw'])
            return response['data']

        health = self.request('/api/health?check=tenders')
        self.assertEqual(health['status'], 200, health['data'])
        self.assertEqual(call('opportunities')['total'], 0)
        self.assertEqual(self.request('/api/tenders/opportunities/detail?id=retired-project',
                                      signed_in=True)['status'], 404)
        created = call('opportunities/save', {
            'project': 'New live project', 'region': 'VIC', 'builders': ['Current Builder'],
            'workingScope': 'Electrical fitout', 'due': '2099-11-25',
            'assignedUserIds': [self.actor_id],
        })['opportunity']
        self.assertEqual(created['status'], 'pending')
        self.assertEqual(created['distanceOrigin'], 'CBD, Melbourne')
        approved = call('opportunities/decision', {
            'id': created['id'], 'revision': created['revision'], 'decision': 'approved',
        })['opportunity']
        self.assertEqual(approved['status'], 'approved')
        self.assertEqual(approved['decisionBy'], self.actor_id)
        unassigned = call('opportunities/assign', {
            'id': approved['id'], 'revision': approved['revision'], 'assignedUserIds': [],
        })['opportunity']
        self.assertEqual(unassigned['assignedUserIds'], [])
        call('opportunities/regions', {'code': 'VIC', 'origin': 'CBD, Melbourne',
                                      'timezone': 'Australia/Melbourne'})

        book = Workbook()
        sheet = book.active
        sheet.title = 'Opportunities'
        sheet.append(['Project Name', 'Region', 'Closing Date', 'Builders', 'Working Scope'])
        sheet.append(['New workbook project', 'NSW', '2099-11-25', 'Current Builder', 'Lighting'])
        sheet.append(['Historical deadline import', 'VIC', '2000-01-01', 'Second Builder', 'Power'])
        stream = io.BytesIO()
        book.save(stream)
        upload = {'filename': 'opportunities.xlsx', 'content': base64.b64encode(stream.getvalue()).decode()}
        preview = call('opportunities/import', {**upload, 'mode': 'preview'})
        imported = call('opportunities/import', {**upload, 'mode': 'apply', 'token': preview['token']})
        self.assertEqual(imported['counts']['added'], 2)
        listed = call('opportunities')
        self.assertEqual(listed['total'], 3)
        self.assertEqual({item['status'] for item in listed['items']}, {'pending', 'approved', 'expired'})
        expired = next(item for item in listed['items'] if item['status'] == 'expired')
        history = call('opportunities/activity?id=' + expired['id'])
        self.assertTrue(any(event['action'] == 'Expired opportunity' for event in history['activity']))
        inbox = call('notifications')
        self.assertGreater(inbox['unread'], 0)
        self.assertEqual(call('notifications/read', {'all': True})['unread'], 0)

        exported = self.request('/api/tenders/opportunities/export.xlsx', signed_in=True)
        self.assertEqual(exported['status'], 200)
        workbook = load_workbook(io.BytesIO(exported['raw']), read_only=True, data_only=True)
        self.assertNotIn('Historical linked project',
                         {cell for row in workbook.active.iter_rows(values_only=True) for cell in row})
        workbook.close()

        contributions = self.request('/api/users/' + str(self.actor_id) + '/performance', signed_in=True)
        self.assertEqual(contributions['status'], 200, contributions['data'])
        self.assertEqual(contributions['data']['opportunities']['rowsAdded'], 2)
        self.assertEqual(contributions['data']['opportunities']['projects'], 4)
        created_count = next(item['count'] for item in contributions['data']['opportunities']['actions']
                             if item['action'] == 'Created opportunity')
        self.assertEqual(created_count, 2)

        store = tenders.get_store()
        with store.connect() as db:
            self.assertEqual(db.execute('PRAGMA foreign_key_check').fetchall(), [])
            self.assertEqual(db.execute('SELECT COUNT(*) FROM historical_assignment_approval').fetchone()[0], 1)
            self.assertEqual(db.execute('SELECT original_source FROM opportunity_activity_legacy_v355').fetchone()[0],
                             'Preserve this legacy field')
            self.assertEqual(db.execute('SELECT actor_name FROM opportunity_activity WHERE id=41').fetchone()[0],
                             'Original actor')
            self.assertGreater(db.execute("SELECT COUNT(*) FROM opportunity_activity WHERE actor_id=0 AND action='Expired opportunity'").fetchone()[0], 0)
            self.assertEqual(db.execute("SELECT tbl_name FROM sqlite_master WHERE name='opportunity_activity_record_v355'").fetchone()[0],
                             'opportunity_activity')
            self.assertEqual(db.execute("SELECT tbl_name FROM sqlite_master WHERE name='opportunity_people_lookup_v355'").fetchone()[0],
                             'opportunity_people')
        tenders.close_store()
        self.assertEqual(self.request('/api/health?check=tenders')['status'], 200)
        self.assertEqual(call('opportunities')['total'], 3)


if __name__ == '__main__':
    unittest.main()
