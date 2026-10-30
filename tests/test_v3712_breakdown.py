"""v3.7.12 atomic area allocation, common/optional cost, persistence and exports."""
import unittest,json,math
from pathlib import Path
from openpyxl import load_workbook
import test_estimate_sections as helpers
from conest_app import projects,estimate_tools,project_json,project_sharing,productivity
from conest_app.database import connect
from conest_app.estimate_areas import save_breakdown,COMMON,allocations
from conest_app.accounts import AccessError
from conest_app.exporter import export_project

class StructuredBreakdownTests(unittest.TestCase):
    setUp=helpers.EstimateSectionTests.setUp
    tearDown=helpers.EstimateSectionTests.tearDown
    prepare=helpers.EstimateSectionTests.prepare

    def body(self,data,allocations_by_id,areas=None,qty_by_id=None):
        return dict(expected_revision=data['project']['revision'],areas=areas or ['Ground Floor','Level 1'],rows=[
            dict(id=l['id'],expected_revision=l['revision'],quantity=(qty_by_id or {}).get(l['id'],l['quantity']),
                 allocations=allocations_by_id.get(l['id'],json.loads(l['area_allocations'] or '[]'))) for l in data['lines']])

    def test_single_multiple_common_unallocated_and_optional_total(self):
        with connect() as c:
            pid,data=self.prepare(c)
            lines=data['lines'];a=lines[0];b=lines[1]
            before=dict(data['totals'])
            body=self.body(data,{a['id']:[{'area':COMMON,'quantity':a['quantity'],'common':True}],
                            b['id']:[{'area':'Ground Floor','quantity':b['quantity']*.4},
                                     {'area':'Level 1','quantity':b['quantity']*.6}]})
            after=save_breakdown(c,pid,self.data['owner'],body)
            self.assertEqual(after['totals'],before)
            self.assertEqual(after['area_names'],['Ground Floor','Level 1'])
            self.assertEqual(next(g for g in after['area_summary'] if g['area']==COMMON)['item_count'],1)
            self.assertEqual(sum(g['included']['quoted'] for g in after['area_summary']),before['quoted'])
            self.assertEqual(len(after['area_summary'][-1]['categories']),1)
            self.assertEqual(len([g for g in after['area_summary'] if g['area']==COMMON]),1)
            copied=estimate_tools.copy_project(c,pid,self.data['owner'],{'name':'Structured copy'})
            self.assertEqual(projects.calculate_project(c,copied)['area_names'],['Ground Floor','Level 1'])
            self.assertEqual(projects.calculate_project(c,copied)['totals'],before)
            raw=project_json.export_json(c,pid,self.data['owner'])
            self.assertEqual(raw['areas'],['Ground Floor','Level 1'])
            clone=project_json.import_json(c,raw,self.data['owner'])
            self.assertEqual(projects.calculate_project(c,clone)['area_names'],raw['areas'])
            self.assertEqual(projects.calculate_project(c,clone)['totals'],before)
            self.assertIn('Common to All Areas',project_sharing.shared_html(after))
            path=export_project(c,pid,Path(self.tmp.name)/'breakdown.xlsx')
        wb=load_workbook(path)
        self.assertIn('STRUCTURED BREAKDOWN',wb.sheetnames)
        ws=wb['STRUCTURED BREAKDOWN']
        self.assertEqual(sum(c.value==COMMON for row in ws for c in row),1)
        self.assertIn('PROJECT TOTAL (COMMON COUNTED ONCE)',[ws.cell(i,1).value for i in range(1,ws.max_row+1)])
        self.assertTrue(any("'AREA ALLOCATIONS'!" in str(c.value) for row in ws for c in row))
        wb.close()
        with connect() as c:
            self.assertEqual(projects.calculate_project(c,pid)['area_names'],['Ground Floor','Level 1'])

    def test_atomic_conflicts_and_invalid_names(self):
        with connect() as c:
            pid,data=self.prepare(c)
            lid=data['lines'][0]['id'];before=projects.calculate_project(c,pid)
            original_alloc=before['lines'][0]['area_allocations']
            bad=self.body(data,{lid:[{'area':'Ground Floor','quantity':999}]})
            with self.assertRaises(ValueError):save_breakdown(c,pid,self.data['owner'],bad)
            self.assertEqual(projects.calculate_project(c,pid)['lines'][0]['area_allocations'],original_alloc)
            self.assertFalse(c.execute('SELECT 1 FROM project_breakdown_areas WHERE project_id=?',(pid,)).fetchone())
            for invalid in (['Common to All Areas'],['Level 1','level 1'],['Unallocated']):
                with self.assertRaises(ValueError):save_breakdown(c,pid,self.data['owner'],{**self.body(data,{}),'areas':invalid})
            with self.assertRaises(AccessError):save_breakdown(c,pid,self.data['viewer'],self.body(data,{}))
            body=self.body(data,{lid:[{'area':'Ground Floor','quantity':before['lines'][0]['quantity']}]})
            saved=save_breakdown(c,pid,self.data['owner'],body)
            with self.assertRaises(ValueError):save_breakdown(c,pid,self.data['owner'],body)
            self.assertEqual(saved['area_names'],['Ground Floor','Level 1'])

    def test_revision_roundtrip_and_optional_cost(self):
        with connect() as c:
            pid,data=self.prepare(c)
            lid=data['lines'][0]['id'];q=data['lines'][0]['quantity']
            rid=productivity.save_revision(c,pid,self.data['owner'],{'label':'Before breakdown'})
            saved=save_breakdown(c,pid,self.data['owner'],self.body(data,{lid:[{'area':COMMON,'quantity':q,'common':True}]}))
            project=projects.calculate_project(c,pid)
            self.assertTrue(any(g['area']==COMMON for g in project['area_summary']))
            self.assertEqual(project['totals'],data['totals'])
            self.assertIn('Common',productivity.save_revision(c,pid,self.data['owner'],{'label':'After breakdown'}) and project_sharing.shared_html(project))
            after_id=c.execute("SELECT id FROM tender_revisions WHERE label='After breakdown'").fetchone()[0]
            self.assertTrue(projects.calculate_project(c,pid)['area_names'])
            productivity.redeploy_revision(c,pid,rid,self.data['owner'],{'expected_revision':c.execute('SELECT revision FROM projects WHERE id=?',(pid,)).fetchone()[0]})
            self.assertFalse(projects.calculate_project(c,pid)['area_names'])
            productivity.redeploy_revision(c,pid,after_id,self.data['owner'],{'expected_revision':c.execute('SELECT revision FROM projects WHERE id=?',(pid,)).fetchone()[0]})
            again=projects.calculate_project(c,pid)
            self.assertTrue(any(g['area']==COMMON for g in again['area_summary']))
            after=helpers.EstimateSectionTests.toggle(self,c,category='Lighting')
            self.assertAlmostEqual(sum(g['included']['quoted'] for g in after['area_summary']),after['totals']['quoted'])
            self.assertGreater(sum(g['optional']['quoted'] for g in after['area_summary']),0)
            self.assertIn('Optional cost',project_sharing.shared_html(after))

    def test_legacy_shared_keeps_pricing(self):
        with connect() as c:
            pid,data=self.prepare(c);lid=data['lines'][0]['id'];qty=data['lines'][0]['quantity']
            distributed=[{'area':'Ground Floor','quantity':qty*.5,'shared':True},{'area':'Level 1','quantity':qty*.5,'shared':True}]
            projects.update_project_line(c,pid,lid,{'area_allocations':distributed})
            old=projects.calculate_project(c,pid)
            payload=self.body(old,{},areas=['Ground Floor','Level 1'])
            saved=save_breakdown(c,pid,self.data['owner'],payload)
            self.assertEqual(saved['totals'],old['totals'])
            self.assertEqual(allocations(next(l for l in saved['lines'] if l['id']==lid)['area_allocations']),distributed)

if __name__=='__main__':unittest.main()
