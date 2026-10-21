"""Opportunity workbook fidelity, append-only imports, and atomic preview checks."""
import io
import json
import os
import struct
import tempfile
import unittest
from contextlib import nullcontext
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook

from conest_app import opportunities, opportunity_http, tenders
from conest_app.database import connect
from conest_app.opportunity_import import (
    EXPORT_HEADERS, export_workbook, import_workbook, initialize_seed, parse_workbook,
)
from conest_app.tender_export import sheet_xml
from conest_app.tenders import ValidationError


def replace_parts(blob, replacements):
    result = io.BytesIO()
    with ZipFile(io.BytesIO(blob)) as source, ZipFile(result, 'w', ZIP_DEFLATED) as destination:
        for name in source.namelist():
            destination.writestr(name, replacements.get(name, source.read(name)))
    return result.getvalue()


def workbook(rows, headers=None):
    headers = headers or EXPORT_HEADERS
    xml = sheet_xml('Opportunities', headers, rows, [30] * len(headers), 'Test workbook', '', '')
    # Fixture structure uses the actual generated export, independently of the
    # user-facing blank template, its shared strings, or its instruction sheet.
    empty_store = SimpleNamespace(connect=lambda: nullcontext(SimpleNamespace(execute=lambda *_: [])))
    with patch('conest_app.opportunities.directory', return_value=([], [])):
        base = export_workbook(empty_store)
    return replace_parts(base, {'xl/worksheets/sheet1.xml': xml})


def row(project='New import project', **kwargs):
    values = {
        'Project Name': project, 'Region': 'NSW', 'Closing Date': '2099-11-25',
        'Builders': 'First Builder\nSecond Builder', 'Budget': '5m-10m',
        'Distance (km)': 20.2, 'Distance Origin': '63-79 Parramatta Rd Silverwater, NSW 2128',
        'Working Scope': 'Lighting, power and communications',
    }
    values.update(kwargs)
    return [values.get(header, '') for header in EXPORT_HEADERS]


class OpportunityImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.environment = patch.dict(os.environ, {'CONEST_DB_PATH': str(Path(self.temp.name) / 'accounts.db')})
        self.environment.start()
        self.addCleanup(self.environment.stop)
        self.today = patch('conest_app.opportunities._today', return_value=date(2026, 10, 21))
        self.today.start()
        self.addCleanup(self.today.stop)
        self.store = tenders.Store(Path(self.temp.name) / 'tenders.db')
        with connect() as db:
            self.actor = dict(db.execute("SELECT id,name,role FROM users WHERE role='admin' LIMIT 1").fetchone())

    def preview(self, blob, **kwargs):
        return import_workbook(self.store, 'import.xlsx', blob, self.actor, **kwargs)

    def apply(self, blob, **kwargs):
        preview = self.preview(blob, **kwargs)
        return import_workbook(self.store, 'import.xlsx', blob, self.actor, expected_token=preview['token'], **kwargs)

    def items(self):
        with self.store.connect() as db:
            return [json.loads(r['body']) for r in db.execute('SELECT body FROM opportunities')]

    def create(self, project='Existing fictional project', **kwargs):
        fields = {'project': project, 'region': 'NSW', 'due': '2099-11-25', 'builders': ['Test Builder']}
        fields.update(kwargs)
        return opportunities.save(self.store, fields, self.actor)['opportunity']

    def test_packaged_register_starts_empty_and_seed_is_idempotent(self):
        self.assertEqual(self.items(), [])
        self.assertEqual(initialize_seed(self.store)['added'], 0)
        fresh = tenders.Store(self.store.path)
        with fresh.connect() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM opportunities').fetchone()[0], 0)

    def test_seed_repeat_does_not_overwrite_or_resurrect_deleted_record(self):
        record = parse_workbook('fictional.xlsx', workbook([row('Fictional seed project', Status='Approved')]))['records'][0]
        fixture_root = Path(self.temp.name) / 'seed_fixture'
        seed_dir = fixture_root / 'opportunities_seed'
        seed_dir.mkdir(parents=True)
        (seed_dir / 'seed.json').write_text(json.dumps({'sources': [
            {'filename': 'fictional.xlsx', 'seedKey': 'test-only-fictional-seed', 'records': [record]},
        ]}), encoding='utf-8')
        with patch('conest_app.opportunity_import.ROOT', fixture_root):
            self.assertEqual(initialize_seed(self.store)['added'], 1)
            original = self.items()[0]
            self.assertEqual(original['sourceStatus'], 'approved')
            self.assertIsNone(original.get('decisionBy'))
            self.assertIsNone(original['createdBy'])
            self.assertEqual(original['createdByName'], 'Workbook import')
            updated = opportunities.save(self.store, {**original, 'remarks': 'Changed inside the app'}, self.actor)['opportunity']
            opportunities.delete(self.store, {'id': updated['id'], 'revision': updated['revision']}, self.actor)
            self.assertEqual(initialize_seed(self.store)['added'], 0)
            fresh = tenders.Store(self.store.path)
            with fresh.connect() as db:
                saved = db.execute('SELECT body,deleted_at FROM opportunities WHERE id=?', (updated['id'],)).fetchone()
                self.assertTrue(saved['deleted_at'])
                self.assertEqual(json.loads(saved['body'])['remarks'], 'Changed inside the app')
                self.assertEqual(db.execute('SELECT COUNT(*) FROM opportunities').fetchone()[0], 1)

    def test_export_roundtrip_preserves_dates_budget_builders_and_literal_notes(self):
        self.apply(workbook([row('Fictional export NSW'), row('Fictional export VIC', Region='VIC', **{'Distance Origin': 'CBD, Melbourne'})]))
        blob = export_workbook(self.store)
        parsed = parse_workbook('export.xlsx', blob)
        self.assertEqual(len(parsed['records']), 2)
        self.assertFalse(parsed['issues'])
        result = self.apply(blob)
        self.assertEqual(result['counts']['duplicates'], 2)
        self.assertEqual(result['counts']['added'], 0)
        book = load_workbook(io.BytesIO(blob), data_only=False)
        self.assertEqual(book['Opportunities'].freeze_panes, 'C6')
        self.assertEqual(book['Opportunities'].auto_filter.ref, 'A5:T7')
        book.close()
        literal = '=HYPERLINK("https://invalid.example","literal text")'
        parsed = parse_workbook('literal.xlsx', workbook([row(Comments=literal)]))
        self.assertEqual(parsed['records'][0]['remarks'], literal)

    def test_add_duplicate_and_stable_export_id_never_overwrite_app_changes(self):
        blob = workbook([row(), row()])
        preview = self.preview(blob)
        self.assertEqual((preview['counts']['new'], preview['counts']['duplicates']), (1, 1))
        result = self.apply(blob)
        self.assertEqual(result['counts']['added'], 1)
        self.assertEqual(len(self.items()), 1)
        item = next(i for i in self.items() if i['project'] == 'New import project')
        changed = opportunities.save(self.store, {**item, 'remarks': 'Retain this edit'}, self.actor)['opportunity']
        renamed_export = workbook([row('Renamed in workbook', **{'Opportunity ID': changed['id'], 'Comments': 'Different'})])
        result = self.apply(renamed_export)
        self.assertEqual(result['counts']['duplicates'], 1)
        self.assertEqual(result['counts']['skippedChanged'], 1)
        self.assertEqual(result['rows'][0]['action'], 'skipped')
        self.assertIn('existing project differs', result['rows'][0]['message'])
        current = next(i for i in self.items() if i['id'] == changed['id'])
        self.assertEqual(current['remarks'], 'Retain this edit')
        self.assertEqual(current['project'], 'New import project')

    def test_invalid_dates_preview_and_conflicting_duplicate_apply_are_atomic(self):
        for blob in (
            workbook([row('Valid project'), row('Bad date', **{'Closing Date': '31/02/2026'})]),
            workbook([row('Conflict'), row('Conflict', Comments='Different data')]),
        ):
            preview = self.preview(blob)
            self.assertEqual(preview['counts']['errors'], 1)
            with self.assertRaisesRegex(ValidationError, 'Correct the highlighted'):
                import_workbook(self.store, 'import.xlsx', blob, self.actor, preview['token'])
            self.assertEqual(len(self.items()), 0)

    def test_formula_without_cache_in_required_field_is_rejected(self):
        blob = workbook([row()])
        with ZipFile(io.BytesIO(blob)) as source:
            root = ET.fromstring(source.read('xl/worksheets/sheet1.xml'))
        ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        target = next(c for c in root.iter(ns + 'c') if c.attrib.get('r') == 'B6')
        target.clear()
        target.set('r', 'B6')
        ET.SubElement(target, ns + 'f').text = '"Missing cached result"'
        bad = replace_parts(blob, {'xl/worksheets/sheet1.xml': ET.tostring(root)})
        preview = self.preview(bad)
        self.assertEqual(preview['counts']['errors'], 1)
        self.assertIn('formula has no valid saved result', preview['issues'][0]['message'])

    def test_supplied_layout_ignores_blank_checkbox_rows_and_remaining_formula(self):
        headers = ['Added Date\nmm/dd/yyyy', 'Closing Date\nmm/dd/yyyy', 'Remaining\nTime', 'Project Name', 'Category', 'Budget', 'Distance from CBD, Melbourne', 'Builder 1', 'Builder 2', 'Builder 3', 'Comments', 'Approved', 'not interested']
        values = [
            ['08/18/2026', '08/31/2026', None, 'Source layout', 'Fit-out', '5m-10m', '1km', 'Builder A', 'Builder B', None, 'Scope note', 'TRUE', 'FALSE'],
            [None, None, 'Expired', None, None, None, None, None, None, None, None, 'FALSE', 'FALSE'],
        ]
        parsed = parse_workbook('renamed.xlsx', workbook(values, headers))
        self.assertEqual(len(parsed['records']), 1)
        self.assertFalse(parsed['issues'])
        item = parsed['records'][0]
        self.assertEqual((item['region'], item['addedDate'], item['due']), ('VIC', '2026-08-18', '2026-08-31'))
        self.assertEqual(item['sourceStatus'], 'approved')
        self.assertEqual(item['distanceKm'], 1)

    def test_uploaded_source_decision_does_not_bypass_approval(self):
        self.apply(workbook([row(Status='Approved')]))
        item = next(i for i in self.items() if i['project'] == 'New import project')
        self.assertEqual(item['sourceStatus'], 'approved')
        self.assertEqual(item['status'], 'pending')
        self.assertIsNone(item.get('decisionBy'))

    def test_builder_company_punctuation_and_long_names_survive_roundtrip(self):
        names = ['Company; a division | NSW, Pty Ltd', 'Long Company ' + 'A' * 230]
        blob = workbook([row(Builders='\n'.join(names))])
        self.apply(blob)
        item = next(i for i in self.items() if i['project'] == 'New import project')
        self.assertEqual(item['builders'], names)
        parsed = parse_workbook('export.xlsx', export_workbook(self.store, ids=[item['id']]))
        self.assertFalse(parsed['issues'])
        self.assertEqual(parsed['records'][0]['builders'], names)

    def test_changed_existing_region_without_origin_is_skipped(self):
        original = self.create()
        blob = workbook([row(), row(original['project'], Region='QLD', **{'Opportunity ID': original['id'], 'Distance Origin': ''})])
        preview = self.preview(blob)
        self.assertEqual(preview['counts']['errors'], 0)
        self.assertEqual(preview['counts']['skippedChanged'], 1)
        self.apply(blob)
        self.assertEqual(len(self.items()), 2)
        self.assertEqual(next(i for i in self.items() if i['id'] == original['id'])['region'], 'NSW')

    def test_new_region_origin_and_metres_are_validated(self):
        bad = workbook([row(Region='QLD', **{'Distance Origin': ''})])
        self.assertEqual(self.preview(bad)['counts']['errors'], 1)
        blob = workbook([
            row(Region='QLD', **{'Distance Origin': 'Brisbane CBD', 'Distance (km)': '1500m'}),
            row('Another QLD project', Region='QLD', **{'Distance Origin': ''}),
        ])
        self.apply(blob)
        item = next(i for i in self.items() if i['region'] == 'QLD')
        self.assertEqual(item['distanceKm'], 1.5)
        self.assertEqual(item['distanceOrigin'], 'Brisbane CBD')
        self.assertEqual(sum(i['region'] == 'QLD' and i['distanceOrigin'] == 'Brisbane CBD' for i in self.items()), 2)
        bad_distance = workbook([row('Negative distance', **{'Distance (km)': '-4km'})])
        self.assertEqual(self.preview(bad_distance)['counts']['errors'], 1)

    def test_token_rejects_data_changes_but_not_notification_reads(self):
        blob = workbook([row()])
        preview = self.preview(blob)
        opportunities.mark_read(self.store, {'ids': []}, self.actor)
        result = import_workbook(self.store, 'import.xlsx', blob, self.actor, preview['token'])
        self.assertTrue(result['applied'])
        second = workbook([row('Second import')])
        preview = self.preview(second)
        item = next(i for i in self.items() if i['region'] == 'NSW')
        opportunities.save(self.store, {**item, 'remarks': 'Concurrent edit'}, self.actor)
        with self.assertRaises(ValidationError) as caught:
            import_workbook(self.store, 'import.xlsx', second, self.actor, preview['token'])
        self.assertEqual(caught.exception.status, 409)

    def test_import_summary_attribution_and_skip_count(self):
        original = self.create()
        blob = workbook([row(), row(original['project'], Region=original['region'])])
        self.apply(blob)
        with self.store.connect() as db:
            activity = db.execute("SELECT * FROM tender_activity WHERE action='Imported opportunities Excel' ORDER BY id DESC LIMIT 1").fetchone()
        details = json.loads(activity['details'])
        self.assertEqual(activity['actor_id'], self.actor['id'])
        self.assertTrue(details['opportunityImport'])
        self.assertEqual((details['added'], details['skipped']), (1, 1))

    def test_filtered_http_export_matches_register(self):
        self.create('Fictional office fit-out')
        self.create('Fictional warehouse', due='2026-10-01')
        self.create('Fictional VIC centre', region='VIC')
        for query, expected in [('region=NSW', 2), ('status=expired', 1), ('q=fit-out', 1)]:
            result = []
            path = '/api/tenders/opportunities/export.xlsx'
            handler = SimpleNamespace(path=path + '?' + query, command='GET', user=self.actor)
            handled = opportunity_http.route(handler, path, self.store, lambda h, body, *args, **kwargs: result.append(body))
            self.assertTrue(handled)
            parsed = parse_workbook('filtered.xlsx', result[0])
            self.assertEqual(len(parsed['records']), expected)
            self.assertFalse(parsed['issues'])

    def test_malformed_uploads_are_explained_without_mutation(self):
        valid = workbook([row()])
        unsupported_compression = bytearray(valid)
        struct.pack_into('<H', unsupported_compression, valid.index(b'PK\x03\x04') + 8, 999)
        struct.pack_into('<H', unsupported_compression, valid.index(b'PK\x01\x02') + 10, 999)
        blobs = [b'not a workbook', valid[:90], replace_parts(valid, {'xl/workbook.xml': b'<broken>'}),
                 bytes(unsupported_compression)]
        for blob in blobs:
            with self.assertRaises(ValidationError):
                self.preview(blob)
        self.assertEqual(len(self.items()), 0)

    def test_sparse_values_beyond_row_limit_are_not_silently_dropped(self):
        blob = workbook([row()])
        with ZipFile(io.BytesIO(blob)) as source:
            sheet = source.read('xl/worksheets/sheet1.xml')
        tail = b'<row r="20000"><c r="B20000" t="inlineStr"><is><t>Far below table</t></is></c></row>'
        oversized = replace_parts(blob, {'xl/worksheets/sheet1.xml': sheet.replace(b'</sheetData>', tail + b'</sheetData>')})
        with self.assertRaisesRegex(ValidationError, 'below row'):
            self.preview(oversized)
        self.assertEqual(len(self.items()), 0)

    def test_contact_fields_import_export_and_changed_duplicate_are_preserved(self):
        blob = workbook([row('Fictional contact job', **{'Contact person': 'Taylor Test', 'Email': 'taylor@example.test'})])
        self.apply(blob)
        item = self.items()[0]
        self.assertEqual((item['contact'], item['email']), ('Taylor Test', 'taylor@example.test'))
        parsed = parse_workbook('export.xlsx', export_workbook(self.store))
        self.assertEqual((parsed['records'][0]['contact'], parsed['records'][0]['email']), ('Taylor Test', 'taylor@example.test'))
        changed = workbook([row('Fictional contact job', **{'Contact person': 'Different contact', 'Email': 'different@example.test'})])
        result = self.apply(changed)
        self.assertEqual(result['counts']['skippedChanged'], 1)
        self.assertIn('contact', result['rows'][0]['changedFields'])
        self.assertIn('email', result['rows'][0]['changedFields'])
        self.assertEqual(self.items()[0]['email'], 'taylor@example.test')

    def test_invalid_contact_email_is_reported_before_apply(self):
        for fields in ({'Email': 'not-an-email'}, {'Contact person': 'A' * 501}):
            blob = workbook([row(**fields)])
            preview = self.preview(blob)
            self.assertEqual(preview['counts']['errors'], 1)
            with self.assertRaisesRegex(ValidationError, 'Correct the highlighted'):
                import_workbook(self.store, 'import.xlsx', blob, self.actor, expected_token=preview['token'])
        self.assertEqual(self.items(), [])

    def test_packaged_template_is_blank_and_uses_supported_columns(self):
        blob = export_workbook(None, template=True)
        book = load_workbook(io.BytesIO(blob))
        self.assertEqual(book.sheetnames, ['Opportunities', 'Instructions'])
        sheet = book['Opportunities']
        self.assertEqual(sheet.freeze_panes, 'C8')
        self.assertEqual(sheet.tables['OpportunityEntries'].ref, 'A7:O107')
        self.assertEqual(sheet['C8'].number_format, 'dd/mm/yyyy')
        self.assertEqual(sheet['N8'].number_format, 'dd/mm/yyyy')
        self.assertEqual(sheet['L7'].value, 'Contact person')
        self.assertEqual(sheet['M7'].value, 'Email')
        self.assertTrue(all(cell.value is None for cells in sheet.iter_rows(min_row=8, max_row=107, max_col=15) for cell in cells))
        self.assertEqual(len(sheet.data_validations.dataValidation), 5)
        book.close()
        with self.assertRaisesRegex(ValidationError, 'No opportunity rows found'):
            parse_workbook('template.xlsx', blob)
        result = []
        path = '/api/tenders/opportunities/template.xlsx'
        handler = SimpleNamespace(path=path, command='GET', user=self.actor)
        self.assertTrue(opportunity_http.route(handler, path, self.store, lambda h, body, *args, **kwargs: result.append(body)))
        self.assertEqual(result[0], blob)


if __name__ == '__main__':
    unittest.main()
