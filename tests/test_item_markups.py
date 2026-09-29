"""Independent item exclusions: arithmetic, formulas, persistence and upgrades."""
import itertools, re, ast, unittest
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter
import test_hours_markup as helpers
from conest_app import projects, productivity, estimate_tools, auto_upgrade, light_prices
from conest_app.database import connect, database_path
from conest_app.release_v31_item_markups import FIELDS
from conest_app.exporter import export_project

class ItemMarkupTests(unittest.TestCase):
    setUp=helpers.HoursMarkupTests.setUp
    tearDown=helpers.HoursMarkupTests.tearDown

    def test_all_combinations_and_other_lines_unchanged(self):
        with connect() as c:
            pid=self.data['project_id'];lid=self.data['line_id']
            projects.update_project(c,pid,{'labor_hours_markup':.2})
            baseline=projects.calculate_project(c,pid);other=baseline['lines'][1]
            for bits in itertools.product([False,True],repeat=3):
                result=projects.update_project_line(c,pid,lid,dict(zip(FIELDS,bits)))
                line=result['lines'][0]
                material=770*(1 if bits[0] else 1.1)
                hours=2.8*(1 if bits[1] else 1.2)
                labor=hours*90+210
                self.assertAlmostEqual(line['total_material'],material)
                self.assertAlmostEqual(line['total_hours'],hours)
                self.assertAlmostEqual(line['total_labor'],labor)
                self.assertAlmostEqual(line['quoted_total'],(material+labor)*(1 if bits[2] else 1.05))
                self.assertEqual(result['lines'][1],other)
                self.assertEqual(projects.project_totals(c,pid),result['totals'])
            for bad in [1,'false',None,[],{}]:
                with self.assertRaises(ValueError):projects.update_project_line(c,pid,lid,{'exclude_material_markup':bad})
            self.assertEqual(projects.calculate_project(c,pid)['lines'][0]['exclude_final_markup'],1)

    def test_manual_save_copy_revision_and_catalog_defaults(self):
        with connect() as c:
            pid=self.data['project_id'];lid=self.data['line_id'];before=projects.calculate_project(c,pid)
            from copy import deepcopy
            legacy=deepcopy(before)
            for row in legacy['lines']:
                for key in FIELDS:row.pop(key)
            self.assertEqual(productivity.compare_snapshots(legacy,before)['changes'],[])
            line=before['lines'][0]
            item={'name':line['name_snapshot'],'category_override':line['category_snapshot'],'work_type':line['work_type_snapshot'],'base_uom':line['uom'],'components':line['components']}
            after=projects.save_manual_estimate_item(c,pid,{'save_mode':'estimate','quantity':7,'item':item,**dict.fromkeys(FIELDS,True)},lid)
            rid=productivity.save_revision(c,pid,self.data['owner'],{'label':'Excluded'})
            snap=productivity.revision_snapshot(c,pid,rid,self.data['owner'])
            self.assertTrue(all(snap['lines'][0][key] for key in FIELDS))
            diff=productivity.compare_snapshots(before,after)['changes'][0]['fields']
            self.assertTrue(set(FIELDS)<=set(diff))
            copied=estimate_tools.copy_project(c,pid,self.data['owner'],{'name':'Excluded copy'})
            self.assertEqual(projects.calculate_project(c,copied)['totals'],after['totals'])
            new=estimate_tools.copy_line(c,pid,lid,{})
            row=next(x for x in projects.calculate_project(c,pid)['lines'] if x['id']==new)
            self.assertTrue(all(row[k] for k in FIELDS));self.assertEqual(row['quoted_total'],after['lines'][0]['quoted_total'])
            added=projects.add_catalog_line(c,pid,self.data['catalog_id'],1)['lines'][-1]
            self.assertTrue(all(not added[k] for k in FIELDS))
            projects.update_project_line(c,pid,lid,{'notes':'Retain choices'})
            self.assertTrue(all(projects.calculate_project(c,pid)['lines'][0][k] for k in FIELDS))

    def test_schema27_upgrade_preserves_totals(self):
        with connect() as c:
            pid=self.data['project_id'];before=projects.calculate_project(c,pid)['totals']
            for field in FIELDS:c.execute(f'ALTER TABLE project_lines DROP COLUMN {field}')
            c.execute('UPDATE schema_version SET version=27')
            key=c.execute('PRAGMA table_info(app_migrations)').fetchone()[1]
            c.execute(f'DELETE FROM app_migrations WHERE "{key}"=?',('v3.1-item-markups',))
        auto_upgrade._ready.pop(str(database_path()),None)
        with connect() as c:
            after=projects.calculate_project(c,pid)
            self.assertEqual(after['totals'],before)
            self.assertTrue(all(not row[field] for row in after['lines'] for field in FIELDS))
            self.assertEqual(c.execute('SELECT version FROM schema_version').fetchone()[0],28)
        self.assertTrue(list((database_path().parent/'upgrade_backups').glob('*before-schema-28.db')))

    def test_export_arithmetic_with_mixed_exclusions_and_fixed_quantity(self):
        def evaluate(w,sheet,cell):
            raw=w[sheet][cell].value
            if not isinstance(raw,str) or not raw.startswith('='):return raw or 0
            expr=raw[1:].replace('$','')
            def rng(m):
                a,b,x,y=range_boundaries(m[1]);return str(sum(evaluate(w,sheet,f'{get_column_letter(i)}{j}') for i in range(a,x+1) for j in range(b,y+1)))
            expr=re.sub(r'SUM\(([A-Z]+\d+:[A-Z]+\d+)\)',rng,expr)
            expr=re.sub(r"'([^']+)'!([A-Z]+\d+)",lambda m:str(evaluate(w,m[1],m[2])),expr)
            expr=re.sub(r'\b([A-Z]+\d+)\b',lambda m:str(evaluate(w,sheet,m[1])),expr)
            expr=re.sub(r'SUM\(([^()]*)\)',lambda m:'('+m[1].replace(',','+')+')',expr)
            tree=ast.parse(expr,mode='eval')
            self.assertTrue(all(isinstance(n,(ast.Expression,ast.Constant,ast.BinOp,ast.UnaryOp,ast.Add,ast.Sub,ast.Mult,ast.Div,ast.USub,ast.UAdd)) for n in ast.walk(tree)),expr)
            return eval(compile(tree,'formula','eval'),{'__builtins__':{}})
        with connect() as c:
            pid=self.data['project_id'];lid=self.data['line_id']
            projects.update_project(c,pid,{'labor_hours_markup':.25})
            c.execute('UPDATE project_line_components SET fixed_quantity=1 WHERE project_line_id=? AND labor_pricing=?',(lid,'fixed'))
            for bits in itertools.product([False,True],repeat=3):
                expected=projects.update_project_line(c,pid,lid,dict(zip(FIELDS,bits)))['totals']
                path=export_project(c,pid,Path(self.tmp.name)/'mixed.xlsx')
                w=load_workbook(path)
                for cell,key in [('M5','material'),('M6','labor'),('M7','hours'),('M8','direct'),('M9','quoted')]:
                    self.assertAlmostEqual(evaluate(w,'COST ESTIMATE',cell),expected[key],places=6)
                w.close()

    def test_multiple_suppliers_union_component_match_and_pagination(self):
        with connect() as c:
            for i,supplier in enumerate(['Second Supplier','Third Supplier']):
                q={'reference':str(i),'quote_date':'2026-09-20','project':'Other','source':'manual','total_ex':10,'lines':[light_prices.normal_line({'Type':'L1','Supplier':supplier,'Qty':1,'Description':'Test light','Unit Ex GST':10,'Total Ex GST':10},1)]}
                light_prices.insert_quote(c,q,light_prices.quote_digest(q))
        rows=light_prices.search(supplier=['Example Supplier','Second Supplier'])
        self.assertEqual(rows['total'],2)
        self.assertEqual(light_prices.search(supplier=['example supplier','SECOND SUPPLIER'],limit=1,offset=1)['total'],2)
        self.assertEqual(len(light_prices.search(supplier=['Example Supplier','Second Supplier'],limit=1,offset=1)['results']),1)
        self.assertEqual(light_prices.search('Orchid',supplier=['Example Supplier','Second Supplier'])['total'],1)
        with connect() as c:c.execute("UPDATE lines SET supplier='Component Vendor' WHERE group_id=?",(self.data['light_id'],))
        self.assertEqual(light_prices.search(supplier=['Component Vendor','Second Supplier'])['total'],2)
